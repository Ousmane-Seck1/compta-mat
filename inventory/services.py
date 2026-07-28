from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from django.contrib.auth import get_user_model
from django.db import IntegrityError
from django.db import transaction
from django.db.models import Max, Prefetch

from openpyxl import load_workbook

from .forms import InternalMovementLineInput, VoucherLineInput
from .models import (
    FiscalYear,
    InternalMovement,
    InternalMovementLine,
    Location,
    Material,
    NomenclatureItem,
    PhysicalInventory,
    SiteSetting,
    Structure,
    Voucher,
    VoucherLine,
)


User = get_user_model()
ZERO = Decimal("0")


@dataclass(slots=True)
class ReportRow:
    account_code: str
    name: str
    unit: str
    group_code: str
    qty_in: Decimal
    amount_in: Decimal
    qty_in_opening: Decimal
    amount_in_opening: Decimal
    qty_in_period: Decimal
    amount_in_period: Decimal
    qty_out: Decimal
    amount_out: Decimal
    qty_out_definitive: Decimal
    qty_out_provisional: Decimal
    pending_qty: Decimal
    in_service_qty: Decimal
    provisional_qty: Decimal
    pv_is_filled: bool
    inventory_total_qty: Decimal
    inventory_unit_price: Decimal
    remaining_qty: Decimal
    remaining_amount: Decimal
    cmup: Decimal
    inventory_gap: Decimal
    gap_plus: Decimal
    gap_minus: Decimal
    inventory_gap_value: Decimal


@dataclass(slots=True)
class LocationInventoryRow:
    material_id: int
    account_code: str
    material_name: str
    inventory_code: str
    quantity: Decimal
    unit: str
    unit_price: Decimal
    amount: Decimal
    last_movement_date: date | None


def _clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip()


def _floatish(value: object) -> Decimal:
    if isinstance(value, bool) or value is None or value == "":
        return ZERO
    return Decimal(str(value))


def _scoped_setting_key(key: str, structure: Structure | None = None) -> str:
    if structure is None:
        return key
    return f"structure:{structure.id}:{key}"


def get_setting(key: str, default: str = "", structure: Structure | None = None) -> str:
    setting = SiteSetting.objects.filter(key=_scoped_setting_key(key, structure)).first()
    return setting.value if setting else default


def set_setting(key: str, value: str, structure: Structure | None = None) -> None:
    SiteSetting.objects.update_or_create(key=_scoped_setting_key(key, structure), defaults={"value": value})


def dashboard_metrics(structure: Structure | None = None, fiscal_year: FiscalYear | None = None) -> dict[str, Decimal | int]:
    summary = compute_stock_summary(structure=structure, fiscal_year=fiscal_year)
    return {
        "materials_count": Material.objects.filter(
            structure=structure, fiscal_year=fiscal_year
        ).count() if structure and fiscal_year else Material.objects.count(),
        "vouchers_count": Voucher.objects.filter(
            structure=structure, fiscal_year=fiscal_year
        ).count() if structure and fiscal_year else Voucher.objects.count(),
        "stock_quantity": sum((item.remaining_qty for item in summary), ZERO),
        "stock_value": sum((item.remaining_amount for item in summary), ZERO),
    }


def next_voucher_number(structure: Structure, fiscal_year: FiscalYear) -> int:
    current = Voucher.objects.filter(
        structure=structure, fiscal_year=fiscal_year
    ).aggregate(max_number=Max("number"))["max_number"] or 0
    return current + 1


def next_internal_movement_number(structure: Structure, fiscal_year: FiscalYear) -> int:
    current = InternalMovement.objects.filter(
        structure=structure,
        fiscal_year=fiscal_year,
    ).aggregate(max_number=Max("number"))["max_number"] or 0
    return current + 1


def _create_voucher_header_with_retry(
    *,
    structure: Structure,
    fiscal_year: FiscalYear,
    voucher_type: str,
    operation_date: date,
    source_or_destination: str,
    storage_location: Location | None,
    comments: str,
    created_by,
) -> Voucher:
    # Retry a few times in case two users generate the same next number concurrently.
    for _attempt in range(3):
        try:
            return Voucher.objects.create(
                structure=structure,
                fiscal_year=fiscal_year,
                number=next_voucher_number(structure, fiscal_year),
                voucher_type=voucher_type,
                operation_date=operation_date,
                source_or_destination=source_or_destination.strip(),
                storage_location=storage_location,
                comments=comments.strip(),
                created_by=created_by,
            )
        except IntegrityError:
            continue
    raise ValueError("Conflit de numerotation detecte. Veuillez reessayer.")


def _create_internal_movement_header_with_retry(
    *,
    structure: Structure,
    fiscal_year: FiscalYear,
    movement_type: str,
    operation_date: date,
    from_location: Location,
    to_location: Location,
    reason: str,
    created_by,
) -> InternalMovement:
    for _attempt in range(3):
        try:
            return InternalMovement.objects.create(
                structure=structure,
                fiscal_year=fiscal_year,
                number=next_internal_movement_number(structure, fiscal_year),
                movement_type=movement_type,
                operation_date=operation_date,
                from_location=from_location,
                to_location=to_location,
                reason=reason.strip(),
                created_by=created_by,
            )
        except IntegrityError:
            continue
    raise ValueError("Conflit de numerotation detecte. Veuillez reessayer.")


def ensure_physical_inventory(material: Material) -> PhysicalInventory:
    inventory, _ = PhysicalInventory.objects.get_or_create(
        material=material,
        defaults={"is_pv_filled": False, "unit_price": Decimal("0")},
    )
    return inventory


def infer_default_store_label(material: Material) -> str:
    group_code = (material.group_code or "").strip().lower()
    if group_code.startswith("1"):
        return "Magasin de materiel et equipement"
    if group_code.startswith("2"):
        return "Magasin de matieres et fournitures"
    return "Magasin de produits"


def _is_consumable_material(material: Material) -> bool:
    """Groupe 2 (nomenclature commençant par '2') : consomptible au premier usage.
    Quand ces matières quittent le magasin elles sont considérées comme consommées,
    contrairement au groupe 1 qui passe en statut 'En service'."""
    return (material.account_code or "").strip().startswith("2")


def infer_entry_destination_for_material_ids(material_ids: list[int]) -> str:
    materials = list(Material.objects.filter(pk__in=material_ids))
    if not materials:
        return "Magasin de rattachement"
    labels = {infer_default_store_label(material) for material in materials}
    if len(labels) == 1:
        return labels.pop()
    return "Magasins de rattachement"


def _account_prefix(material: Material) -> str:
    normalized = (material.account_code or "").replace(" ", "").strip()
    return normalized[:1] if normalized else ""


def _expected_store_location_type(material: Material) -> str:
    prefix = _account_prefix(material)
    if prefix == "1":
        return Location.TYPE_STORE_EQUIPMENT
    if prefix == "2":
        return Location.TYPE_STORE_SUPPLIES
    return Location.TYPE_STORE_PRODUCTS


def _default_store_name_for_type(location_type: str) -> str:
    if location_type == Location.TYPE_STORE_EQUIPMENT:
        return "Magasin de materiel et equipement"
    if location_type == Location.TYPE_STORE_SUPPLIES:
        return "Magasin de matieres et fournitures"
    return "Magasin de produits"


def _resolve_voucher_store_for_materials(
    *,
    materials: list[Material],
    structure: Structure,
) -> Location:
    if not materials:
        raise ValueError("Le bon doit contenir au moins une ligne de matiere.")

    expected_types = {_expected_store_location_type(material) for material in materials}
    if len(expected_types) > 1:
        raise ValueError(
            "Un meme bon ne peut pas melanger des matieres relevant de magasins differents. "
            "Veuillez separer les lignes par magasin (materiel / fournitures / produits)."
        )

    expected_type = expected_types.pop()
    store = Location.objects.filter(
        structure=structure,
        is_active=True,
        location_type=expected_type,
    ).order_by("name").first()
    if store:
        return store

    # Compatibilite et robustesse: si aucun magasin n'existe encore, on cree le magasin attendu.
    default_name = _default_store_name_for_type(expected_type)
    existing_same_name = Location.objects.filter(structure=structure, name=default_name).first()
    if existing_same_name:
        if existing_same_name.location_type != expected_type:
            raise ValueError(
                f"La localisation '{default_name}' existe deja avec un type incompatible. "
                "Corrigez son type ou creez un autre magasin de stockage."
            )
        if not existing_same_name.is_active:
            existing_same_name.is_active = True
            existing_same_name.save(update_fields=["is_active", "updated_at"])
        return existing_same_name

    return Location.objects.create(
        structure=structure,
        name=default_name,
        location_type=expected_type,
        responsible_name="",
        details="Cree automatiquement lors de la saisie d'un bon.",
        is_active=True,
    )


def _apply_entry_to_physical_inventory(material: Material, quantity: Decimal, unit_price: Decimal) -> None:
    inventory = ensure_physical_inventory(material)
    inventory.pending_qty += quantity
    if unit_price >= ZERO:
        inventory.unit_price = unit_price
    inventory.save(update_fields=["pending_qty", "unit_price", "is_pv_filled", "updated_at"])


def _validate_internal_locations(movement_type: str, from_location: Location, to_location: Location) -> None:
    if from_location.structure_id != to_location.structure_id:
        raise ValueError("Les localisations doivent appartenir au meme service.")
    if movement_type == InternalMovement.TYPE_ASSIGNMENT:
        if not from_location.is_store:
            raise ValueError("Une affectation doit partir d'un magasin.")
        if to_location.is_store:
            raise ValueError("Une affectation doit arriver vers une localisation en service.")
    elif movement_type == InternalMovement.TYPE_UNASSIGNMENT:
        if from_location.is_store:
            raise ValueError("Une desaffectation doit partir d'une localisation en service.")
        if not to_location.is_store:
            raise ValueError("Une desaffectation doit revenir vers un magasin.")
    elif movement_type == InternalMovement.TYPE_TRANSFER:
        if from_location.is_store or to_location.is_store:
            raise ValueError("Une mutation doit se faire entre localisations en service.")


def _validate_movement_line_group(
    movement_type: str,
    from_location: Location,
    material: Material,
) -> None:
    """Valide la cohérence groupe matière / type de mouvement / magasin source."""
    is_consumable = _is_consumable_material(material)

    # Groupe 2 (consomptibles) : ne peuvent pas faire l'objet d'une désaffectation ou mutation
    if is_consumable and movement_type in (
        InternalMovement.TYPE_UNASSIGNMENT,
        InternalMovement.TYPE_TRANSFER,
    ):
        raise ValueError(
            f"La matiere '{material.name}' est consomptible (groupe 2 — fournitures) et ne peut "
            f"pas faire l'objet d'une desaffectation ou mutation : elle est consommee au premier usage."
        )

    # Lors d'une affectation, vérifier que le magasin source correspond au groupe
    if movement_type == InternalMovement.TYPE_ASSIGNMENT:
        if (
            from_location.location_type == Location.TYPE_STORE_EQUIPMENT
            and is_consumable
        ):
            raise ValueError(
                f"La matiere '{material.name}' (groupe 2, fournitures) ne peut pas sortir "
                f"du magasin principal. Elle doit provenir du magasin de fournitures."
            )
        if (
            from_location.location_type == Location.TYPE_STORE_SUPPLIES
            and not is_consumable
        ):
            raise ValueError(
                f"La matiere '{material.name}' (groupe 1, materiel) ne peut pas sortir "
                f"du magasin de fournitures. Elle doit provenir du magasin principal."
            )


def list_internal_movements(structure: Structure, fiscal_year: FiscalYear):
    return InternalMovement.objects.filter(
        structure=structure,
        fiscal_year=fiscal_year,
    ).select_related("from_location", "to_location", "created_by").prefetch_related("lines__material")


def _reconcile_pending_qty_from_accounting(
    *,
    inventory: PhysicalInventory,
    material: Material,
    stock_row: ReportRow | None,
) -> None:
    """Rattrape pending_qty si la table physique est en retard par rapport au journal.

    Cas couvert: anciennes donnees avec bons d'entree existants mais pending_qty reste a zero.
    Le pending attendu est approxime par: restant comptable - en_service - provisoire.
    """
    if stock_row is None:
        return

    # Ne pas toucher aux fiches de PV deja valides: elles representent une saisie manuelle de reference.
    if inventory.is_pv_filled:
        return

    expected_pending = stock_row.remaining_qty - inventory.in_service_qty - inventory.provisional_qty
    if expected_pending < ZERO:
        expected_pending = ZERO

    # Ne corriger que dans le sens d'un rattrapage pour eviter d'effacer des corrections manuelles.
    if expected_pending > inventory.pending_qty:
        inventory.pending_qty = expected_pending
        inventory.unit_price = stock_row.cmup if stock_row.cmup > ZERO else inventory.unit_price
        inventory.save(update_fields=["pending_qty", "unit_price", "is_pv_filled", "updated_at"])


@transaction.atomic
def delete_internal_movement(movement: InternalMovement) -> None:
    for line in movement.lines.select_related("material"):
        inventory = ensure_physical_inventory(line.material)
        if movement.movement_type == InternalMovement.TYPE_ASSIGNMENT:
            # Annulation affectation : on remet en magasin
            inventory.pending_qty += line.quantity
            if not _is_consumable_material(line.material):
                # Groupe 1 : retirer de in_service_qty
                inventory.in_service_qty -= line.quantity
            # Groupe 2 (consomptibles) : avaient été consommés, on restitue juste pending_qty
        elif movement.movement_type == InternalMovement.TYPE_UNASSIGNMENT:
            inventory.pending_qty -= line.quantity
            inventory.in_service_qty += line.quantity
        inventory.save(update_fields=["pending_qty", "in_service_qty", "is_pv_filled", "updated_at"])
    movement.delete()


@transaction.atomic
def create_internal_movement(
    *,
    user,
    movement_type: str,
    operation_date: date,
    from_location: Location,
    to_location: Location,
    reason: str,
    lines: list[InternalMovementLineInput],
    structure: Structure,
    fiscal_year: FiscalYear,
) -> InternalMovement:
    if movement_type not in dict(InternalMovement.TYPE_CHOICES):
        raise ValueError("Type de mouvement interne invalide.")
    if fiscal_year.is_closed:
        raise ValueError(f"L'exercice {fiscal_year.year} est cloture. Creation de mouvement impossible.")
    if from_location.structure_id != structure.id or to_location.structure_id != structure.id:
        raise ValueError("Les localisations selectionnees ne correspondent pas au service courant.")
    _validate_internal_locations(movement_type, from_location, to_location)

    stock_rows_by_code = {
        row.account_code: row
        for row in compute_stock_summary(structure=structure, fiscal_year=fiscal_year)
    }

    movement = _create_internal_movement_header_with_retry(
        structure=structure,
        fiscal_year=fiscal_year,
        movement_type=movement_type,
        operation_date=operation_date,
        from_location=from_location,
        to_location=to_location,
        reason=reason,
        created_by=user if getattr(user, "is_authenticated", False) else None,
    )

    for line in lines:
        material = Material.objects.get(
            pk=line.material_id,
            structure=structure,
            fiscal_year=fiscal_year,
        )
        inventory = ensure_physical_inventory(material)

        # Rattrapage auto pour les donnees historiques non synchronisees.
        _reconcile_pending_qty_from_accounting(
            inventory=inventory,
            material=material,
            stock_row=stock_rows_by_code.get(material.account_code),
        )

        # Validation groupe matière vs type de mouvement vs magasin source
        _validate_movement_line_group(movement_type, from_location, material)

        if movement_type == InternalMovement.TYPE_ASSIGNMENT and line.quantity > inventory.pending_qty:
            raise ValueError(f"La quantite a affecter depasse le stock en attente pour {material.name}.")
        if movement_type == InternalMovement.TYPE_UNASSIGNMENT and line.quantity > inventory.in_service_qty:
            raise ValueError(f"La quantite a desaffecter depasse le stock en service pour {material.name}.")
        if movement_type == InternalMovement.TYPE_TRANSFER and line.quantity > inventory.in_service_qty:
            raise ValueError(f"La quantite a muter depasse le stock en service pour {material.name}.")

        if movement_type == InternalMovement.TYPE_ASSIGNMENT:
            inventory.pending_qty -= line.quantity
            if not _is_consumable_material(material):
                # Groupe 1 (matériel) : passe en statut "En service"
                inventory.in_service_qty += line.quantity
            # Groupe 2 (fournitures) : consommé au premier usage, aucun in_service_qty
        elif movement_type == InternalMovement.TYPE_UNASSIGNMENT:
            # Désaffectation : uniquement possible pour groupe 1
            inventory.in_service_qty -= line.quantity
            inventory.pending_qty += line.quantity
        inventory.unit_price = line.unit_price
        inventory.save(update_fields=["pending_qty", "in_service_qty", "unit_price", "is_pv_filled", "updated_at"])

        amount = line.quantity * line.unit_price
        InternalMovementLine.objects.create(
            movement=movement,
            material=material,
            material_name=material.name,
            inventory_code=line.inventory_code,
            quantity=line.quantity,
            unit=material.unit,
            unit_price=line.unit_price,
            amount=amount,
            observations=line.observations,
        )

    return movement


def build_location_inventory_report(
    location: Location,
    fiscal_year: FiscalYear,
    *,
    include_prior_year: bool = True,
) -> list[LocationInventoryRow]:
    if location.is_store:
        def _normalized_account_prefix(account_code: str) -> str:
            normalized = (account_code or "").replace(" ", "").strip()
            return normalized[:1] if normalized else ""

        def _belongs_to_store(material: Material, store_location: Location) -> bool:
            prefix = _normalized_account_prefix(material.account_code)
            if store_location.location_type == Location.TYPE_STORE_EQUIPMENT:
                return prefix == "1"
            if store_location.location_type == Location.TYPE_STORE_SUPPLIES:
                return prefix == "2"
            if store_location.location_type == Location.TYPE_STORE_PRODUCTS:
                return prefix not in {"1", "2"}
            return False

        rows: dict[tuple[int, str], LocationInventoryRow] = {}

        voucher_lines = VoucherLine.objects.filter(
            voucher__structure=location.structure,
            voucher__fiscal_year=fiscal_year,
            voucher__storage_location=location,
        ).select_related("voucher", "material")

        for line in voucher_lines.order_by("voucher__operation_date", "voucher__number", "id"):
            delta = ZERO
            amount_delta = ZERO
            if line.voucher.voucher_type == Voucher.TYPE_ENTRY:
                delta = line.quantity
                amount_delta = line.amount
            elif line.voucher.voucher_type in {Voucher.TYPE_FINAL_EXIT, Voucher.TYPE_TEMP_EXIT}:
                delta = -line.quantity
                amount_delta = -line.amount

            if delta == ZERO:
                continue

            key = (line.material_id, "")
            row = rows.get(key)
            if row is None:
                row = LocationInventoryRow(
                    material_id=line.material_id,
                    account_code=line.material.account_code,
                    material_name=line.material_name,
                    inventory_code="",
                    quantity=ZERO,
                    unit=line.unit,
                    unit_price=line.unit_price,
                    amount=ZERO,
                    last_movement_date=line.voucher.operation_date,
                )
                rows[key] = row

            row.quantity += delta
            row.amount += amount_delta
            if row.quantity > ZERO:
                row.unit_price = row.amount / row.quantity
            if row.last_movement_date is None or line.voucher.operation_date > row.last_movement_date:
                row.last_movement_date = line.voucher.operation_date

        movement_lines = InternalMovementLine.objects.filter(
            movement__structure=location.structure,
            movement__fiscal_year=fiscal_year,
            movement__from_location_id=location.id,
        ).select_related("movement", "material")

        for line in movement_lines.order_by("movement__operation_date", "movement__number", "id"):
            key = (line.material_id, line.inventory_code or "")
            row = rows.get(key)
            if row is None:
                row = LocationInventoryRow(
                    material_id=line.material_id,
                    account_code=line.material.account_code,
                    material_name=line.material_name,
                    inventory_code=line.inventory_code,
                    quantity=ZERO,
                    unit=line.unit,
                    unit_price=line.unit_price,
                    amount=ZERO,
                    last_movement_date=line.movement.operation_date,
                )
                rows[key] = row

            row.quantity -= line.quantity
            row.amount -= line.amount
            if row.quantity > ZERO:
                row.unit_price = row.amount / row.quantity
            if row.last_movement_date is None or line.movement.operation_date > row.last_movement_date:
                row.last_movement_date = line.movement.operation_date

        incoming_movement_lines = InternalMovementLine.objects.filter(
            movement__structure=location.structure,
            movement__fiscal_year=fiscal_year,
            movement__to_location_id=location.id,
        ).select_related("movement", "material")

        for line in incoming_movement_lines.order_by("movement__operation_date", "movement__number", "id"):
            key = (line.material_id, line.inventory_code or "")
            row = rows.get(key)
            if row is None:
                row = LocationInventoryRow(
                    material_id=line.material_id,
                    account_code=line.material.account_code,
                    material_name=line.material_name,
                    inventory_code=line.inventory_code,
                    quantity=ZERO,
                    unit=line.unit,
                    unit_price=line.unit_price,
                    amount=ZERO,
                    last_movement_date=line.movement.operation_date,
                )
                rows[key] = row

            row.quantity += line.quantity
            row.amount += line.amount
            if row.quantity > ZERO:
                row.unit_price = row.amount / row.quantity
            if row.last_movement_date is None or line.movement.operation_date > row.last_movement_date:
                row.last_movement_date = line.movement.operation_date

        # Fallback historique: conserver l'ancien stock deja saisi dans la fiche physique
        # pour les matieres qui n'ont pas encore de bon/localisation explicite.
        stock_rows_by_code = {
            row.account_code: row
            for row in compute_stock_summary(structure=location.structure, fiscal_year=fiscal_year)
        }
        materials = Material.objects.filter(
            structure=location.structure,
            fiscal_year=fiscal_year,
        ).prefetch_related("physical_inventory")

        for material in materials:
            if not _belongs_to_store(material, location):
                continue

            inventory = getattr(material, "physical_inventory", None)
            stock_row = stock_rows_by_code.get(material.account_code)
            if inventory and inventory.is_pv_filled:
                pending_qty = inventory.pending_qty
            else:
                # In magasin view, fallback must reflect only "en attente" stock,
                # not the whole accounting stock of the structure.
                pending_qty = stock_row.pending_qty if stock_row else ZERO

            if pending_qty <= ZERO:
                continue

            # Quantite deja localisee pour cette matiere dans ce magasin
            # (lignes normales + mouvements internes), potentiellement partielle.
            current_qty = sum(
                (existing_row.quantity for (material_id, _), existing_row in rows.items() if material_id == material.id),
                ZERO,
            )
            # Ne pas melanger les anciens stocks avec des lignes deja localisees.
            if current_qty > ZERO:
                continue
            delta_qty = pending_qty

            unit_price = ZERO
            if inventory and inventory.unit_price > ZERO:
                unit_price = inventory.unit_price
            elif stock_row and stock_row.cmup > ZERO:
                unit_price = stock_row.cmup

            key = (material.id, "")
            row = rows.get(key)
            if row is None:
                row = LocationInventoryRow(
                    material_id=material.id,
                    account_code=material.account_code,
                    material_name=material.name,
                    inventory_code="",
                    quantity=ZERO,
                    unit=material.unit,
                    unit_price=unit_price,
                    amount=ZERO,
                    last_movement_date=None,
                )
                rows[key] = row

            row.quantity += delta_qty
            row.amount += delta_qty * unit_price
            if row.quantity > ZERO:
                row.unit_price = row.amount / row.quantity

        return sorted(
            [row for row in rows.values() if row.quantity > ZERO],
            key=lambda item: (item.account_code, item.inventory_code, item.material_name),
        )

    movement_lines = InternalMovementLine.objects.filter(
        movement__structure=location.structure,
        movement__fiscal_year=fiscal_year,
    ).select_related("movement", "material", "movement__from_location", "movement__to_location")

    rows: dict[tuple[int, str], LocationInventoryRow] = {}
    for line in movement_lines:
        delta = ZERO
        if line.movement.to_location_id == location.id:
            delta += line.quantity
        if line.movement.from_location_id == location.id:
            delta -= line.quantity
        if delta == ZERO:
            continue

        key = (line.material_id, line.inventory_code or "")
        row = rows.get(key)
        if row is None:
            row = LocationInventoryRow(
                material_id=line.material_id,
                account_code=line.material.account_code,
                material_name=line.material_name,
                inventory_code=line.inventory_code,
                quantity=ZERO,
                unit=line.unit,
                unit_price=line.unit_price,
                amount=ZERO,
                last_movement_date=line.movement.operation_date,
            )
            rows[key] = row
        row.quantity += delta
        row.amount += line.amount if delta > ZERO else -line.amount
        if row.quantity > ZERO:
            row.unit_price = row.amount / row.quantity
        if row.last_movement_date is None or line.movement.operation_date > row.last_movement_date:
            row.last_movement_date = line.movement.operation_date

    if include_prior_year:
        has_report = InternalMovementLine.objects.filter(
            movement__structure=location.structure,
            movement__fiscal_year=fiscal_year,
            movement__to_location_id=location.id,
            observations="Report affectation exercice precedent",
        ).exists()
        if not has_report:
            previous_fy = FiscalYear.objects.filter(
                structure=location.structure,
                year=fiscal_year.year - 1,
            ).first()
            if previous_fy:
                material_map = {
                    material.account_code: material
                    for material in Material.objects.filter(
                        structure=location.structure,
                        fiscal_year=fiscal_year,
                    )
                }
                previous_rows = build_location_inventory_report(
                    location,
                    previous_fy,
                    include_prior_year=False,
                )
                for prev_row in previous_rows:
                    new_material = material_map.get(prev_row.account_code)
                    if new_material is None:
                        continue
                    key = (new_material.id, prev_row.inventory_code or "")
                    row = rows.get(key)
                    if row is None:
                        row = LocationInventoryRow(
                            material_id=new_material.id,
                            account_code=new_material.account_code,
                            material_name=new_material.name,
                            inventory_code=prev_row.inventory_code,
                            quantity=ZERO,
                            unit=new_material.unit,
                            unit_price=prev_row.unit_price,
                            amount=ZERO,
                            last_movement_date=prev_row.last_movement_date or previous_fy.end_date,
                        )
                        rows[key] = row
                    row.quantity += prev_row.quantity
                    row.amount += prev_row.amount
                    if row.quantity > ZERO:
                        row.unit_price = row.amount / row.quantity
                    if prev_row.last_movement_date and (
                        row.last_movement_date is None or prev_row.last_movement_date > row.last_movement_date
                    ):
                        row.last_movement_date = prev_row.last_movement_date

    return sorted(
        [row for row in rows.values() if row.quantity > ZERO],
        key=lambda item: (item.account_code, item.inventory_code, item.material_name),
    )


@transaction.atomic
def propagate_nomenclature_add(item: NomenclatureItem) -> int:
    """Create the corresponding Material in all active structures x open fiscal years.

    If the item has a structure_type, only targets structures of that type.
    If structure_type is null, targets all active structures (global account).
    """
    count = 0
    structures_qs = Structure.objects.filter(is_active=True)
    if item.structure_type_id is not None:
        structures_qs = structures_qs.filter(structure_type_id=item.structure_type_id)
    for structure in structures_qs:
        for fy in FiscalYear.objects.filter(structure=structure, is_closed=False):
            _, created = Material.objects.get_or_create(
                structure=structure,
                fiscal_year=fy,
                account_code=item.account_code,
                defaults={"name": item.name, "unit": item.unit, "group_code": item.group_code},
            )
            if created:
                count += 1
    return count


@transaction.atomic
def propagate_nomenclature_delete(account_code: str) -> None:
    """Remove all Materials with this account_code if none has VoucherLine references."""
    if VoucherLine.objects.filter(material__account_code=account_code).exists():
        raise ValueError(
            f"Impossible de supprimer '{account_code}' : des bons referencent ce compte dans certains services."
        )
    Material.objects.filter(account_code=account_code).delete()


@transaction.atomic
def sync_all_nomenclature_to_structure(structure: Structure) -> int:
    """Ensure every applicable NomenclatureItem exists as Material for all open FYs of this structure.

    Applies global items (structure_type=null) + items matching structure.structure_type.
    """
    from django.db.models import Q
    count = 0
    nomenclature_qs = NomenclatureItem.objects.filter(
        Q(structure_type__isnull=True) | Q(structure_type=structure.structure_type)
    )
    for fy in FiscalYear.objects.filter(structure=structure, is_closed=False):
        for item in nomenclature_qs:
            mat, created = Material.objects.get_or_create(
                structure=structure,
                fiscal_year=fy,
                account_code=item.account_code,
                defaults={"name": item.name, "unit": item.unit, "group_code": item.group_code},
            )
            if created:
                count += 1
            else:
                # Update fields if they differ from nomenclature
                updated = False
                if mat.name != item.name:
                    mat.name = item.name
                    updated = True
                if mat.unit != item.unit:
                    mat.unit = item.unit
                    updated = True
                if mat.group_code != item.group_code:
                    mat.group_code = item.group_code
                    updated = True
                if updated:
                    mat.save(update_fields=["name", "unit", "group_code"])
    return count


@transaction.atomic
@transaction.atomic
def create_voucher(
    *,
    user,
    voucher_type: str,
    operation_date: date,
    source_or_destination: str,
    storage_location: Location | None = None,
    comments: str,
    lines: list[VoucherLineInput],
    structure: Structure,
    fiscal_year: FiscalYear,
    apply_to_physical_inventory: bool = True,
) -> list[Voucher]:
    if voucher_type not in dict(Voucher.TYPE_CHOICES):
        raise ValueError("Type de bon invalide.")
    if fiscal_year.is_closed:
        raise ValueError(f"L'exercice {fiscal_year.year} est cloture. Creation de bon impossible.")
    material_ids = [line.material_id for line in lines]
    materials = list(
        Material.objects.filter(
            pk__in=material_ids,
            structure=structure,
            fiscal_year=fiscal_year,
        )
    )
    materials_by_id = {material.id: material for material in materials}
    if len(materials_by_id) != len(set(material_ids)):
        raise ValueError("Une ou plusieurs matieres du bon sont invalides pour le service/exercice courant.")

    lines_by_store_type: dict[str, list[VoucherLineInput]] = {}
    for line in lines:
        material = materials_by_id[line.material_id]
        prefix = _account_prefix(material)
        if prefix == "1":
            store_type = Location.TYPE_STORE_EQUIPMENT
        elif prefix == "2":
            store_type = Location.TYPE_STORE_SUPPLIES
        else:
            raise ValueError(
                f"La matiere '{material.name}' ne correspond pas aux magasins definis (codes 1 ou 2)."
            )
        lines_by_store_type.setdefault(store_type, []).append(line)

    if storage_location is not None and len(lines_by_store_type) > 1:
        raise ValueError(
            "Le bon contient des matieres de magasins differents. "
            "Laissez le magasin vide pour une affectation automatique."
        )

    created_vouchers: list[Voucher] = []

    def _create_single_voucher(group_lines: list[VoucherLineInput]) -> Voucher:
        group_materials = [materials_by_id[item.material_id] for item in group_lines]
        effective_source_or_destination = source_or_destination
        effective_storage_location: Location | None = storage_location

        if voucher_type in {Voucher.TYPE_ENTRY, Voucher.TYPE_FINAL_EXIT, Voucher.TYPE_TEMP_EXIT}:
            # Regle metier: ces bons proviennent toujours d'un magasin de stockage.
            # On valide (et cree si besoin) le magasin rattache aux matieres de la ligne.
            expected_store = _resolve_voucher_store_for_materials(materials=group_materials, structure=structure)

            if effective_storage_location is None:
                effective_storage_location = expected_store
            else:
                if effective_storage_location.structure_id != structure.id:
                    raise ValueError("Le magasin selectionne n'appartient pas au service courant.")
                if not effective_storage_location.is_store:
                    raise ValueError("Le lieu selectionne doit etre un magasin (en attente d'affectation).")
                if not effective_storage_location.is_active:
                    raise ValueError("Le magasin selectionne est inactif.")
                if effective_storage_location.location_type != expected_store.location_type:
                    raise ValueError(
                        "Le magasin selectionne ne correspond pas au type attendu pour les matieres du bon."
                    )

            if voucher_type != Voucher.TYPE_ENTRY and not effective_source_or_destination.strip():
                # Pour les sorties, conserver la destination si l'utilisateur la renseigne,
                # sinon remplir le magasin de provenance.
                effective_source_or_destination = effective_storage_location.name

        voucher = _create_voucher_header_with_retry(
            structure=structure,
            fiscal_year=fiscal_year,
            voucher_type=voucher_type,
            operation_date=operation_date,
            source_or_destination=effective_source_or_destination,
            storage_location=effective_storage_location,
            comments=comments,
            created_by=user if getattr(user, "is_authenticated", False) else None,
        )
        for line in group_lines:
            material = materials_by_id[line.material_id]
            summary = get_material_stock(material.account_code, structure, fiscal_year)
            if voucher_type != Voucher.TYPE_ENTRY and summary and line.quantity > summary.remaining_qty:
                raise ValueError(f"La quantite sortante depasse le stock disponible pour {material.name}.")
            effective_unit_price = line.unit_price
            if voucher_type != Voucher.TYPE_ENTRY:
                effective_unit_price = summary.cmup if summary and summary.cmup > ZERO else ZERO
            amount = line.quantity * effective_unit_price
            VoucherLine.objects.create(
                voucher=voucher,
                material=material,
                material_name=material.name,
                specification=line.specification,
                quantity=line.quantity,
                unit=material.unit,
                unit_price=effective_unit_price,
                amount=amount,
            )
            if voucher_type == Voucher.TYPE_ENTRY and apply_to_physical_inventory:
                _apply_entry_to_physical_inventory(material, line.quantity, effective_unit_price)
        return voucher

    for group_lines in lines_by_store_type.values():
        created_vouchers.append(_create_single_voucher(group_lines))

    return created_vouchers


def compute_stock_summary(structure: Structure | None = None, fiscal_year: FiscalYear | None = None) -> list[ReportRow]:
    materials_qs = Material.objects.all()
    lines_qs = VoucherLine.objects.select_related("voucher", "material")
    opening_voucher_ids: set[int] = set()

    if structure and fiscal_year:
        materials_qs = materials_qs.filter(structure=structure, fiscal_year=fiscal_year)
        lines_qs = lines_qs.filter(voucher__structure=structure, voucher__fiscal_year=fiscal_year)
        opening_voucher_ids = set(
            Voucher.objects.filter(
                structure=structure,
                fiscal_year=fiscal_year,
                voucher_type=Voucher.TYPE_ENTRY,
                source_or_destination="Balance d'entree",
            ).values_list("id", flat=True)
        )
        if not opening_voucher_ids:
            opening_voucher_ids = set(
                Voucher.objects.filter(
                    structure=structure,
                    fiscal_year=fiscal_year,
                    voucher_type=Voucher.TYPE_ENTRY,
                    number=1,
                ).values_list("id", flat=True)
            )

    materials = list(materials_qs.prefetch_related("physical_inventory"))
    lines = lines_qs.order_by("voucher__operation_date", "voucher__number", "id")

    rows: dict[str, ReportRow] = {}
    for material in materials:
        physical = getattr(material, "physical_inventory", None)
        pv_is_filled = bool(physical and physical.is_pv_filled)
        group_code = (material.group_code or "").strip()
        is_groupe_2 = group_code.startswith("2")
        # Par défaut, tout à zéro
        pending_qty = ZERO
        in_service_qty = ZERO
        provisional_qty = ZERO
        if physical:
            if is_groupe_2:
                # Pour les fournitures, pending_qty sera calculé après (= remaining_qty)
                pending_qty = ZERO
                in_service_qty = ZERO
                provisional_qty = ZERO
            else:
                # Groupe 1 : utiliser les valeurs du PV tel quel (pas de forçage depuis les bons)
                pending_qty = physical.pending_qty
                in_service_qty = physical.in_service_qty
                provisional_qty = physical.provisional_qty
        rows[material.account_code] = ReportRow(
            account_code=material.account_code,
            name=material.name,
            unit=material.unit,
            group_code=material.group_code,
            qty_in=ZERO,
            amount_in=ZERO,
            qty_in_opening=ZERO,
            amount_in_opening=ZERO,
            qty_in_period=ZERO,
            amount_in_period=ZERO,
            qty_out=ZERO,
            amount_out=ZERO,
            qty_out_definitive=ZERO,
            qty_out_provisional=ZERO,
            pending_qty=pending_qty,
            in_service_qty=in_service_qty,
            provisional_qty=provisional_qty,
            pv_is_filled=pv_is_filled,
            inventory_total_qty=ZERO,
            inventory_unit_price=ZERO,
            remaining_qty=ZERO,
            remaining_amount=ZERO,
            cmup=ZERO,
            inventory_gap=ZERO,
            gap_plus=ZERO,
            gap_minus=ZERO,
            inventory_gap_value=ZERO,
        )

    for line in lines:
        row = rows[line.material.account_code]
        if line.voucher.voucher_type == Voucher.TYPE_ENTRY:
            row.qty_in += line.quantity
            row.amount_in += line.amount
            if opening_voucher_ids and line.voucher_id in opening_voucher_ids:
                row.qty_in_opening += line.quantity
                row.amount_in_opening += line.amount
            else:
                row.qty_in_period += line.quantity
                row.amount_in_period += line.amount
        elif line.voucher.voucher_type == Voucher.TYPE_TEMP_EXIT:
            # Sortie provisoire: tracee au journal, sans impact sur l'existant comptable.
            row.qty_out_provisional += line.quantity
        else:
            row.qty_out += line.quantity
            row.amount_out += line.amount
            row.qty_out_definitive += line.quantity

    result: list[ReportRow] = []
    for row in rows.values():
        row.remaining_qty = row.qty_in - row.qty_out
        row.remaining_amount = row.amount_in - row.amount_out
        group_code = (row.group_code or "").strip()
        is_groupe_2 = group_code.startswith("2")
        if is_groupe_2:
            # Fournitures : tout le stock restant est en attente d'affectation
            row.pending_qty = row.remaining_qty
            row.in_service_qty = ZERO
            row.provisional_qty = ZERO
        else:
            # Rattrapage historique: si la fiche physique est absente/incomplete,
            # on impute le reliquat au statut "en attente d'affectation".
            classified_qty = row.pending_qty + row.in_service_qty + row.provisional_qty
            if classified_qty < row.remaining_qty:
                row.pending_qty += (row.remaining_qty - classified_qty)
        if row.remaining_qty > ZERO:
            row.cmup = (row.remaining_amount / row.remaining_qty).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        else:
            row.cmup = ZERO
        if row.pv_is_filled:
            row.inventory_total_qty = row.pending_qty + row.in_service_qty + row.provisional_qty
            row.inventory_gap = row.inventory_total_qty - row.remaining_qty
            row.gap_plus = row.inventory_gap if row.inventory_gap > ZERO else ZERO
            row.gap_minus = -row.inventory_gap if row.inventory_gap < ZERO else ZERO
            row.inventory_unit_price = row.cmup
            row.inventory_gap_value = row.inventory_gap * row.cmup
        result.append(row)
    return sorted(result, key=lambda item: item.account_code)


def get_material_stock(
    account_code: str,
    structure: Structure | None = None,
    fiscal_year: FiscalYear | None = None,
) -> ReportRow | None:
    for row in compute_stock_summary(structure=structure, fiscal_year=fiscal_year):
        if row.account_code == account_code:
            return row
    return None


def build_grand_ledger(
    account_code: str,
    structure: Structure | None = None,
    fiscal_year: FiscalYear | None = None,
) -> list[dict]:
    lines_qs = VoucherLine.objects.select_related("voucher", "material").filter(
        material__account_code=account_code
    )
    if structure and fiscal_year:
        lines_qs = lines_qs.filter(voucher__structure=structure, voucher__fiscal_year=fiscal_year)
    
    lines = lines_qs.order_by("voucher__operation_date", "voucher__number", "id")
    running_qty = ZERO
    running_amount = ZERO
    ledger = []
    for line in lines:
        entry_qty = line.quantity if line.voucher.voucher_type == Voucher.TYPE_ENTRY else ZERO
        exit_qty = ZERO if line.voucher.voucher_type == Voucher.TYPE_ENTRY else line.quantity
        impacts_stock = line.voucher.voucher_type != Voucher.TYPE_TEMP_EXIT
        running_qty += entry_qty - (exit_qty if impacts_stock else ZERO)
        running_amount += line.amount if line.voucher.voucher_type == Voucher.TYPE_ENTRY else (-line.amount if impacts_stock else ZERO)
        ledger.append(
            {
                "date": line.voucher.operation_date,
                "voucher_number": line.voucher.number,
                "voucher_type": line.voucher.get_voucher_type_display(),
                "designation": line.voucher.source_or_destination,
                "entry_qty": entry_qty,
                "exit_qty": exit_qty,
                "unit_price": line.unit_price,
                "running_qty": running_qty,
                "running_amount": running_amount,
            }
        )
    return ledger


@transaction.atomic
def import_workbook(
    workbook_path: str | Path,
    *,
    replace_existing: bool = True,
    structure: Structure | None = None,
    fiscal_year: FiscalYear | None = None,
) -> dict[str, int]:
    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Classeur introuvable: {workbook_path}")
    
    if not structure or not fiscal_year:
        raise ValueError("Structure et exercice sont obligatoires pour l'import.")
    if fiscal_year.is_closed:
        raise ValueError(f"L'exercice {fiscal_year.year} est cloture. Import impossible.")

    workbook = load_workbook(workbook_path, data_only=True, keep_vba=True)
    if replace_existing:
        VoucherLine.objects.filter(voucher__structure=structure, voucher__fiscal_year=fiscal_year).delete()
        Voucher.objects.filter(structure=structure, fiscal_year=fiscal_year).delete()
        PhysicalInventory.objects.filter(material__structure=structure, material__fiscal_year=fiscal_year).delete()
        Material.objects.filter(structure=structure, fiscal_year=fiscal_year).delete()

    counters = {"materials": 0, "vouchers": 0, "lines": 0, "inventory_rows": 0}

    if "Menu" in workbook.sheetnames:
        menu = workbook["Menu"]
        for key, value in {
            "structure_label": _clean(menu["B3"].value),
            "structure_value": _clean(menu["D3"].value),
            "district_label": _clean(menu["B4"].value),
            "district_value": _clean(menu["D4"].value),
            "fiscal_year": _clean(menu["F6"].value),
        }.items():
            if value:
                set_setting(key, value)

    seen_accounts: set[str] = set()
    nomenclature = workbook["Nomenclature"]
    for row_index in range(6, nomenclature.max_row + 1):
        for account_col, name_col, unit_col in ((1, 2, 3), (5, 6, 7)):
            account_code = _clean(nomenclature.cell(row_index, account_col).value)
            name = _clean(nomenclature.cell(row_index, name_col).value)
            unit = _clean(nomenclature.cell(row_index, unit_col).value)
            if not account_code or not name or not unit or account_code in seen_accounts:
                continue
            Material.objects.update_or_create(
                structure=structure,
                fiscal_year=fiscal_year,
                account_code=account_code,
                defaults={"name": name, "unit": unit, "group_code": account_code[:2]},
            )
            # Also upsert the global nomenclature reference
            NomenclatureItem.objects.update_or_create(
                account_code=account_code,
                defaults={"name": name, "unit": unit, "group_code": account_code[:2]},
            )
            seen_accounts.add(account_code)
            counters["materials"] += 1

    journal = workbook["LivreJournal"]
    vouchers: dict[tuple[int, str], Voucher] = {}
    for row_index in range(8, journal.max_row + 1):
        operation_date = journal.cell(row_index, 1).value
        account_code = _clean(journal.cell(row_index, 2).value)
        material_name = _clean(journal.cell(row_index, 3).value)
        entry_voucher = journal.cell(row_index, 4).value
        entry_qty = _floatish(journal.cell(row_index, 5).value)
        exit_voucher = journal.cell(row_index, 7).value
        exit_qty = _floatish(journal.cell(row_index, 8).value)
        unit = _clean(journal.cell(row_index, 6).value or journal.cell(row_index, 9).value) or "N"
        unit_price = _floatish(journal.cell(row_index, 10).value)
        origin = _clean(journal.cell(row_index, 14).value)
        provisional_marker = journal.cell(row_index, 13).value

        if not account_code or not material_name:
            continue

        material, created = Material.objects.get_or_create(
            structure=structure,
            fiscal_year=fiscal_year,
            account_code=account_code,
            defaults={"name": material_name, "unit": unit, "group_code": account_code[:2]},
        )
        if created:
            counters["materials"] += 1

        def get_or_create_voucher(number_value, voucher_type: str, voucher_material: Material) -> Voucher:
            number = int(number_value) if number_value not in (None, "") else next_voucher_number(structure, fiscal_year)
            key = (number, voucher_type)
            expected_store = _resolve_voucher_store_for_materials(materials=[voucher_material], structure=structure)
            if key not in vouchers:
                vouchers[key] = Voucher.objects.create(
                    structure=structure,
                    fiscal_year=fiscal_year,
                    number=number,
                    voucher_type=voucher_type,
                    operation_date=operation_date.date() if isinstance(operation_date, datetime) else operation_date,
                    source_or_destination=origin,
                    storage_location=expected_store,
                    comments="",
                )
                counters["vouchers"] += 1
            elif vouchers[key].storage_location_id is None:
                vouchers[key].storage_location = expected_store
                vouchers[key].save(update_fields=["storage_location", "updated_at"])
            return vouchers[key]

        if entry_qty > ZERO:
            voucher = get_or_create_voucher(entry_voucher, Voucher.TYPE_ENTRY, material)
            VoucherLine.objects.create(
                voucher=voucher,
                material=material,
                material_name=material.name,
                specification="",
                quantity=entry_qty,
                unit=material.unit,
                unit_price=unit_price,
                amount=entry_qty * unit_price,
            )
            counters["lines"] += 1

        if exit_qty > ZERO:
            voucher_type = Voucher.TYPE_TEMP_EXIT if provisional_marker not in (None, "") else Voucher.TYPE_FINAL_EXIT
            voucher = get_or_create_voucher(exit_voucher, voucher_type, material)
            VoucherLine.objects.create(
                voucher=voucher,
                material=material,
                material_name=material.name,
                specification="",
                quantity=exit_qty,
                unit=material.unit,
                unit_price=unit_price,
                amount=exit_qty * unit_price,
            )
            counters["lines"] += 1

    if "Releve recapitulatif" in workbook.sheetnames:
        recap = workbook["Releve recapitulatif"]
        for row_index in range(7, recap.max_row + 1):
            account_code = _clean(recap.cell(row_index, 1).value)
            if not account_code:
                continue
            material = Material.objects.filter(
                structure=structure,
                fiscal_year=fiscal_year,
                account_code=account_code
            ).first()
            if not material:
                continue
            pending_qty = _floatish(recap.cell(row_index, 11).value)
            in_service_qty = _floatish(recap.cell(row_index, 12).value)
            provisional_qty = _floatish(recap.cell(row_index, 13).value)
            if pending_qty or in_service_qty or provisional_qty:
                PhysicalInventory.objects.update_or_create(
                    material=material,
                    defaults={
                        "pending_qty": pending_qty,
                        "in_service_qty": in_service_qty,
                        "provisional_qty": provisional_qty,
                    },
                )
                counters["inventory_rows"] += 1

    set_setting("source_workbook", str(workbook_path))
    return counters


def default_site_context(*, structure: Structure | None = None, fiscal_year: FiscalYear | None = None) -> dict[str, str]:
    return {
        "country": get_setting("country", "REPUBLIQUE DU SENEGAL", structure=structure),
        "ministry": get_setting("ministry", "MINISTERE DE LA SANTE ET DE L'ACTION SOCIALE", structure=structure),
        "structure_label": get_setting("structure_label", "REGION MEDICALE DE", structure=structure),
        "structure_value": get_setting("structure_value", structure.code if structure else "", structure=structure),
        "address": get_setting("address", "", structure=structure),
        "phone": get_setting("phone", "", structure=structure),
        "email": get_setting("email", "", structure=structure),
        "fiscal_year": get_setting("fiscal_year", str(fiscal_year.year) if fiscal_year else "", structure=structure),
    }


@transaction.atomic
def carry_forward_to_new_year(
    *,
    structure: Structure,
    current_fy: FiscalYear,
    user,
) -> tuple[FiscalYear, list[Voucher]]:
    """
    Cloture current_fy, cree l'exercice N+1, copie la nomenclature
    et genere le Bon N°1 "Balance d'entree" avec les qtés finales.
    """
    # Lock the current fiscal year row to avoid two concurrent closure/report operations.
    current_fy = FiscalYear.objects.select_for_update().get(pk=current_fy.pk)

    if current_fy.is_closed:
        raise ValueError(f"L'exercice {current_fy.year} est deja cloture.")

    if Voucher.objects.filter(structure=structure, fiscal_year=current_fy).exists() is False:
        pass  # pas de bons, on peut quand même reporter (stock vide)

    summary = compute_stock_summary(structure=structure, fiscal_year=current_fy)

    new_year = current_fy.year + 1
    new_start = current_fy.end_date + timedelta(days=1)
    new_end = new_start.replace(year=new_start.year + 1) - timedelta(days=1)

    new_fy, created = FiscalYear.objects.get_or_create(
        structure=structure,
        year=new_year,
        defaults={
            "start_date": new_start,
            "end_date": new_end,
            "is_active": True,
            "is_closed": False,
        },
    )

    if not created and new_fy.is_closed:
        raise ValueError(f"L'exercice {new_year} existe deja et est deja cloture.")

    if not created and Voucher.objects.filter(structure=structure, fiscal_year=new_fy).exists():
        raise ValueError(
            f"L'exercice {new_year} a deja des bons enregistres. Report impossible."
        )

    if not created and InternalMovement.objects.filter(structure=structure, fiscal_year=new_fy).exists():
        raise ValueError(
            f"L'exercice {new_year} a deja des bordereaux internes enregistres. Report impossible."
        )

    # Copier la nomenclature vers le nouvel exercice
    material_map: dict[str, Material] = {}
    for row in summary:
        old_material = Material.objects.filter(
            structure=structure,
            fiscal_year=current_fy,
            account_code=row.account_code,
        ).first()
        if old_material is None:
            continue
        new_material, _ = Material.objects.get_or_create(
            structure=structure,
            fiscal_year=new_fy,
            account_code=row.account_code,
            defaults={
                "name": row.name,
                "unit": row.unit,
                "group_code": old_material.group_code,
            },
        )
        material_map[row.account_code] = new_material

    # Ensure all global nomenclature items also exist in the new fiscal year
    for nom_item in NomenclatureItem.objects.all():
        new_material, _ = Material.objects.get_or_create(
            structure=structure,
            fiscal_year=new_fy,
            account_code=nom_item.account_code,
            defaults={"name": nom_item.name, "unit": nom_item.unit, "group_code": nom_item.group_code},
        )
        material_map.setdefault(nom_item.account_code, new_material)

    # Copy physical inventory records from the current fiscal year to the new fiscal year.
    # We preserve the existence of individual inventory sheets, but values must start empty
    # so the new fiscal year PV is filled manually by the user.
    for old_material in Material.objects.filter(structure=structure, fiscal_year=current_fy):
        old_inventory = getattr(old_material, "physical_inventory", None)
        if old_inventory is None:
            continue
        new_material = material_map.get(old_material.account_code)
        if new_material is None:
            continue
        PhysicalInventory.objects.update_or_create(
            material=new_material,
            defaults={
                "pending_qty": ZERO,
                "in_service_qty": ZERO,
                "provisional_qty": ZERO,
                "unit_price": ZERO,
                "is_pv_filled": False,
            },
        )

    # Construire les lignes de la balance d'entree
    opening_lines: list[VoucherLineInput] = []
    for row in summary:
        reg_qty = row.inventory_gap if row.pv_is_filled else ZERO
        final_qty = row.remaining_qty + reg_qty
        if final_qty <= ZERO:
            continue
        new_material = material_map.get(row.account_code)
        if new_material is None:
            continue
        opening_lines.append(
            VoucherLineInput(
                material_id=new_material.id,
                specification="Report de balance",
                quantity=final_qty,
                unit_price=row.cmup,
            )
        )

    # Cloturer l'exercice courant
    current_fy.is_closed = True
    current_fy.is_active = False
    current_fy.save(update_fields=["is_closed", "is_active"])

    # Creer le Bon N°1 (balance d'entree)
    opening_vouchers: list[Voucher] = []
    if opening_lines:
        opening_vouchers = create_voucher(
            user=user,
            voucher_type=Voucher.TYPE_ENTRY,
            operation_date=new_fy.start_date,
            source_or_destination="Balance d'entree",
            comments=f"Report de la balance generale de l'exercice {current_fy.year}",
            lines=opening_lines,
            structure=structure,
            fiscal_year=new_fy,
            apply_to_physical_inventory=False,
        )

    # Reporter les situations de depart par localisation (inventaire individuel).
    for location in Location.objects.filter(structure=structure).order_by("name"):
        if location.is_store:
            continue
        rows = build_location_inventory_report(location, current_fy)
        if not rows:
            continue

        lines_by_store_type: dict[str, list[tuple[Material, LocationInventoryRow]]] = {}
        for row in rows:
            if row.quantity <= ZERO:
                continue
            old_material = Material.objects.filter(pk=row.material_id).first()
            if old_material is None:
                continue
            new_material = material_map.get(old_material.account_code)
            if new_material is None:
                continue
            store_type = _expected_store_location_type(old_material)
            lines_by_store_type.setdefault(store_type, []).append((new_material, row))

        for group in lines_by_store_type.values():
            materials = [item[0] for item in group]
            expected_store = _resolve_voucher_store_for_materials(materials=materials, structure=structure)
            movement = _create_internal_movement_header_with_retry(
                structure=structure,
                fiscal_year=new_fy,
                movement_type=InternalMovement.TYPE_ASSIGNMENT,
                operation_date=new_fy.start_date,
                from_location=expected_store,
                to_location=location,
                reason=f"Report des affectations depuis {current_fy.year}",
                created_by=user if getattr(user, "is_authenticated", False) else None,
            )
            for new_material, row in group:
                amount = row.quantity * row.unit_price
                InternalMovementLine.objects.create(
                    movement=movement,
                    material=new_material,
                    material_name=new_material.name,
                    inventory_code=row.inventory_code,
                    quantity=row.quantity,
                    unit=new_material.unit,
                    unit_price=row.unit_price,
                    amount=amount,
                    observations="Report affectation exercice precedent",
                )

    return new_fy, opening_vouchers
