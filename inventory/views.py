from __future__ import annotations

import csv
from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path
from uuid import uuid4
import subprocess
from urllib.parse import urlencode

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.contrib.auth.views import LoginView
from django.core.cache import cache
from django.core.files.base import ContentFile
from django.core.paginator import Paginator
from django.db import IntegrityError
from django.db.models import Count, Prefetch, Q
from django.forms import ValidationError
from django.http import FileResponse, HttpRequest, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse, reverse_lazy
from django.utils import timezone

try:
    from num2words import num2words as _num2words
except Exception:  # noqa: BLE001
    _num2words = None

from .forms import (
    InternalMovementHeaderForm,
    LocationForm,
    MaterialForm,
    NomenclatureItemForm,
    PhysicalInventoryForm,
    ProfileForm,
    UserUpdateWithProfileForm,
    QuarterlyReportForm,
    parse_internal_movement_lines,
    VoucherLineInput,
    UserCreationWithProfileForm,
    VoucherHeaderForm,
    parse_voucher_lines,
)
from .models import AuditLog, FiscalYear, InternalMovement, Location, Material, NomenclatureItem, PhysicalInventory, QuarterlyReport, SiteSetting, Structure, StructureType, UserProfile, Voucher
from .services import build_grand_ledger, build_location_inventory_report, carry_forward_to_new_year, compute_stock_summary, create_internal_movement, create_voucher, dashboard_metrics, default_site_context, delete_internal_movement, get_setting, list_internal_movements, next_internal_movement_number, propagate_nomenclature_add, propagate_nomenclature_delete, set_setting, sync_all_nomenclature_to_structure


User = get_user_model()


def _get_or_create_profile(user) -> UserProfile | None:
    if not getattr(user, "is_authenticated", False):
        return None
    profile = getattr(user, "profile", None)
    if profile is not None:
        return profile
    profile, _ = UserProfile.objects.get_or_create(user=user)
    return profile


def _query_without_page(request: HttpRequest) -> str:
    params = request.GET.copy()
    params.pop("page", None)
    return params.urlencode()


def _number_to_french_words(value: Decimal) -> str:
    """Convertit un nombre en lettres francaises avec fallback robuste."""
    def _int_to_french_words(n: int) -> str:
        if n == 0:
            return "zero"

        units = [
            "zero", "un", "deux", "trois", "quatre", "cinq", "six", "sept", "huit", "neuf",
            "dix", "onze", "douze", "treize", "quatorze", "quinze", "seize",
        ]
        tens_words = {
            20: "vingt",
            30: "trente",
            40: "quarante",
            50: "cinquante",
            60: "soixante",
        }

        def under_hundred(x: int) -> str:
            if x < 17:
                return units[x]
            if x < 20:
                return "dix-" + units[x - 10]
            if x < 70:
                ten = (x // 10) * 10
                unit = x % 10
                ten_word = tens_words[ten]
                if unit == 0:
                    return ten_word
                if unit == 1:
                    return ten_word + " et un"
                return ten_word + "-" + units[unit]
            if x < 80:
                rest = x - 60
                if rest == 11:
                    return "soixante et onze"
                return "soixante-" + under_hundred(rest)
            if x == 80:
                return "quatre-vingts"
            return "quatre-vingt-" + under_hundred(x - 80)

        def under_thousand(x: int) -> str:
            if x < 100:
                return under_hundred(x)
            hundreds = x // 100
            rest = x % 100
            if hundreds == 1:
                prefix = "cent"
            else:
                prefix = units[hundreds] + " cent"
            if rest == 0 and hundreds > 1:
                return prefix + "s"
            if rest == 0:
                return prefix
            return prefix + " " + under_hundred(rest)

        parts: list[str] = []
        billions = n // 1_000_000_000
        n %= 1_000_000_000
        millions = n // 1_000_000
        n %= 1_000_000
        thousands = n // 1000
        rest = n % 1000

        if billions:
            if billions == 1:
                parts.append("un milliard")
            else:
                parts.append(f"{under_thousand(billions)} milliards")
        if millions:
            if millions == 1:
                parts.append("un million")
            else:
                parts.append(f"{under_thousand(millions)} millions")
        if thousands:
            if thousands == 1:
                parts.append("mille")
            else:
                parts.append(f"{under_thousand(thousands)} mille")
        if rest:
            parts.append(under_thousand(rest))

        return " ".join(parts)

    try:
        normalized = Decimal(value)
    except Exception:  # noqa: BLE001
        return str(value)

    if _num2words is None:
        sign = "moins " if normalized < 0 else ""
        absolute = abs(normalized)
        integer_part = int(absolute)
        decimal_part = int((absolute - integer_part) * 100)
        words = _int_to_french_words(integer_part)
        if decimal_part:
            words = f"{words} virgule {_int_to_french_words(decimal_part)}"
        return f"{sign}{words}"

    try:
        return _num2words(float(normalized), lang="fr")
    except Exception:  # noqa: BLE001
        sign = "moins " if normalized < 0 else ""
        absolute = abs(normalized)
        integer_part = int(absolute)
        decimal_part = int((absolute - integer_part) * 100)
        words = _int_to_french_words(integer_part)
        if decimal_part:
            words = f"{words} virgule {_int_to_french_words(decimal_part)}"
        return f"{sign}{words}"


def _fmt_pdf_int(value: Decimal | int | str) -> str:
    try:
        n = int(Decimal(str(value)))
    except Exception:  # noqa: BLE001
        return str(value)
    return f"{n:,}".replace(",", " ")


def _fmt_pdf_amount(value: Decimal | int | str) -> str:
    try:
        n = Decimal(str(value))
    except Exception:  # noqa: BLE001
        return str(value)
    formatted = f"{n:,.2f}"
    return formatted.replace(",", " ").replace(".", ",")

class ServiceAwareLoginView(LoginView):
    template_name = "registration/login.html"

    def form_valid(self, form):
        response = super().form_valid(form)
        user = self.request.user
        profile = _get_or_create_profile(user)

        if profile:
            selected_structure = None
            available_structures = profile.accessible_structures_qs().order_by("code")
            structure_id = self.request.session.get("structure_id")
            if structure_id:
                selected_structure = available_structures.filter(id=structure_id).first()
            if not selected_structure and profile.default_structure and profile.default_structure.is_active:
                if available_structures.filter(id=profile.default_structure.id).exists():
                    selected_structure = profile.default_structure
            if not selected_structure:
                selected_structure = available_structures.first()

            if selected_structure is None:
                messages.error(self.request, "Aucun service actif n'est assigne a cet utilisateur.")
                return response

            if selected_structure:
                self.request.session["structure_id"] = selected_structure.id
                fy = FiscalYear.objects.filter(structure=selected_structure, year=2025).first()
                if not fy:
                    fy = FiscalYear.objects.filter(
                        structure=selected_structure,
                        is_active=True,
                        is_closed=False,
                    ).order_by("-year").first()
                if not fy:
                    fy = FiscalYear.objects.filter(structure=selected_structure).order_by("-year").first()
                if fy:
                    self.request.session["fiscal_year_id"] = fy.id

        return response

    def get_success_url(self):
        profile = _get_or_create_profile(self.request.user)
        if profile and profile.role != "admin":
            return str(reverse_lazy("screen_operations"))
        return str(reverse_lazy("menu"))


def _can_edit(user) -> bool:
    if not user.is_authenticated:
        return False
    profile = getattr(user, "profile", None)
    return user.is_superuser or user.is_staff or (profile and profile.role in {"admin", "comptable"})


def _is_admin_user(user) -> bool:
    if not user.is_authenticated:
        return False
    profile = getattr(user, "profile", None)
    return user.is_superuser or bool(profile and profile.role == "admin")


def _can_close_fiscal_year(user) -> bool:
    if not user.is_authenticated:
        return False
    profile = getattr(user, "profile", None)
    return user.is_superuser or bool(profile and profile.role in {"admin", "comptable"})


def _ensure_open_fiscal_year(request: HttpRequest, *, redirect_to: str = "screen_operations") -> bool:
    """Retourne False et affiche un message si l'exercice courant est cloture."""
    fy = getattr(request, "current_fiscal_year", None)
    if fy and fy.is_closed:
        messages.error(request, f"L'exercice {fy.year} est cloture. Aucune modification n'est autorisee.")
        return False
    return True


def _audit(
    request: HttpRequest,
    *,
    action: str,
    entity: str,
    entity_id: str = "",
    description: str = "",
    metadata: dict | None = None,
) -> None:
    AuditLog.objects.create(
        user=request.user if request.user.is_authenticated else None,
        structure=getattr(request, "current_structure", None),
        fiscal_year=getattr(request, "current_fiscal_year", None),
        action=action,
        entity=entity,
        entity_id=entity_id,
        description=description,
        metadata=metadata or {},
    )


@login_required
def menu_view(request: HttpRequest) -> HttpResponse:
    profile = _get_or_create_profile(request.user)
    context = {
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
        "is_admin": bool(profile and profile.role == "admin"),
    }
    return render(request, "inventory/menu.html", context)


@login_required
def screen_admin_view(request: HttpRequest) -> HttpResponse:
    profile = _get_or_create_profile(request.user)
    recent_security_rejections_count = 0
    if profile and profile.role == "admin":
        cutoff = timezone.now() - timezone.timedelta(hours=24)
        recent_security_rejections_count = AuditLog.objects.filter(
            entity="DocumentSecurity",
            created_at__gte=cutoff,
        ).count()

    context = {
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
        "is_admin": bool(profile and profile.role == "admin"),
        "recent_security_rejections_count": recent_security_rejections_count,
    }
    return render(request, "inventory/screen_admin.html", context)


@login_required
def screen_operations_view(request: HttpRequest) -> HttpResponse:
    context = {
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
    }
    return render(request, "inventory/screen_operations.html", context)


@login_required
def screen_internal_movements_view(request: HttpRequest) -> HttpResponse:
    context = {
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
        "locations_count": Location.objects.filter(structure=request.current_structure, is_active=True).count() if request.current_structure else 0,
        "movements_count": InternalMovement.objects.filter(
            structure=request.current_structure,
            fiscal_year=request.current_fiscal_year,
        ).count() if request.current_structure and request.current_fiscal_year else 0,
    }
    return render(request, "inventory/screen_internal_movements.html", context)


@login_required
def location_list_view(request: HttpRequest) -> HttpResponse:
    if request.current_structure is None:
        messages.error(request, "Aucun service n'est selectionne.")
        return redirect("menu")

    edit_location = None
    if request.method == "POST":
        if not _can_edit(request.user):
            messages.error(request, "Vous n'avez pas les droits de saisie.")
            return redirect("location_list")
        location_id = request.POST.get("location_id")
        if location_id:
            edit_location = get_object_or_404(Location, pk=location_id, structure=request.current_structure)
        form = LocationForm(request.POST, instance=edit_location)
        if form.is_valid():
            location = form.save(commit=False)
            location.structure = request.current_structure
            is_create = location.pk is None
            location.save()
            _audit(
                request,
                action=AuditLog.ACTION_CREATE if is_create else AuditLog.ACTION_UPDATE,
                entity="Location",
                entity_id=str(location.id),
                description=f"{'Creation' if is_create else 'Mise a jour'} localisation {location.name}",
                metadata={"type": location.location_type},
            )
            messages.success(request, "Localisation enregistree.")
            return redirect("location_list")
    else:
        edit_id = request.GET.get("edit")
        if edit_id:
            edit_location = get_object_or_404(Location, pk=edit_id, structure=request.current_structure)
        form = LocationForm(instance=edit_location)

    locations = Location.objects.filter(structure=request.current_structure).order_by("location_type", "name")
    context = {
        "form": form,
        "locations": locations,
        "edit_location": edit_location,
        "can_edit": _can_edit(request.user),
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
    }
    return render(request, "inventory/location_list.html", context)


@login_required
def location_delete_view(request: HttpRequest, pk: int) -> HttpResponse:
    if request.current_structure is None:
        messages.error(request, "Aucun service n'est selectionne.")
        return redirect("menu")
    if request.method != "POST":
        return redirect("location_list")
    if not _can_edit(request.user):
        messages.error(request, "Vous n'avez pas les droits de suppression.")
        return redirect("location_list")

    location = get_object_or_404(Location, pk=pk, structure=request.current_structure)

    has_internal_movements = InternalMovement.objects.filter(
        Q(from_location=location) | Q(to_location=location)
    ).exists()
    has_vouchers = Voucher.objects.filter(storage_location=location).exists()

    if has_internal_movements or has_vouchers:
        messages.error(
            request,
            "Suppression impossible: cette localisation contient des donnees (bordereaux ou bons).",
        )
        return redirect("location_list")

    location_name = location.name
    location_id = location.id
    location.delete()
    _audit(
        request,
        action=AuditLog.ACTION_DELETE,
        entity="Location",
        entity_id=str(location_id),
        description=f"Suppression localisation {location_name}",
    )
    messages.success(request, "Localisation supprimee.")
    return redirect("location_list")


@login_required
def internal_movement_list_view(request: HttpRequest) -> HttpResponse:
    q = (request.GET.get("q") or "").strip()
    movement_type = (request.GET.get("movement_type") or "").strip()
    from_location_id = (request.GET.get("from_location") or "").strip()
    to_location_id = (request.GET.get("to_location") or "").strip()
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    sort = request.GET.get("sort") or "date_desc"

    movements_page = []
    locations = Location.objects.none()
    if request.current_structure and request.current_fiscal_year:
        locations = Location.objects.filter(structure=request.current_structure).order_by("name")
        movements_qs = InternalMovement.objects.filter(
            structure=request.current_structure,
            fiscal_year=request.current_fiscal_year,
        ).select_related("from_location", "to_location", "created_by").prefetch_related("lines__material")

        if movement_type in dict(InternalMovement.TYPE_CHOICES):
            movements_qs = movements_qs.filter(movement_type=movement_type)
        if from_location_id.isdigit():
            movements_qs = movements_qs.filter(from_location_id=int(from_location_id))
        if to_location_id.isdigit():
            movements_qs = movements_qs.filter(to_location_id=int(to_location_id))
        if date_from:
            movements_qs = movements_qs.filter(operation_date__gte=date_from)
        if date_to:
            movements_qs = movements_qs.filter(operation_date__lte=date_to)
        if q:
            query_filter = Q(reason__icontains=q) | Q(from_location__name__icontains=q) | Q(to_location__name__icontains=q)
            if q.isdigit():
                query_filter |= Q(number=int(q))
            movements_qs = movements_qs.filter(query_filter)

        sort_mapping = {
            "date_desc": ["-operation_date", "-number"],
            "date_asc": ["operation_date", "number"],
            "number_desc": ["-number"],
            "number_asc": ["number"],
        }
        movements_qs = movements_qs.order_by(*sort_mapping.get(sort, ["-operation_date", "-number"]))
        paginator = Paginator(movements_qs, 20)
        movements_page = paginator.get_page(request.GET.get("page"))

    context = {
        "movements": movements_page,
        "locations": locations,
        "movement_type_choices": InternalMovement.TYPE_CHOICES,
        "filters": {
            "q": q,
            "movement_type": movement_type,
            "from_location": from_location_id,
            "to_location": to_location_id,
            "date_from": date_from,
            "date_to": date_to,
            "sort": sort,
        },
        "query_without_page": _query_without_page(request),
        "can_edit": _can_edit(request.user),
        "is_fiscal_year_closed": bool(request.current_fiscal_year and request.current_fiscal_year.is_closed),
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
    }
    return render(request, "inventory/internal_movement_list.html", context)


@login_required
def internal_movement_create_view(request: HttpRequest) -> HttpResponse:
    if request.current_structure is None or request.current_fiscal_year is None:
        messages.error(request, "Selectionnez d'abord un service et un exercice.")
        return redirect("menu")

    active_locations = Location.objects.filter(structure=request.current_structure, is_active=True)
    if not active_locations.exists():
        messages.error(request, "Creez d'abord au moins une localisation avant de saisir un bordereau interne.")
        return redirect("location_list")

    materials = Material.objects.filter(
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    ).order_by("account_code")
    summary_by_code = {
        row.account_code: row
        for row in compute_stock_summary(
            structure=request.current_structure,
            fiscal_year=request.current_fiscal_year,
        )
    }
    materials_with_cmup = [
        {
            "id": material.id,
            "account_code": material.account_code,
            "name": material.name,
            "unit": material.unit,
            "cmup": summary_by_code.get(material.account_code).cmup if summary_by_code.get(material.account_code) else Decimal("0"),
            "pending_qty": summary_by_code.get(material.account_code).pending_qty if summary_by_code.get(material.account_code) else Decimal("0"),
            "in_service_qty": summary_by_code.get(material.account_code).in_service_qty if summary_by_code.get(material.account_code) else Decimal("0"),
        }
        for material in materials
    ]
    locations_data = [
        {
            "id": str(loc.id),
            "name": loc.name,
            "is_store": loc.is_store,
            "type_label": loc.get_location_type_display(),
        }
        for loc in active_locations.order_by("name")
    ]
    next_number = next_internal_movement_number(request.current_structure, request.current_fiscal_year)

    line_rows = []
    initial_line_count = 4
    if request.method == "POST":
        if not _can_edit(request.user):
            messages.error(request, "Vous n'avez pas les droits de saisie.")
            return redirect("internal_movement_create")
        if not _ensure_open_fiscal_year(request, redirect_to="internal_movement_create"):
            return redirect("internal_movement_create")
        form = InternalMovementHeaderForm(request.POST, structure=request.current_structure)
        try:
            line_count = int(request.POST.get("line_count") or initial_line_count)
        except (TypeError, ValueError):
            line_count = initial_line_count
        for index in range(line_count):
            line_rows.append(
                {
                    "index": index,
                    "material_id": request.POST.get(f"line_material_{index}", ""),
                    "inventory_code": request.POST.get(f"line_inventory_code_{index}", ""),
                    "quantity": request.POST.get(f"line_quantity_{index}", ""),
                    "unit_price": request.POST.get(f"line_unit_price_{index}", ""),
                    "observations": request.POST.get(f"line_observations_{index}", ""),
                }
            )
        if form.is_valid():
            try:
                lines = parse_internal_movement_lines(request.POST)
                movement = create_internal_movement(
                    user=request.user,
                    movement_type=form.cleaned_data["movement_type"],
                    operation_date=form.cleaned_data["operation_date"],
                    from_location=form.cleaned_data["from_location"],
                    to_location=form.cleaned_data["to_location"],
                    reason=form.cleaned_data["reason"],
                    lines=lines,
                    structure=request.current_structure,
                    fiscal_year=request.current_fiscal_year,
                )
            except (ValueError, ValidationError) as error:
                messages.error(request, str(error))
            else:
                _audit(
                    request,
                    action=AuditLog.ACTION_CREATE,
                    entity="InternalMovement",
                    entity_id=str(movement.id),
                    description=f"Creation bordereau interne {movement.number}",
                    metadata={"type": movement.movement_type, "lines": movement.lines.count()},
                )
                _invalidate_central_recap_cache()
                messages.success(request, f"Bordereau interne {movement.number} enregistre.")
                return redirect("internal_movement_detail", pk=movement.pk)
    else:
        form = InternalMovementHeaderForm(
            initial={"operation_date": date.today()},
            structure=request.current_structure,
        )
        for index in range(initial_line_count):
            line_rows.append(
                {
                    "index": index,
                    "material_id": "",
                    "inventory_code": "",
                    "quantity": "",
                    "unit_price": "",
                    "observations": "",
                }
            )

    context = {
        "form": form,
        "materials_with_cmup": materials_with_cmup,
        "locations_data": locations_data,
        "next_number": next_number,
        "line_rows": line_rows,
        "can_edit": _can_edit(request.user),
        "is_fiscal_year_closed": bool(request.current_fiscal_year and request.current_fiscal_year.is_closed),
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
    }
    return render(request, "inventory/internal_movement_form.html", context)


@login_required
def internal_movement_detail_view(request: HttpRequest, pk: int) -> HttpResponse:
    movement = get_object_or_404(
        InternalMovement.objects.select_related("from_location", "to_location", "created_by").prefetch_related("lines__material"),
        pk=pk,
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    )
    total_quantity = sum((line.quantity for line in movement.lines.all()), Decimal("0"))
    context = {
        "movement": movement,
        "total_quantity": total_quantity,
        "can_edit": _can_edit(request.user),
        "is_fiscal_year_closed": bool(request.current_fiscal_year and request.current_fiscal_year.is_closed),
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
    }
    return render(request, "inventory/internal_movement_detail.html", context)


@login_required
def internal_movement_pdf_view(request: HttpRequest, pk: int) -> HttpResponse:
    movement = get_object_or_404(
        InternalMovement.objects.select_related("from_location", "to_location", "created_by").prefetch_related("lines__material"),
        pk=pk,
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    )

    headers = [
        "Compte",
        "Designation",
        "Code immatriculation",
        "Quantite",
        "PU",
        "Montant",
    ]
    rows = [
        [
            line.material.account_code,
            line.material_name,
            line.inventory_code or "-",
            _fmt_pdf_int(line.quantity),
            _fmt_pdf_amount(line.unit_price),
            _fmt_pdf_amount(line.amount),
        ]
        for line in movement.lines.all()
    ]

    total_qty = sum((line.quantity for line in movement.lines.all()), Decimal("0"))
    total_amount = sum((line.amount for line in movement.lines.all()), Decimal("0"))
    receipt_verb = "remis" if movement.movement_type == InternalMovement.TYPE_UNASSIGNMENT else "recu"
    movement_doc_label = {
        InternalMovement.TYPE_ASSIGNMENT: "d'affectation",
        InternalMovement.TYPE_TRANSFER: "de mutation",
        InternalMovement.TYPE_UNASSIGNMENT: "de desaffectation",
    }.get(movement.movement_type, "")
    declaration = (
        "Je, soussigne, reconnais avoir "
        f"{receipt_verb} les matieres portees au present bordereau {movement_doc_label} de "
        f"{_fmt_pdf_int(total_qty)} unites representant une valeur de {_fmt_pdf_amount(total_amount)} F CFA."
    )
    totals_line = (
        f"Total quantites: {_fmt_pdf_int(total_qty)} unites. "
        f"Total montants: {_fmt_pdf_amount(total_amount)} F CFA."
    )
    
    # Model numbers: 7 for affectation/desaffectation, 8 for transfer (mutation)
    model_ref = "Modele 8" if movement.movement_type == InternalMovement.TYPE_TRANSFER else "Modele 7"

    return _build_pdf_response(
        title=f"Bordereau interne N° {movement.number}",
        headers=headers,
        rows=rows,
        filename=f"bordereau_interne_{movement.number}.pdf",
        signatures=[
            "Ordonnateur des matieres",
            "Comptable des matieres",
            "Detenteur",
        ],
        context_lines=_pdf_context_lines(request),
        document_ref=model_ref,
        subtitle=(
            f"{movement.get_movement_type_display()} du {movement.operation_date.strftime('%d/%m/%Y')} "
            f"de {movement.from_location.name} vers {movement.to_location.name}"
        ),
        logo_path=_pdf_logo_path(request),
        extra_paragraphs=[totals_line, declaration],
    )


@login_required
def internal_movement_delete_view(request: HttpRequest, pk: int) -> HttpResponse:
    movement = get_object_or_404(
        InternalMovement,
        pk=pk,
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    )
    if request.method != "POST":
        return redirect("internal_movement_detail", pk=movement.pk)
    if not _can_edit(request.user):
        messages.error(request, "Vous n'avez pas les droits de suppression.")
        return redirect("internal_movement_detail", pk=movement.pk)
    if not _ensure_open_fiscal_year(request, redirect_to="internal_movement_detail"):
        return redirect("internal_movement_detail", pk=movement.pk)

    movement_number = movement.number
    delete_internal_movement(movement)
    _audit(
        request,
        action=AuditLog.ACTION_DELETE,
        entity="InternalMovement",
        entity_id=str(pk),
        description=f"Suppression bordereau interne {movement_number}",
    )
    _invalidate_central_recap_cache()
    messages.success(request, f"Bordereau interne {movement_number} supprime.")
    return redirect("internal_movement_list")


@login_required
def contradictory_inventory_view(request: HttpRequest) -> HttpResponse:
    if request.current_structure is None or request.current_fiscal_year is None:
        messages.error(request, "Selectionnez d'abord un service et un exercice.")
        return redirect("menu")

    # Inclure toutes les localisations actives, y compris les magasins.
    locations = Location.objects.filter(
        structure=request.current_structure,
        is_active=True,
    ).order_by("name")

    selected_location = None
    if locations.exists():
        location_id = request.GET.get("location_id")
        if location_id:
            selected_location = get_object_or_404(locations, pk=location_id)
        else:
            selected_location = locations.first()

    rows = build_location_inventory_report(selected_location, request.current_fiscal_year) if selected_location else []
    total_quantity = sum((row.quantity for row in rows), Decimal("0"))
    total_amount = sum((row.amount for row in rows), Decimal("0"))
    context = {
        "locations": locations,
        "selected_location": selected_location,
        "rows": rows,
        "total_quantity": total_quantity,
        "total_amount": total_amount,
        "generated_on": timezone.now(),
        "comptable_name": getattr(getattr(request.user, "profile", None), "display_name", "") or request.user.get_full_name() or request.user.username,
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
    }
    return render(request, "inventory/contradictory_inventory.html", context)


@login_required
def contradictory_inventory_pdf_view(request: HttpRequest) -> HttpResponse:
    if request.current_structure is None or request.current_fiscal_year is None:
        messages.error(request, "Selectionnez d'abord un service et un exercice.")
        return redirect("menu")

    locations = Location.objects.filter(
        structure=request.current_structure,
        is_active=True,
    ).order_by("name")

    selected_location = None
    if locations.exists():
        location_id = request.GET.get("location_id")
        if location_id:
            selected_location = get_object_or_404(locations, pk=location_id)
        else:
            selected_location = locations.first()

    if selected_location is None:
        messages.error(request, "Aucune localisation disponible pour l'impression.")
        return redirect("contradictory_inventory")

    rows_data = build_location_inventory_report(selected_location, request.current_fiscal_year)
    headers = [
        "Compte",
        "Designation",
        "Code immatriculation",
        "Quantite",
        "Unite",
        "PU",
        "Montant",
        "Dernier mouvement",
    ]
    rows = [
        [
            row.account_code,
            row.material_name,
            row.inventory_code or "-",
            _fmt_pdf_int(row.quantity),
            row.unit,
            _fmt_pdf_amount(row.unit_price),
            _fmt_pdf_amount(row.amount),
            row.last_movement_date.strftime("%d/%m/%Y") if row.last_movement_date else "-",
        ]
        for row in rows_data
    ]
    total_qty = sum((row.quantity for row in rows_data), Decimal("0"))
    total_amount = sum((row.amount for row in rows_data), Decimal("0"))
    totals_line = (
        f"Total quantites: {_fmt_pdf_int(total_qty)} unites. "
        f"Total montants: {_fmt_pdf_amount(total_amount)} F CFA."
    )

    return _build_pdf_response(
        title="Fiche d'inventaire individuel contradictoire",
        headers=headers,
        rows=rows,
        filename=f"fiche_inventaire_{selected_location.id}.pdf",
        signatures=[
            "Ordonnateur des matieres",
            "Comptable des matieres",
            "Detenteur",
        ],
        context_lines=_pdf_context_lines(request),
        document_ref="Modele 13",
        subtitle=f"Localisation: {selected_location.name} - Responsable: {selected_location.responsible_name or '-'}",
        logo_path=_pdf_logo_path(request),
        extra_paragraphs=[totals_line],
    )


_ALLOWED_UPLOAD_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".xlsx", ".xls", ".odt", ".ods",
    ".txt", ".csv", ".png", ".jpg", ".jpeg",
}
_MAX_DOCUMENT_UPLOAD_BYTES = int(getattr(settings, "DOCUMENT_UPLOAD_MAX_BYTES", 10 * 1024 * 1024))
_CENTRAL_RECAP_CACHE_VERSION_KEY = "central_recap:version"
_ALLOWED_UPLOAD_CONTENT_TYPES = set(
    getattr(
        settings,
        "DOCUMENT_UPLOAD_ALLOWED_CONTENT_TYPES",
        (
            "application/pdf",
            "application/msword",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/vnd.ms-excel",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.oasis.opendocument.text",
            "application/vnd.oasis.opendocument.spreadsheet",
            "text/plain",
            "text/csv",
            "image/png",
            "image/jpeg",
        ),
    )
)


def _safe_filename(name: str) -> str | None:
    """Retourne le nom nettoyé ou None si le nom contient des séquences dangereuses."""
    import re
    # Interdire les séquences de traversée et les caractères spéciaux de chemins
    if ".." in name or "/" in name or "\\" in name:
        return None
    sanitized = re.sub(r"[^\w.\-]", "_", name)
    return sanitized or None


def _is_safe_filename(name: str) -> bool:
    if not name:
        return False
    if ".." in name or "/" in name or "\\" in name:
        return False
    return True


def _unique_filename(path: Path) -> str:
    if not path.exists():
        return path.name
    suffix = path.suffix
    stem = path.stem
    unique = uuid4().hex[:8]
    return f"{stem}_{unique}{suffix}"


def _csv_safe(value: object) -> str:
    text = "" if value is None else str(value)
    if text and text[0] in ("=", "+", "-", "@"):
        return "'" + text
    return text


def _request_ip(request: HttpRequest) -> str:
    forwarded = (request.META.get("HTTP_X_FORWARDED_FOR") or "").strip()
    if forwarded:
        return forwarded.split(",")[0].strip()
    return (request.META.get("REMOTE_ADDR") or "").strip()


def _audit_document_security_rejection(
    request: HttpRequest,
    *,
    reason: str,
    filename: str,
    content_type: str,
    size: int,
) -> None:
    _audit(
        request,
        action=AuditLog.ACTION_UPDATE,
        entity="DocumentSecurity",
        entity_id=filename,
        description=f"Rejet upload document: {reason}",
        metadata={
            "reason": reason,
            "filename": filename,
            "content_type": content_type,
            "size": size,
            "ip": _request_ip(request),
        },
    )


def _get_central_recap_cache_version() -> int:
    version = cache.get(_CENTRAL_RECAP_CACHE_VERSION_KEY)
    if version is None:
        version = 1
        cache.set(_CENTRAL_RECAP_CACHE_VERSION_KEY, version, timeout=None)
    return int(version)


def _invalidate_central_recap_cache() -> None:
    current = _get_central_recap_cache_version()
    cache.set(_CENTRAL_RECAP_CACHE_VERSION_KEY, current + 1, timeout=None)


def _scan_uploaded_document(file_path: Path) -> tuple[bool, str]:
    if not bool(getattr(settings, "DOCUMENT_UPLOAD_SCAN_ENABLED", False)):
        return True, ""

    command_template = tuple(getattr(settings, "DOCUMENT_UPLOAD_SCAN_COMMAND", ()) or ())
    if not command_template:
        return False, "Scan document active, mais DOCUMENT_UPLOAD_SCAN_COMMAND est vide."

    command = [str(part).replace("{path}", str(file_path)) for part in command_template]
    try:
        result = subprocess.run(
            command,
            check=False,
            capture_output=True,
            text=True,
            timeout=30,
        )
    except Exception as exc:  # noqa: BLE001
        return False, f"Erreur de scan document: {exc}"

    if result.returncode != 0:
        details = (result.stderr or result.stdout or "").strip()
        return False, details[:300] if details else "Le scan antivirus a echoue."

    return True, ""


@login_required
def screen_documents_view(request: HttpRequest) -> HttpResponse:
    media_documents_path = Path(settings.MEDIA_ROOT) / "documents"
    media_documents_path.mkdir(parents=True, exist_ok=True)
    is_admin = _is_admin_user(request.user)

    if request.method == "POST" and not is_admin:
        messages.error(request, "Seuls les admins peuvent ajouter des documents.")
        return redirect("screen_documents")

    if request.method == "POST" and is_admin:
        uploaded = request.FILES.get("document")
        if not uploaded:
            messages.error(request, "Aucun fichier sélectionné.")
        else:
            original_name = uploaded.name
            ext = Path(original_name).suffix.lower()
            content_type = (uploaded.content_type or "").lower()
            if ext not in _ALLOWED_UPLOAD_EXTENSIONS:
                _audit_document_security_rejection(
                    request,
                    reason="extension_not_allowed",
                    filename=original_name,
                    content_type=content_type,
                    size=uploaded.size,
                )
                messages.error(request, f"Extension non autorisée : {ext}. Extensions acceptées : {', '.join(sorted(_ALLOWED_UPLOAD_EXTENSIONS))}")
            elif content_type and content_type not in _ALLOWED_UPLOAD_CONTENT_TYPES:
                _audit_document_security_rejection(
                    request,
                    reason="mime_not_allowed",
                    filename=original_name,
                    content_type=content_type,
                    size=uploaded.size,
                )
                messages.error(request, f"Type MIME non autorisé : {content_type}.")
            elif uploaded.size > _MAX_DOCUMENT_UPLOAD_BYTES:
                _audit_document_security_rejection(
                    request,
                    reason="size_exceeded",
                    filename=original_name,
                    content_type=content_type,
                    size=uploaded.size,
                )
                max_mb = _MAX_DOCUMENT_UPLOAD_BYTES / (1024 * 1024)
                messages.error(request, f"Fichier trop volumineux. Taille maximale autorisée : {max_mb:.0f} MB.")
            else:
                safe_name = _safe_filename(original_name)
                if not safe_name:
                    _audit_document_security_rejection(
                        request,
                        reason="invalid_filename",
                        filename=original_name,
                        content_type=content_type,
                        size=uploaded.size,
                    )
                    messages.error(request, "Nom de fichier invalide.")
                else:
                    dest = media_documents_path / safe_name
                    if dest.exists():
                        safe_name = _unique_filename(dest)
                        dest = media_documents_path / safe_name
                    with dest.open("wb") as fout:
                        for chunk in uploaded.chunks():
                            fout.write(chunk)
                    scan_ok, scan_message = _scan_uploaded_document(dest)
                    if not scan_ok:
                        _audit_document_security_rejection(
                            request,
                            reason="scan_failed",
                            filename=safe_name,
                            content_type=content_type,
                            size=uploaded.size,
                        )
                        try:
                            dest.unlink(missing_ok=True)
                        except Exception:  # noqa: BLE001
                            pass
                        messages.error(request, f"Document rejete par controle de securite. {scan_message}")
                        return redirect("screen_documents")
                    _audit(
                        request,
                        action=AuditLog.ACTION_CREATE,
                        entity="Document",
                        entity_id=safe_name,
                        description=f"Document téléversé : {safe_name}",
                    )
                    messages.success(request, f"Document « {safe_name} » ajouté.")
        return redirect("screen_documents")

    documents = []
    for file_path in sorted(media_documents_path.glob("*")):
        if file_path.is_file():
            size_kb = file_path.stat().st_size / 1024
            documents.append({
                "name": file_path.name,
                "size": f"{size_kb:.1f} KB",
                "url": reverse("document_download", args=[file_path.name]),
            })

    return render(request, "inventory/screen_documents.html", {
        "documents": documents,
        "is_admin": is_admin,
    })


@login_required
def document_delete_view(request: HttpRequest) -> HttpResponse:
    if not _is_admin_user(request.user):
        messages.error(request, "Seuls les admins peuvent supprimer des documents.")
        return redirect("screen_documents")
    if request.method != "POST":
        return redirect("screen_documents")

    filename = (request.POST.get("filename") or "").strip()
    if not _is_safe_filename(filename):
        messages.error(request, "Nom de fichier invalide.")
        return redirect("screen_documents")

    media_documents_path = Path(settings.MEDIA_ROOT) / "documents"
    target = media_documents_path / filename
    # S'assurer que le fichier résolu est bien dans le dossier documents (anti path-traversal)
    try:
        target.resolve().relative_to(media_documents_path.resolve())
    except ValueError:
        messages.error(request, "Opération non autorisée.")
        return redirect("screen_documents")

    if target.is_file():
        target.unlink()
        _audit(
            request,
            action=AuditLog.ACTION_DELETE,
            entity="Document",
            entity_id=filename,
            description=f"Document supprimé : {filename}",
        )
        messages.success(request, f"Document « {filename} » supprimé.")
    else:
        messages.error(request, "Document introuvable.")
    return redirect("screen_documents")


@login_required
def document_download_view(request: HttpRequest, filename: str) -> HttpResponse:
    filename = (filename or "").strip()
    if not _is_safe_filename(filename):
        messages.error(request, "Nom de fichier invalide.")
        return redirect("screen_documents")

    media_documents_path = Path(settings.MEDIA_ROOT) / "documents"
    target = media_documents_path / filename
    try:
        target.resolve().relative_to(media_documents_path.resolve())
    except ValueError:
        messages.error(request, "Operation non autorisee.")
        return redirect("screen_documents")

    if not target.is_file():
        return HttpResponse("Document introuvable.", status=404, content_type="text/plain; charset=utf-8")

    return FileResponse(target.open("rb"), as_attachment=True, filename=filename)


@login_required
def user_guide_view(request: HttpRequest) -> HttpResponse:
    guide_source = (request.GET.get("source") or "").strip().lower()
    context = {
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
        "guide_source": guide_source,
    }
    return render(request, "inventory/user_guide.html", context)


@login_required
def dashboard(request: HttpRequest) -> HttpResponse:
    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    pv_filled_count = sum((1 for item in summary if item.pv_is_filled), 0)
    pv_missing_count = sum((1 for item in summary if not item.pv_is_filled), 0)
    nonzero_gap_count = sum((1 for item in summary if item.pv_is_filled and item.inventory_gap != Decimal("0")), 0)
    negative_stock_count = sum((1 for item in summary if item.remaining_qty < Decimal("0")), 0)
    cmup_zero_with_stock_count = sum((1 for item in summary if item.remaining_qty > Decimal("0") and item.cmup <= Decimal("0")), 0)
    
    # Filtre PV
    pv_filter = request.GET.get("pv", "all")
    if pv_filter == "filled":
        filtered_summary = [item for item in summary if item.pv_is_filled]
    elif pv_filter == "missing":
        filtered_summary = [item for item in summary if not item.pv_is_filled]
    else:
        filtered_summary = summary
    
    context = {
        "metrics": dashboard_metrics(structure=request.current_structure, fiscal_year=request.current_fiscal_year),
        "summary": filtered_summary,
        "pv_filter": pv_filter,
        "quality_metrics": {
            "pv_filled_count": pv_filled_count,
            "pv_missing_count": pv_missing_count,
            "nonzero_gap_count": nonzero_gap_count,
            "negative_stock_count": negative_stock_count,
            "cmup_zero_with_stock_count": cmup_zero_with_stock_count,
        },
        "current_structure": request.current_structure,
        "current_fiscal_year": request.current_fiscal_year,
    }
    return render(request, "inventory/dashboard.html", context)


@login_required
def profile_view(request: HttpRequest) -> HttpResponse:
    profile = _get_or_create_profile(request.user)
    if profile is None:
        messages.error(request, "Profil utilisateur introuvable.")
        return redirect("menu")
    if request.method == "POST":
        form = ProfileForm(request.POST, instance=profile)
        if form.is_valid():
            form.save()
            messages.success(request, "Profil mis a jour.")
            return redirect("profile")
    else:
        form = ProfileForm(instance=profile)
    return render(request, "inventory/profile.html", {"form": form})


@login_required
def user_management_view(request: HttpRequest) -> HttpResponse:
    if not _is_admin_user(request.user):
        messages.error(request, "Seuls les admins peuvent gerer les utilisateurs.")
        return redirect("screen_admin")

    for user in User.objects.all().only("id"):
        UserProfile.objects.get_or_create(user=user)

    editing_user = None
    edit_form = None

    if request.method == "POST":
        action = (request.POST.get("action") or "create").strip()
        if action == "create":
            form = UserCreationWithProfileForm(request.POST)
            if form.is_valid():
                created_user = form.save()
                messages.success(request, f"Utilisateur {created_user.username} cree.")
                return redirect("user_management")
        elif action == "update":
            user_id = request.POST.get("user_id")
            editing_user = get_object_or_404(User, pk=user_id)
            edit_form = UserUpdateWithProfileForm(request.POST, user_instance=editing_user)
            form = UserCreationWithProfileForm()
            if edit_form.is_valid():
                updated_user = edit_form.save()
                messages.success(request, f"Utilisateur {updated_user.username} mis a jour.")
                return redirect("user_management")
        elif action == "delete":
            user_id = request.POST.get("user_id")
            target_user = get_object_or_404(User, pk=user_id)
            if target_user.pk == request.user.pk:
                messages.error(request, "Vous ne pouvez pas supprimer votre propre compte.")
            else:
                username = target_user.username
                target_user.delete()
                messages.success(request, f"Utilisateur {username} supprime.")
            return redirect("user_management")
        else:
            form = UserCreationWithProfileForm()
    else:
        form = UserCreationWithProfileForm()
        edit_id = (request.GET.get("edit") or "").strip()
        if edit_id:
            editing_user = get_object_or_404(User, pk=edit_id)
            edit_form = UserUpdateWithProfileForm(user_instance=editing_user)

    users = User.objects.select_related(
        "profile",
        "profile__default_structure",
        "profile__assigned_structure_type",
    ).prefetch_related("profile__assigned_structures").order_by("username")
    return render(
        request,
        "inventory/user_management.html",
        {
            "form": form,
            "users": users,
            "editing_user": editing_user,
            "edit_form": edit_form,
        },
    )


@login_required
def document_security_logs_view(request: HttpRequest) -> HttpResponse:
    if not _is_admin_user(request.user):
        messages.error(request, "Seuls les admins peuvent consulter ce journal.")
        return redirect("screen_admin")

    filters = _document_security_log_filters_from_request(request)
    logs_qs = _build_document_security_logs_queryset(filters=filters)

    reasons = [
        value
        for value in AuditLog.objects.filter(entity="DocumentSecurity")
        .values_list("metadata__reason", flat=True)
        .distinct()
        .order_by("metadata__reason")
        if value
    ]

    paginator = Paginator(logs_qs, 30)
    logs = paginator.get_page(request.GET.get("page"))
    retention_days_default = int(getattr(settings, "DOCUMENT_SECURITY_LOG_RETENTION_DAYS", 180))

    return render(
        request,
        "inventory/document_security_logs.html",
        {
            "logs": logs,
            "reasons": reasons,
            "filters": filters,
            "query_without_page": _query_without_page(request),
            "retention_days_default": retention_days_default,
            "current_structure": request.current_structure,
            "current_fiscal_year": request.current_fiscal_year,
        },
    )


@login_required
def document_security_logs_csv_view(request: HttpRequest) -> HttpResponse:
    if not _is_admin_user(request.user):
        messages.error(request, "Seuls les admins peuvent exporter ce journal.")
        return redirect("screen_admin")

    filters = _document_security_log_filters_from_request(request)
    logs_qs = _build_document_security_logs_queryset(filters=filters)

    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="journal_securite_documents.csv"'

    writer = csv.writer(response, delimiter=";")
    writer.writerow(["Date", "Utilisateur", "Raison", "Fichier", "MIME", "Taille", "IP"])
    for log in logs_qs:
        writer.writerow(
            [
                _csv_safe(timezone.localtime(log.created_at).strftime("%Y-%m-%d %H:%M:%S")),
                _csv_safe(log.user.username if log.user else ""),
                _csv_safe(log.metadata.get("reason", "")),
                _csv_safe(log.metadata.get("filename", "")),
                _csv_safe(log.metadata.get("content_type", "")),
                _csv_safe(log.metadata.get("size", "")),
                _csv_safe(log.metadata.get("ip", "")),
            ]
        )
    return response


@login_required
def document_security_logs_purge_view(request: HttpRequest) -> HttpResponse:
    if not _is_admin_user(request.user):
        messages.error(request, "Seuls les admins peuvent purger ce journal.")
        return redirect("screen_admin")
    if request.method != "POST":
        return redirect("document_security_logs")

    retention_days_raw = (request.POST.get("retention_days") or "").strip()
    try:
        retention_days = int(retention_days_raw or int(getattr(settings, "DOCUMENT_SECURITY_LOG_RETENTION_DAYS", 180)))
    except ValueError:
        messages.error(request, "Valeur de retention invalide.")
        return redirect("document_security_logs")

    if retention_days < 1:
        messages.error(request, "La retention doit etre >= 1 jour.")
        return redirect("document_security_logs")

    cutoff = timezone.now() - timezone.timedelta(days=retention_days)
    to_delete_qs = AuditLog.objects.filter(entity="DocumentSecurity", created_at__lt=cutoff)
    deleted_count = to_delete_qs.count()
    to_delete_qs.delete()

    _audit(
        request,
        action=AuditLog.ACTION_DELETE,
        entity="DocumentSecurity",
        entity_id="bulk_purge",
        description=f"Purge des logs securite documents > {retention_days} jours",
        metadata={
            "retention_days": retention_days,
            "deleted_count": deleted_count,
            "cutoff": cutoff.isoformat(),
        },
    )
    messages.success(request, f"Purge terminee: {deleted_count} evenement(s) supprime(s).")
    return redirect("document_security_logs")


def _document_security_log_filters_from_request(request: HttpRequest) -> dict[str, str]:
    return {
        "reason": (request.GET.get("reason") or "").strip(),
        "username": (request.GET.get("username") or "").strip(),
        "ip": (request.GET.get("ip") or "").strip(),
        "date_from": (request.GET.get("date_from") or "").strip(),
        "date_to": (request.GET.get("date_to") or "").strip(),
    }


def _build_document_security_logs_queryset(*, filters: dict[str, str]):
    logs_qs = (
        AuditLog.objects.filter(entity="DocumentSecurity")
        .select_related("user", "structure", "fiscal_year")
        .order_by("-created_at")
    )

    reason = filters.get("reason", "")
    username = filters.get("username", "")
    ip = filters.get("ip", "")
    date_from = filters.get("date_from", "")
    date_to = filters.get("date_to", "")

    if reason:
        logs_qs = logs_qs.filter(metadata__reason=reason)
    if username:
        logs_qs = logs_qs.filter(user__username__icontains=username)
    if ip:
        logs_qs = logs_qs.filter(metadata__ip__icontains=ip)
    if date_from:
        logs_qs = logs_qs.filter(created_at__date__gte=date_from)
    if date_to:
        logs_qs = logs_qs.filter(created_at__date__lte=date_to)

    return logs_qs



@login_required
def materials_view(request: HttpRequest) -> HttpResponse:
    """Read-only view of the nomenclature for the current structure/fiscal year."""
    q = (request.GET.get("q") or "").strip()
    sort = request.GET.get("sort") or "code_asc"
    sort_mapping = {
        "code_asc": "account_code",
        "code_desc": "-account_code",
        "name_asc": "name",
        "name_desc": "-name",
        "group_asc": "group_code",
        "group_desc": "-group_code",
    }
    order_by = sort_mapping.get(sort, "account_code")

    materials_qs = Material.objects.filter(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    if q:
        materials_qs = materials_qs.filter(
            Q(account_code__icontains=q) | Q(name__icontains=q) | Q(group_code__icontains=q)
        )
    materials_qs = materials_qs.order_by(order_by)
    paginator = Paginator(materials_qs, 25)
    materials = paginator.get_page(request.GET.get("page"))

    return render(
        request,
        "inventory/materials.html",
        {
            "materials": materials,
            "filters": {"q": q, "sort": sort},
            "query_without_page": _query_without_page(request),
            "is_admin": _is_admin_user(request.user),
        },
    )


@login_required
def nomenclature_admin_view(request: HttpRequest) -> HttpResponse:
    """Admin-only: manage the global nomenclature and propagate changes to all structures."""
    if not _is_admin_user(request.user):
        messages.error(request, "Acces reserve aux administrateurs.")
        return redirect("screen_admin")

    edit_item = None
    form = NomenclatureItemForm()

    # Handle GET for edit
    edit_code = request.GET.get("edit")
    if edit_code:
        edit_item = get_object_or_404(NomenclatureItem, account_code=edit_code)
        form = NomenclatureItemForm(instance=edit_item)

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "add":
            form = NomenclatureItemForm(request.POST)
            if form.is_valid():
                item = form.save()
                count = propagate_nomenclature_add(item)
                _audit(
                    request,
                    action=AuditLog.ACTION_CREATE,
                    entity="NomenclatureItem",
                    entity_id=str(item.id),
                    description=f"Ajout nomenclature {item.account_code} - propagation: {count} fiche(s)",
                )
                messages.success(request, f"Compte '{item.account_code}' ajoute et propage vers {count} fiche(s).")
                return redirect("nomenclature_admin")
            messages.error(request, "Impossible d'enregistrer ce compte. Corrigez les erreurs du formulaire.")

        elif action == "edit":
            account_code = request.POST.get("account_code", "").strip()
            edit_item = get_object_or_404(NomenclatureItem, account_code=account_code)
            form = NomenclatureItemForm(request.POST, instance=edit_item)
            if form.is_valid():
                form.save()
                _audit(
                    request,
                    action=AuditLog.ACTION_UPDATE,
                    entity="NomenclatureItem",
                    entity_id=str(edit_item.id),
                    description=f"Modification nomenclature {edit_item.account_code}",
                )
                messages.success(request, f"Compte '{edit_item.account_code}' modifié avec succès.")
                return redirect("nomenclature_admin")
            messages.error(request, "Impossible de modifier ce compte. Corrigez les erreurs du formulaire.")

        elif action == "delete":
            account_code = request.POST.get("account_code", "").strip()
            item = get_object_or_404(NomenclatureItem, account_code=account_code)
            try:
                propagate_nomenclature_delete(account_code)
                _audit(
                    request,
                    action=AuditLog.ACTION_DELETE,
                    entity="NomenclatureItem",
                    entity_id=str(item.id),
                    description=f"Suppression nomenclature {account_code}",
                )
                item.delete()
                messages.success(request, f"Compte '{account_code}' supprime de la nomenclature et de tous les services.")
            except ValueError as exc:
                messages.error(request, str(exc))
            return redirect("nomenclature_admin")

        elif action == "sync":
            total = 0
            for structure in Structure.objects.filter(is_active=True):
                total += sync_all_nomenclature_to_structure(structure)
            _audit(
                request,
                action=AuditLog.ACTION_UPDATE,
                entity="NomenclatureItem",
                description=f"Synchronisation nomenclature: {total} fiche(s) creee(s)",
            )
            messages.success(request, f"Synchronisation terminee : {total} fiche(s) creee(s).")
            return redirect("nomenclature_admin")

    q = (request.GET.get("q") or "").strip()
    items_qs = NomenclatureItem.objects.all()
    if q:
        items_qs = items_qs.filter(
            Q(account_code__icontains=q) | Q(name__icontains=q) | Q(group_code__icontains=q)
        )
    paginator = Paginator(items_qs, 25)
    items = paginator.get_page(request.GET.get("page"))

    return render(
        request,
        "inventory/nomenclature.html",
        {
            "form": form,
            "items": items,
            "q": q,
            "query_without_page": _query_without_page(request),
            "edit_item": edit_item,
        },
    )


@login_required
def nomenclature_import_view(request: HttpRequest) -> HttpResponse:
    """Admin-only: import nomenclature items from an Excel (.xlsx), Word (.docx) or CSV file.

    N'effectue PAS de propagation automatique vers les structures.
    L'administrateur utilise ensuite le bouton 'Synchroniser' pour propager.
    """
    if not _is_admin_user(request.user):
        messages.error(request, "Accès réservé aux administrateurs.")
        return redirect("nomenclature_admin")

    if request.method != "POST":
        return redirect("nomenclature_admin")

    uploaded = request.FILES.get("import_file")
    if not uploaded:
        messages.error(request, "Aucun fichier sélectionné.")
        return redirect("nomenclature_admin")

    filename = uploaded.name.lower()
    rows: list[dict] = []

    try:
        if filename.endswith(".xlsx") or filename.endswith(".xls"):
            rows = _parse_nomenclature_excel(uploaded)
        elif filename.endswith(".docx"):
            rows = _parse_nomenclature_docx(uploaded)
        elif filename.endswith(".csv"):
            rows = _parse_nomenclature_csv(uploaded)
        else:
            messages.error(request, "Format non supporté. Utilisez un fichier .xlsx, .docx ou .csv.")
            return redirect("nomenclature_admin")
    except Exception as exc:  # noqa: BLE001
        messages.error(request, f"Erreur lors de la lecture du fichier : {exc}")
        return redirect("nomenclature_admin")

    if not rows:
        messages.error(request, "Aucune donnée trouvée dans le fichier. Vérifiez le format.")
        return redirect("nomenclature_admin")

    created_count = 0
    updated_count = 0
    error_count = 0

    for row in rows:
        code = (row.get("account_code") or "").strip()
        name = (row.get("name") or "").strip()
        unit = (row.get("unit") or "").strip()
        group_code = (row.get("group_code") or "").strip()

        if not code or not name:
            error_count += 1
            continue

        _, created = NomenclatureItem.objects.update_or_create(
            account_code=code,
            defaults={"name": name, "unit": unit or "U", "group_code": group_code},
        )
        if created:
            created_count += 1
        else:
            updated_count += 1

    _audit(
        request,
        action=AuditLog.ACTION_CREATE,
        entity="NomenclatureItem",
        description=(
            f"Import nomenclature depuis '{uploaded.name}' : "
            f"{created_count} créé(s), {updated_count} mis à jour, {error_count} ignoré(s)."
        ),
    )

    messages.success(
        request,
        f"Import terminé : {created_count} compte(s) créé(s), {updated_count} mis à jour, "
        f"{error_count} ligne(s) ignorée(s). "
        "Utilisez 'Synchroniser vers tous les services' pour propager vers les structures.",
    )
    return redirect("nomenclature_admin")


def _parse_nomenclature_excel(file_obj) -> list[dict]:
    """Parse un fichier Excel et retourne une liste de dicts avec les colonnes de nomenclature.

    Colonnes attendues (insensible à la casse) :
      code_compte / compte / account_code
      intitule / designation / nom / name
      unite / unit / unité
      groupe / group_code (optionnel)
    """
    import openpyxl  # noqa: PLC0415

    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)

    # Détection des en-têtes sur la première ligne non vide
    header_row = None
    data_rows = []
    for row in rows_iter:
        if any(cell is not None for cell in row):
            if header_row is None:
                header_row = [str(c).strip().lower() if c else "" for c in row]
            else:
                data_rows.append(row)

    if not header_row:
        return []

    def _col(header: list[str], *names: str) -> int | None:
        for name in names:
            if name in header:
                return header.index(name)
        return None

    idx_code = _col(header_row, "code_compte", "compte", "account_code", "code")
    idx_name = _col(header_row, "intitule", "intitulé", "designation", "désignation", "nom", "name", "libelle", "libellé")
    idx_unit = _col(header_row, "unite", "unité", "unit", "u")
    idx_group = _col(header_row, "groupe", "group_code", "group", "code_groupe")

    if idx_code is None or idx_name is None:
        raise ValueError(
            "Colonnes 'code_compte' et 'intitulé' introuvables. "
            "Assurez-vous que la première ligne contient les en-têtes."
        )

    result = []
    for row in data_rows:
        code = str(row[idx_code]).strip() if row[idx_code] is not None else ""
        name = str(row[idx_name]).strip() if row[idx_name] is not None else ""
        unit = str(row[idx_unit]).strip() if idx_unit is not None and row[idx_unit] is not None else ""
        group = str(row[idx_group]).strip() if idx_group is not None and row[idx_group] is not None else ""
        if code and name:
            result.append({"account_code": code, "name": name, "unit": unit, "group_code": group})
    return result


def _parse_nomenclature_docx(file_obj) -> list[dict]:
    """Parse un fichier Word (.docx) : recherche la première table avec des colonnes de nomenclature.

    La première ligne de la table doit contenir les en-têtes (même noms qu'Excel).
    """
    from docx import Document  # noqa: PLC0415

    doc = Document(file_obj)

    if not doc.tables:
        raise ValueError("Aucun tableau trouvé dans le document Word.")

    # Chercher la table qui contient les colonnes code + nom
    for table in doc.tables:
        if not table.rows:
            continue

        header_cells = [cell.text.strip().lower() for cell in table.rows[0].cells]

        def _col(header: list[str], *names: str) -> int | None:
            for name in names:
                if name in header:
                    return header.index(name)
            return None

        idx_code = _col(header_cells, "code_compte", "compte", "account_code", "code")
        idx_name = _col(header_cells, "intitule", "intitulé", "designation", "désignation", "nom", "name", "libelle", "libellé")
        idx_unit = _col(header_cells, "unite", "unité", "unit", "u")
        idx_group = _col(header_cells, "groupe", "group_code", "group", "code_groupe")

        if idx_code is None or idx_name is None:
            continue  # essayer la prochaine table

        result = []
        for row in table.rows[1:]:
            cells = row.cells
            code = cells[idx_code].text.strip() if idx_code < len(cells) else ""
            name = cells[idx_name].text.strip() if idx_name < len(cells) else ""
            unit = cells[idx_unit].text.strip() if idx_unit is not None and idx_unit < len(cells) else ""
            group = cells[idx_group].text.strip() if idx_group is not None and idx_group < len(cells) else ""
            if code and name:
                result.append({"account_code": code, "name": name, "unit": unit, "group_code": group})
        return result

    raise ValueError(
        "Aucun tableau avec les colonnes 'code_compte' et 'intitulé' trouvé dans le document Word."
    )


def _parse_nomenclature_csv(file_obj) -> list[dict]:
    """Parse un fichier CSV (séparateur virgule ou point-virgule, encodage UTF-8 ou latin-1)."""
    import io  # noqa: PLC0415

    raw = file_obj.read()
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Encodage du fichier CSV non reconnu (essayez UTF-8 ou Latin-1).")

    # Détection du séparateur
    sep = ";" if text.count(";") >= text.count(",") else ","

    reader = csv.DictReader(io.StringIO(text), delimiter=sep)
    # Normaliser les clés
    normalized_fieldnames = {k: k.strip().lower() for k in (reader.fieldnames or [])}

    def _resolve(row: dict, *names: str) -> str:
        for name in names:
            for original_key, norm_key in normalized_fieldnames.items():
                if norm_key in names or norm_key == name:
                    return str(row.get(original_key) or "").strip()
        return ""

    result = []
    for row in reader:
        code = _resolve(row, "code_compte", "compte", "account_code", "code")
        name = _resolve(row, "intitule", "intitulé", "designation", "désignation", "nom", "name", "libelle", "libellé")
        unit = _resolve(row, "unite", "unité", "unit", "u")
        group = _resolve(row, "groupe", "group_code", "group", "code_groupe")
        if code and name:
            result.append({"account_code": code, "name": name, "unit": unit, "group_code": group})
    return result


@login_required
def voucher_create_view(request: HttpRequest) -> HttpResponse:
    store_location_types = [
        Location.TYPE_STORE_EQUIPMENT,
        Location.TYPE_STORE_SUPPLIES,
        Location.TYPE_STORE_PRODUCTS,
    ]
    store_locations = list(
        Location.objects.filter(
            structure=request.current_structure,
            is_active=True,
            location_type__in=store_location_types,
        ).order_by("name")
    )
    selected_store_location_id = request.POST.get("storage_location_id", "") if request.method == "POST" else ""

    materials = Material.objects.filter(
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    ).order_by("account_code")
    summary_by_code = {
        row.account_code: row
        for row in compute_stock_summary(
            structure=request.current_structure,
            fiscal_year=request.current_fiscal_year,
        )
    }
    materials_with_cmup = [
        {
            "id": material.id,
            "account_code": material.account_code,
            "name": material.name,
            "cmup": summary_by_code.get(material.account_code).cmup if summary_by_code.get(material.account_code) else Decimal("0"),
        }
        for material in materials
    ]
    line_range = range(3)
    if request.method == "POST":
        if not _can_edit(request.user):
            messages.error(request, "Vous n'avez pas les droits de saisie.")
            return redirect("voucher_create")
        if not _ensure_open_fiscal_year(request, redirect_to="voucher_create"):
            return redirect("voucher_create")
        form = VoucherHeaderForm(request.POST)
        # Préparer les valeurs déjà saisies pour chaque ligne
        posted_lines = []
        for i in line_range:
            posted_lines.append({
                "material_id": request.POST.get(f"line_material_{i}", ""),
                "specification": request.POST.get(f"line_specification_{i}", ""),
                "quantity": request.POST.get(f"line_quantity_{i}", ""),
                "unit_price": request.POST.get(f"line_unit_price_{i}", ""),
            })
        if form.is_valid():
            try:
                lines = parse_voucher_lines(request.POST)
                storage_location = None
                storage_location_id = (request.POST.get("storage_location_id") or "").strip()
                if storage_location_id:
                    storage_location = get_object_or_404(
                        Location,
                        pk=storage_location_id,
                        structure=request.current_structure,
                        is_active=True,
                    )
                vouchers = create_voucher(
                    user=request.user,
                    voucher_type=form.cleaned_data["voucher_type"],
                    operation_date=form.cleaned_data["operation_date"],
                    source_or_destination=form.cleaned_data["source_or_destination"],
                    storage_location=storage_location,
                    comments=form.cleaned_data["comments"],
                    lines=lines,
                    structure=request.current_structure,
                    fiscal_year=request.current_fiscal_year,
                )
            except (ValueError, ValidationError) as error:
                messages.error(request, str(error))
            else:
                for voucher in vouchers:
                    _audit(
                        request,
                        action=AuditLog.ACTION_CREATE,
                        entity="Voucher",
                        entity_id=str(voucher.id),
                        description=f"Creation bon {voucher.number}",
                        metadata={"type": voucher.voucher_type, "lines": voucher.lines.count()},
                    )
                _invalidate_central_recap_cache()
                if len(vouchers) == 1:
                    voucher = vouchers[0]
                    messages.success(request, f"Bon {voucher.number} enregistre.")
                    return redirect("voucher_detail", pk=voucher.pk)
                nums = ", ".join(f"N° {v.number}" for v in vouchers)
                messages.success(request, f"{len(vouchers)} bons enregistres : {nums}.")
                return redirect("voucher_list")
        else:
            # En cas d'erreur, on conserve les valeurs postées
            context = {
                "form": form,
                "materials_with_cmup": materials_with_cmup,
                "line_range": line_range,
                "can_edit": _can_edit(request.user),
                "is_fiscal_year_closed": bool(request.current_fiscal_year and request.current_fiscal_year.is_closed),
                "prefill": prefill,
                "selected_material_ids": [l["material_id"] for l in posted_lines],
                "posted_lines": posted_lines,
                "store_locations": store_locations,
                "selected_store_location_id": selected_store_location_id,
            }
            return render(request, "inventory/voucher_form.html", context)
    else:
        initial = {}
        prefill = {
            "material_id": request.GET.get("material_id", ""),
            "quantity": request.GET.get("quantity", ""),
            "unit_price": request.GET.get("unit_price", ""),
            "regulation_reason": request.GET.get("regulation_reason", ""),
        }
        voucher_type = request.GET.get("voucher_type", "")
        if voucher_type in {Voucher.TYPE_ENTRY, Voucher.TYPE_FINAL_EXIT, Voucher.TYPE_TEMP_EXIT}:
            initial["voucher_type"] = voucher_type
        if prefill["regulation_reason"]:
            initial["comments"] = prefill["regulation_reason"]
            initial["source_or_destination"] = "Regularisation"
        form = VoucherHeaderForm(initial=initial)
    if request.method == "POST":
        prefill = {"material_id": "", "quantity": "", "unit_price": "", "regulation_reason": ""}
    # Prépare la liste des matières sélectionnées pour chaque ligne
    selected_material_ids = []
    if request.method == "POST":
        for i in line_range:
            selected_material_ids.append(request.POST.get(f"line_material_{i}", ""))
    else:
        # Préremplissage GET ou vide
        for i in line_range:
            if i == 0 and prefill["material_id"]:
                selected_material_ids.append(str(prefill["material_id"]))
            else:
                selected_material_ids.append("")
    context = {
        "form": form,
        "materials_with_cmup": materials_with_cmup,
        "line_range": line_range,
        "can_edit": _can_edit(request.user),
        "is_fiscal_year_closed": bool(request.current_fiscal_year and request.current_fiscal_year.is_closed),
        "prefill": prefill,
        "selected_material_ids": selected_material_ids,
        "store_locations": store_locations,
        "selected_store_location_id": selected_store_location_id,
    }
    return render(request, "inventory/voucher_form.html", context)


@login_required
def voucher_list_view(request: HttpRequest) -> HttpResponse:
    q = (request.GET.get("q") or "").strip()
    voucher_type = (request.GET.get("voucher_type") or "").strip()
    temp_return_status = (request.GET.get("temp_return_status") or "").strip()
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    sort = request.GET.get("sort") or "date_desc"
    sort_mapping = {
        "date_desc": "-operation_date",
        "date_asc": "operation_date",
        "number_desc": "-number",
        "number_asc": "number",
    }
    order_by = sort_mapping.get(sort, "-operation_date")

    vouchers_qs = (
        Voucher.objects.filter(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
        .select_related("created_by")
        .prefetch_related("lines", "lines__material")
    )
    if voucher_type:
        vouchers_qs = vouchers_qs.filter(voucher_type=voucher_type)
    if temp_return_status in {"en_cours", "retourne"}:
        vouchers_qs = vouchers_qs.filter(voucher_type=Voucher.TYPE_TEMP_EXIT)
        if temp_return_status == "retourne":
            vouchers_qs = vouchers_qs.filter(temp_return_date__isnull=False)
        else:
            vouchers_qs = vouchers_qs.filter(temp_return_date__isnull=True)
    if date_from:
        vouchers_qs = vouchers_qs.filter(operation_date__gte=date_from)
    if date_to:
        vouchers_qs = vouchers_qs.filter(operation_date__lte=date_to)
    if q:
        query_filter = Q(source_or_destination__icontains=q) | Q(comments__icontains=q)
        if q.isdigit():
            query_filter |= Q(number=int(q))
        vouchers_qs = vouchers_qs.filter(query_filter)

    vouchers_qs = vouchers_qs.order_by(order_by, "-number")
    paginator = Paginator(vouchers_qs, 20)
    vouchers = paginator.get_page(request.GET.get("page"))

    return render(
        request,
        "inventory/voucher_list.html",
        {
            "vouchers": vouchers,
            "filters": {
                "q": q,
                "voucher_type": voucher_type,
                "temp_return_status": temp_return_status,
                "date_from": date_from,
                "date_to": date_to,
                "sort": sort,
            },
            "query_without_page": _query_without_page(request),
        },
    )

@login_required
def voucher_detail_view(request: HttpRequest, pk: int) -> HttpResponse:
    if request.current_structure is None or request.current_fiscal_year is None:
        messages.error(request, "Selectionnez d'abord un service et un exercice.")
        return redirect("menu")
    voucher = get_object_or_404(
        Voucher.objects.select_related("created_by").prefetch_related("lines", "lines__material"),
        pk=pk,
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    )
    can_update_temp_return = (
        _can_edit(request.user)
        and voucher.voucher_type == Voucher.TYPE_TEMP_EXIT
        and not bool(request.current_fiscal_year and request.current_fiscal_year.is_closed)
    )

    if request.method == "POST":
        if not can_update_temp_return:
            messages.error(request, "Vous n'avez pas le droit de modifier le statut de retour de ce bon.")
            return redirect("voucher_detail", pk=voucher.pk)

        action = (request.POST.get("action") or "").strip()
        if action != "update_temp_return_status":
            messages.error(request, "Action non reconnue.")
            return redirect("voucher_detail", pk=voucher.pk)

        status = (request.POST.get("temp_return_status") or "").strip()
        return_note = (request.POST.get("temp_return_note") or "").strip()
        return_date_raw = (request.POST.get("temp_return_date") or "").strip()

        if status not in {"en_cours", "retourne"}:
            messages.error(request, "Statut de retour invalide.")
            return redirect("voucher_detail", pk=voucher.pk)

        if status == "retourne":
            if not return_date_raw:
                messages.error(request, "La date de retour est obligatoire quand le statut est 'retourne'.")
                return redirect("voucher_detail", pk=voucher.pk)
            try:
                return_date = date.fromisoformat(return_date_raw)
            except ValueError:
                messages.error(request, "La date de retour est invalide.")
                return redirect("voucher_detail", pk=voucher.pk)
        else:
            return_date = None

        voucher.temp_return_date = return_date
        voucher.temp_return_note = return_note
        voucher.save(update_fields=["temp_return_date", "temp_return_note", "updated_at"])
        _audit(
            request,
            action=AuditLog.ACTION_UPDATE,
            entity="Voucher",
            entity_id=str(voucher.id),
            description="Mise a jour statut retour sortie provisoire",
            metadata={
                "voucher_number": voucher.number,
                "status": status,
                "return_date": return_date.isoformat() if return_date else "",
            },
        )
        messages.success(request, "Statut de retour mis a jour.")
        return redirect("voucher_detail", pk=voucher.pk)

    total_quantity = sum(line.quantity for line in voucher.lines.all())
    total_amount = voucher.total_amount
    total_quantity_words = _number_to_french_words(total_quantity)
    total_amount_words = _number_to_french_words(total_amount)
    context = {
        "voucher": voucher,
        "total_quantity": total_quantity,
        "total_amount": total_amount,
        "total_quantity_words": total_quantity_words,
        "total_amount_words": total_amount_words,
        "can_update_temp_return": can_update_temp_return,
        "print_logo_url": request.current_structure.logo.url if request.current_structure and request.current_structure.logo else "",
        **default_site_context(structure=request.current_structure, fiscal_year=request.current_fiscal_year),
    }
    return render(request, "inventory/voucher_detail.html", context)


@login_required
def journal_view(request: HttpRequest) -> HttpResponse:
    q = (request.GET.get("q") or "").strip()
    voucher_type = (request.GET.get("voucher_type") or "").strip()
    date_from = (request.GET.get("date_from") or "").strip()
    date_to = (request.GET.get("date_to") or "").strip()
    sort = request.GET.get("sort") or "date_asc"

    vouchers = Voucher.objects.filter(structure=request.current_structure, fiscal_year=request.current_fiscal_year).prefetch_related("lines", "lines__material")
    if voucher_type:
        vouchers = vouchers.filter(voucher_type=voucher_type)
    if date_from:
        vouchers = vouchers.filter(operation_date__gte=date_from)
    if date_to:
        vouchers = vouchers.filter(operation_date__lte=date_to)
    vouchers = vouchers.order_by("operation_date", "number")

    rows = []
    for voucher in vouchers:
        for line in voucher.lines.all():
            rows.append(
                {
                    "date": voucher.operation_date,
                    "number": voucher.number,
                    "type": voucher.get_voucher_type_display(),
                    "account_code": line.material.account_code,
                    "material_name": line.material_name,
                    "quantity": line.quantity,
                    "unit": line.unit,
                    "unit_price": line.unit_price,
                    "amount": line.amount,
                    "designation": voucher.source_or_destination,
                }
            )

    if q:
        q_l = q.lower()
        rows = [
            row
            for row in rows
            if q_l in str(row["number"]).lower()
            or q_l in row["account_code"].lower()
            or q_l in row["material_name"].lower()
            or q_l in (row["designation"] or "").lower()
        ]

    reverse_sort = sort == "date_desc"
    rows.sort(key=lambda item: (item["date"], item["number"]), reverse=reverse_sort)

    paginator = Paginator(rows, 40)
    rows_page = paginator.get_page(request.GET.get("page"))

    return render(
        request,
        "inventory/journal.html",
        {
            "rows": rows_page,
            "filters": {
                "q": q,
                "voucher_type": voucher_type,
                "date_from": date_from,
                "date_to": date_to,
                "sort": sort,
            },
            "query_without_page": _query_without_page(request),
        },
    )


@login_required
def ledger_view(request: HttpRequest) -> HttpResponse:
    materials = Material.objects.filter(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    selected_code = request.GET.get("material") or (materials.first().account_code if materials else "")
    ledger = build_grand_ledger(selected_code, structure=request.current_structure, fiscal_year=request.current_fiscal_year) if selected_code else []
    return render(
        request,
        "inventory/ledger.html",
        {"materials": materials, "selected_code": selected_code, "ledger": ledger},
    )


@login_required
def recap_view(request: HttpRequest) -> HttpResponse:
    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    from decimal import Decimal
    ZERO = Decimal("0")
    quarterly_reports = []
    report_form = None
    suggested_quarter = 1
    if request.current_structure and request.current_fiscal_year:
        quarterly_reports = list(
            QuarterlyReport.objects.filter(
                structure=request.current_structure,
                fiscal_year=request.current_fiscal_year,
            ).order_by("quarter")
        )
        used_quarters = {report.quarter for report in quarterly_reports}
        for quarter in (1, 2, 3):
            if quarter not in used_quarters:
                suggested_quarter = quarter
                break
        report_form = QuarterlyReportForm(
            initial={"quarter": suggested_quarter, "name": f"Rapport trim {suggested_quarter}"}
        )
    totals = {
        "qty_in": sum((r.qty_in for r in summary), ZERO),
        "amount_in": sum((r.amount_in for r in summary), ZERO),
        "qty_in_opening": sum((r.qty_in_opening for r in summary), ZERO),
        "amount_in_opening": sum((r.amount_in_opening for r in summary), ZERO),
        "qty_in_period": sum((r.qty_in_period for r in summary), ZERO),
        "amount_in_period": sum((r.amount_in_period for r in summary), ZERO),
        "qty_out": sum((r.qty_out for r in summary), ZERO),
        "amount_out": sum((r.amount_out for r in summary), ZERO),
        "remaining_qty": sum((r.remaining_qty for r in summary), ZERO),
        "remaining_amount": sum((r.remaining_amount for r in summary), ZERO),
        "pending_qty": sum((r.pending_qty for r in summary), ZERO),
        "in_service_qty": sum((r.in_service_qty for r in summary), ZERO),
        "provisional_qty": sum((r.provisional_qty for r in summary), ZERO),
    }
    return render(
        request,
        "inventory/recap.html",
        {
            "summary": summary,
            "totals": totals,
            "print_logo_url": request.current_structure.logo.url if request.current_structure and request.current_structure.logo else "",
            "quarterly_reports": quarterly_reports,
            "report_form": report_form,
            "suggested_quarter": suggested_quarter,
            **default_site_context(structure=request.current_structure, fiscal_year=request.current_fiscal_year),
        },
    )


@login_required
def quarterly_report_create_view(request: HttpRequest) -> HttpResponse:
    if request.method != "POST":
        return redirect("recap")

    if not _can_edit(request.user):
        messages.error(request, "Vous n'avez pas les droits pour creer un rapport trimestriel.")
        return redirect("recap")

    if not request.current_structure or not request.current_fiscal_year:
        messages.error(request, "Veuillez selectionner une structure et un exercice.")
        return redirect("select_structure")

    form = QuarterlyReportForm(request.POST)
    if not form.is_valid():
        messages.error(request, "Formulaire de rapport trimestriel invalide.")
        return redirect("recap")

    quarter = int(form.cleaned_data["quarter"])
    name = form.cleaned_data["name"].strip()
    if quarter not in {1, 2, 3}:
        messages.error(request, "Le rapport trimestriel est disponible uniquement pour T1, T2 et T3.")
        return redirect("recap")

    if QuarterlyReport.objects.filter(
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
        quarter=quarter,
    ).exists():
        messages.error(request, f"Un rapport T{quarter} existe deja pour cet exercice.")
        return redirect("recap")

    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    headers = [
        "Compte",
        "Matiere",
        "Existant debut gestion",
        "Entree en cours gestion",
        "Qte sortie",
        "Total",
        "En attente affectation",
        "En service",
        "En sortie provisoire",
        "Total",
        "Montant stock",
        "CMUP",
    ]
    from decimal import Decimal
    ZERO = Decimal("0")
    rows = [
        [
            row.account_code,
            row.name,
            _fmt_pdf_int(row.qty_in_opening),
            _fmt_pdf_int(row.qty_in_period),
            _fmt_pdf_int(row.qty_out),
            _fmt_pdf_int(row.qty_in - row.qty_out),
            _fmt_pdf_int(row.pending_qty),
            _fmt_pdf_int(row.in_service_qty),
            _fmt_pdf_int(row.provisional_qty),
            _fmt_pdf_int(row.remaining_qty),
            _fmt_pdf_amount(row.remaining_amount),
            _fmt_pdf_amount(row.cmup),
        ]
        for row in summary
    ]
    rows.append([
        "",
        "TOTAL",
        _fmt_pdf_int(sum((r.qty_in_opening for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.qty_in_period for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.qty_out for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.qty_in - r.qty_out for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.pending_qty for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.in_service_qty for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.provisional_qty for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.remaining_qty for r in summary), ZERO)),
        _fmt_pdf_amount(sum((r.remaining_amount for r in summary), ZERO)),
        "",
    ])

    snapshot_rows = []
    for row in summary:
        snapshot_rows.append(
            {
                "account_code": row.account_code,
                "name": row.name,
                "unit": row.unit,
                "qty_in_opening": str(row.qty_in_opening),
                "qty_in_period": str(row.qty_in_period),
                "qty_out": str(row.qty_out),
                "remaining_qty": str(row.remaining_qty),
                "pending_qty": str(row.pending_qty),
                "in_service_qty": str(row.in_service_qty),
                "provisional_qty": str(row.provisional_qty),
                "remaining_amount": str(row.remaining_amount),
                "cmup": str(row.cmup),
            }
        )
    snapshot_totals = {
        "qty_in_opening": str(sum((r.qty_in_opening for r in summary), ZERO)),
        "qty_in_period": str(sum((r.qty_in_period for r in summary), ZERO)),
        "qty_out": str(sum((r.qty_out for r in summary), ZERO)),
        "remaining_qty": str(sum((r.remaining_qty for r in summary), ZERO)),
        "pending_qty": str(sum((r.pending_qty for r in summary), ZERO)),
        "in_service_qty": str(sum((r.in_service_qty for r in summary), ZERO)),
        "provisional_qty": str(sum((r.provisional_qty for r in summary), ZERO)),
        "remaining_amount": str(sum((r.remaining_amount for r in summary), ZERO)),
    }

    pdf_response = _build_pdf_response(
        title=f"Rapport trimestriel - Trimestre {quarter}",
        headers=headers,
        rows=rows,
        filename="rapport_trimestriel.pdf",
        signatures=[
            "Ordonnateur des matieres",
            "Comptable des matieres",
        ],
        landscape_mode=True,
        context_lines=_pdf_context_lines(request),
        document_ref="Modele 19",
        subtitle=f"Etat de synthese des mouvements et du stock par compte matiere - T{quarter}",
        logo_path=_pdf_logo_path(request),
    )
    excel_response = _build_excel_response(
        title=f"Rapport trimestriel - Trimestre {quarter}",
        headers=headers,
        rows=[
            [
                row.account_code,
                row.name,
                row.qty_in_opening,
                row.qty_in_period,
                row.qty_out,
                row.qty_in - row.qty_out,
                row.pending_qty,
                row.in_service_qty,
                row.provisional_qty,
                row.remaining_qty,
                row.remaining_amount,
                row.cmup,
            ]
            for row in summary
        ] + [[
            "",
            "TOTAL",
            sum((r.qty_in_opening for r in summary), ZERO),
            sum((r.qty_in_period for r in summary), ZERO),
            sum((r.qty_out for r in summary), ZERO),
            sum((r.qty_in - r.qty_out for r in summary), ZERO),
            sum((r.pending_qty for r in summary), ZERO),
            sum((r.in_service_qty for r in summary), ZERO),
            sum((r.provisional_qty for r in summary), ZERO),
            sum((r.remaining_qty for r in summary), ZERO),
            sum((r.remaining_amount for r in summary), ZERO),
            "",
        ]],
        filename="rapport_trimestriel.xlsx",
        context_lines=_pdf_context_lines(request),
    )

    report = QuarterlyReport.objects.create(
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
        quarter=quarter,
        name=name,
        created_by=request.user,
        snapshot_data={"rows": snapshot_rows, "totals": snapshot_totals},
    )

    pdf_name = f"rapport_trimestriel_T{quarter}_{request.current_structure.code}_{request.current_fiscal_year.year}.pdf"
    excel_name = f"rapport_trimestriel_T{quarter}_{request.current_structure.code}_{request.current_fiscal_year.year}.xlsx"
    report.pdf_file.save(pdf_name, ContentFile(pdf_response.content))
    report.excel_file.save(excel_name, ContentFile(excel_response.content))

    messages.success(request, f"Rapport trimestriel T{quarter} enregistre.")
    return redirect("recap")


@login_required
def quarterly_report_pdf_view(request: HttpRequest, pk: int) -> HttpResponse:
    report = get_object_or_404(
        QuarterlyReport,
        pk=pk,
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    )
    if not report.pdf_file:
        return HttpResponse("Fichier PDF introuvable.", status=404, content_type="text/plain; charset=utf-8")
    return FileResponse(report.pdf_file.open("rb"), as_attachment=True, filename=Path(report.pdf_file.name).name)


@login_required
def quarterly_report_excel_view(request: HttpRequest, pk: int) -> HttpResponse:
    report = get_object_or_404(
        QuarterlyReport,
        pk=pk,
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    )
    if not report.excel_file:
        return HttpResponse("Fichier Excel introuvable.", status=404, content_type="text/plain; charset=utf-8")
    return FileResponse(report.excel_file.open("rb"), as_attachment=True, filename=Path(report.excel_file.name).name)


def _quarterly_decimal(value: str | int | float | None) -> Decimal:
    if value is None:
        return Decimal("0")
    return Decimal(str(value))


def _get_quarterly_consolidation_data(*, quarter: int, fiscal_year: int, structure_type_id: int | None = None) -> dict:
    reports_qs = QuarterlyReport.objects.filter(
        fiscal_year__year=fiscal_year,
        quarter=quarter,
    )
    if structure_type_id is not None:
        reports_qs = reports_qs.filter(structure__structure_type_id=structure_type_id)
    reports = list(reports_qs.select_related("structure", "fiscal_year", "structure__structure_type"))

    aggregate_data = []
    aggregate_totals = {
        "structures_count": 0,
        "comptable_qty": Decimal("0"),
        "pending_qty": Decimal("0"),
        "in_service_qty": Decimal("0"),
        "provisional_qty": Decimal("0"),
        "final_value": Decimal("0"),
    }

    rows_by_account: dict[str, dict] = {}
    for report in reports:
        snapshot = report.snapshot_data or {}
        totals = snapshot.get("totals") or {}
        comptable_qty = _quarterly_decimal(totals.get("remaining_qty"))
        pending_qty = _quarterly_decimal(totals.get("pending_qty"))
        in_service_qty = _quarterly_decimal(totals.get("in_service_qty"))
        provisional_qty = _quarterly_decimal(totals.get("provisional_qty"))
        final_value = _quarterly_decimal(totals.get("remaining_amount"))

        aggregate_data.append(
            {
                "structure": report.structure,
                "fiscal_year": report.fiscal_year,
                "quarterly_report": report,
                "comptable_qty": comptable_qty,
                "pending_qty": pending_qty,
                "in_service_qty": in_service_qty,
                "provisional_qty": provisional_qty,
                "final_value": final_value,
            }
        )

        aggregate_totals["structures_count"] += 1
        aggregate_totals["comptable_qty"] += comptable_qty
        aggregate_totals["pending_qty"] += pending_qty
        aggregate_totals["in_service_qty"] += in_service_qty
        aggregate_totals["provisional_qty"] += provisional_qty
        aggregate_totals["final_value"] += final_value

        for row in snapshot.get("rows", []):
            account_code = row.get("account_code") or ""
            if not account_code:
                continue
            item = rows_by_account.get(account_code)
            if item is None:
                item = {
                    "account_code": account_code,
                    "name": row.get("name") or "",
                    "unit": row.get("unit") or "",
                    "qty_in_opening": Decimal("0"),
                    "qty_in_period": Decimal("0"),
                    "qty_out": Decimal("0"),
                    "remaining_qty": Decimal("0"),
                    "pending_qty": Decimal("0"),
                    "in_service_qty": Decimal("0"),
                    "provisional_qty": Decimal("0"),
                    "remaining_amount": Decimal("0"),
                }
                rows_by_account[account_code] = item
            item["qty_in_opening"] += _quarterly_decimal(row.get("qty_in_opening"))
            item["qty_in_period"] += _quarterly_decimal(row.get("qty_in_period"))
            item["qty_out"] += _quarterly_decimal(row.get("qty_out"))
            item["remaining_qty"] += _quarterly_decimal(row.get("remaining_qty"))
            item["pending_qty"] += _quarterly_decimal(row.get("pending_qty"))
            item["in_service_qty"] += _quarterly_decimal(row.get("in_service_qty"))
            item["provisional_qty"] += _quarterly_decimal(row.get("provisional_qty"))
            item["remaining_amount"] += _quarterly_decimal(row.get("remaining_amount"))

    consolidated_rows = []
    for item in rows_by_account.values():
        remaining_qty = item["remaining_qty"]
        remaining_amount = item["remaining_amount"]
        cmup = remaining_amount / remaining_qty if remaining_qty > Decimal("0") else Decimal("0")
        consolidated_rows.append(
            {
                "account_code": item["account_code"],
                "name": item["name"],
                "unit": item["unit"],
                "qty_in_opening": item["qty_in_opening"],
                "qty_in_period": item["qty_in_period"],
                "qty_out": item["qty_out"],
                "remaining_qty": remaining_qty,
                "pending_qty": item["pending_qty"],
                "in_service_qty": item["in_service_qty"],
                "provisional_qty": item["provisional_qty"],
                "remaining_amount": remaining_amount,
                "cmup": cmup,
            }
        )
    consolidated_rows.sort(key=lambda row: row["account_code"])

    return {
        "aggregate_data": aggregate_data,
        "aggregate_totals": aggregate_totals,
        "consolidated_rows": consolidated_rows,
    }


@login_required
def quarterly_central_recap_view(request: HttpRequest) -> HttpResponse:
    profile = _get_or_create_profile(request.user)
    if profile is None:
        messages.error(request, "Profil utilisateur introuvable.")
        return redirect("dashboard")
    if profile.role != "admin":
        messages.error(request, "Seuls les admins peuvent voir le rapport trimestriel consolide.")
        return redirect("dashboard")

    quarter_param = request.GET.get("quarter") or "1"
    fiscal_year_param = request.GET.get("fiscal_year")
    structure_type_param = (request.GET.get("structure_type") or "").strip()
    try:
        quarter = int(quarter_param)
    except ValueError:
        quarter = 1
    if quarter not in {1, 2, 3}:
        quarter = 1

    all_fiscal_years = list(
        FiscalYear.objects.values_list("year", flat=True).distinct().order_by("-year")
    )
    all_structure_types = list(StructureType.objects.order_by("code"))
    selected_fiscal_year = int(fiscal_year_param) if fiscal_year_param else (all_fiscal_years[0] if all_fiscal_years else None)
    selected_structure_type_id: int | None = None
    if structure_type_param:
        try:
            selected_structure_type_id = int(structure_type_param)
        except ValueError:
            selected_structure_type_id = None
    selected_structure_type = None
    if selected_structure_type_id is not None:
        selected_structure_type = next((t for t in all_structure_types if t.id == selected_structure_type_id), None)
        if selected_structure_type is None:
            selected_structure_type_id = None

    data = {"aggregate_data": [], "aggregate_totals": {}, "consolidated_rows": []}
    if selected_fiscal_year is not None:
        data = _get_quarterly_consolidation_data(
            quarter=quarter,
            fiscal_year=selected_fiscal_year,
            structure_type_id=selected_structure_type_id,
        )

    return render(
        request,
        "inventory/quarterly_central_recap.html",
        {
            "aggregate_data": data["aggregate_data"],
            "aggregate_totals": data["aggregate_totals"],
            "consolidated_rows": data["consolidated_rows"],
            "all_fiscal_years": all_fiscal_years,
            "all_structure_types": all_structure_types,
            "selected_fiscal_year": selected_fiscal_year,
            "selected_quarter": quarter,
            "selected_structure_type_id": selected_structure_type_id,
            "selected_structure_type": selected_structure_type,
        },
    )


@login_required
def quarterly_central_recap_excel_view(request: HttpRequest) -> HttpResponse:
    profile = _get_or_create_profile(request.user)
    if profile is None:
        messages.error(request, "Profil utilisateur introuvable.")
        return redirect("dashboard")
    if profile.role != "admin":
        messages.error(request, "Seuls les admins peuvent exporter le rapport trimestriel consolide.")
        return redirect("dashboard")

    quarter_param = request.GET.get("quarter") or "1"
    fiscal_year_param = request.GET.get("fiscal_year")
    structure_type_param = (request.GET.get("structure_type") or "").strip()
    try:
        quarter = int(quarter_param)
    except ValueError:
        quarter = 1
    if quarter not in {1, 2, 3}:
        quarter = 1
    if not fiscal_year_param:
        return HttpResponse("Exercice non selectionne.", status=400, content_type="text/plain; charset=utf-8")

    structure_type_id: int | None = None
    selected_structure_type = None
    if structure_type_param:
        try:
            structure_type_id = int(structure_type_param)
            selected_structure_type = StructureType.objects.filter(id=structure_type_id).first()
        except ValueError:
            structure_type_id = None

    data = _get_quarterly_consolidation_data(
        quarter=quarter,
        fiscal_year=int(fiscal_year_param),
        structure_type_id=structure_type_id,
    )
    headers = [
        "Compte",
        "Matiere",
        "Existant debut gestion",
        "Entree en cours gestion",
        "Qte sortie",
        "Qté comptable",
        "En attente affectation",
        "En service",
        "En sortie provisoire",
        "Montant restant",
        "CMUP",
    ]
    rows = [
        [
            row["account_code"],
            row["name"],
            row["qty_in_opening"],
            row["qty_in_period"],
            row["qty_out"],
            row["remaining_qty"],
            row["pending_qty"],
            row["in_service_qty"],
            row["provisional_qty"],
            row["remaining_amount"],
            row["cmup"],
        ]
        for row in data["consolidated_rows"]
    ]

    return _build_excel_response(
        title=f"Consolidation trimestrielle T{quarter}",
        headers=headers,
        rows=rows,
        filename=f"consolidation_trimestrielle_T{quarter}.xlsx",
        context_lines=[
            f"Exercice : {fiscal_year_param}",
            f"Type de service : {selected_structure_type.code} - {selected_structure_type.name}" if selected_structure_type else "Type de service : Tous",
        ],
    )


@login_required
def quarterly_central_recap_pdf_view(request: HttpRequest) -> HttpResponse:
    profile = _get_or_create_profile(request.user)
    if profile is None:
        messages.error(request, "Profil utilisateur introuvable.")
        return redirect("dashboard")
    if profile.role != "admin":
        messages.error(request, "Seuls les admins peuvent exporter le rapport trimestriel consolide.")
        return redirect("dashboard")

    quarter_param = request.GET.get("quarter") or "1"
    fiscal_year_param = request.GET.get("fiscal_year")
    structure_type_param = (request.GET.get("structure_type") or "").strip()
    try:
        quarter = int(quarter_param)
    except ValueError:
        quarter = 1
    if quarter not in {1, 2, 3}:
        quarter = 1
    if not fiscal_year_param:
        return HttpResponse("Exercice non selectionne.", status=400, content_type="text/plain; charset=utf-8")

    structure_type_id: int | None = None
    selected_structure_type = None
    if structure_type_param:
        try:
            structure_type_id = int(structure_type_param)
            selected_structure_type = StructureType.objects.filter(id=structure_type_id).first()
        except ValueError:
            structure_type_id = None

    data = _get_quarterly_consolidation_data(
        quarter=quarter,
        fiscal_year=int(fiscal_year_param),
        structure_type_id=structure_type_id,
    )
    headers = [
        "Compte",
        "Matiere",
        "Existant debut gestion",
        "Entree en cours gestion",
        "Qte sortie",
        "Qté comptable",
        "En attente affectation",
        "En service",
        "En sortie provisoire",
        "Montant restant",
        "CMUP",
    ]
    rows = [
        [
            row["account_code"],
            row["name"],
            _fmt_pdf_int(row["qty_in_opening"]),
            _fmt_pdf_int(row["qty_in_period"]),
            _fmt_pdf_int(row["qty_out"]),
            _fmt_pdf_int(row["remaining_qty"]),
            _fmt_pdf_int(row["pending_qty"]),
            _fmt_pdf_int(row["in_service_qty"]),
            _fmt_pdf_int(row["provisional_qty"]),
            _fmt_pdf_amount(row["remaining_amount"]),
            _fmt_pdf_amount(row["cmup"]),
        ]
        for row in data["consolidated_rows"]
    ]

    return _build_pdf_response(
        title=f"Consolidation trimestrielle T{quarter}",
        headers=headers,
        rows=rows,
        filename=f"consolidation_trimestrielle_T{quarter}.pdf",
        signatures=["Ordonnateur des matieres", "Comptable des matieres"],
        context_lines=[
            f"Exercice : {fiscal_year_param}",
            f"Type de service : {selected_structure_type.code} - {selected_structure_type.name}" if selected_structure_type else "Type de service : Tous",
        ],
        document_ref="Modele 19",
        subtitle="Consolidation des releves trimestriels",
        logo_path=_pdf_logo_path(request),
    )


# ---------------------------------------------------------------------------
# Rapport consolide par type de structure
# ---------------------------------------------------------------------------

def _get_consolidation_by_type_data(*, structure_type_id: int, fiscal_year: int) -> dict:
    """Aggregation du releve recapitulatif par type de structure et exercice."""
    from .services import compute_stock_summary

    structures = Structure.objects.filter(structure_type_id=structure_type_id, is_active=True)
    rows_by_account: dict = {}
    aggregate_data = []

    for structure in structures:
        try:
            fy = FiscalYear.objects.get(structure=structure, year=fiscal_year)
        except FiscalYear.DoesNotExist:
            continue
        summary = compute_stock_summary(structure=structure, fiscal_year=fy)
        structure_total = Decimal("0")
        for row in summary:
            code = row.account_code
            if not code:
                continue
            item = rows_by_account.get(code)
            if item is None:
                item = {
                    "account_code": code,
                    "name": row.material_name,
                    "unit": row.unit,
                    "qty_in_opening": Decimal("0"),
                    "qty_in_period": Decimal("0"),
                    "qty_out": Decimal("0"),
                    "remaining_qty": Decimal("0"),
                    "remaining_amount": Decimal("0"),
                }
                rows_by_account[code] = item
            item["qty_in_opening"] += row.qty_in_opening
            item["qty_in_period"] += row.qty_in_period
            item["qty_out"] += row.qty_out
            item["remaining_qty"] += row.remaining_qty
            item["remaining_amount"] += row.remaining_amount
            structure_total += row.remaining_amount
        aggregate_data.append({"structure": structure, "total": structure_total})

    consolidated = []
    grand_total = Decimal("0")
    for item in sorted(rows_by_account.values(), key=lambda r: r["account_code"]):
        remaining_qty = item["remaining_qty"]
        remaining_amount = item["remaining_amount"]
        cmup = remaining_amount / remaining_qty if remaining_qty > Decimal("0") else Decimal("0")
        consolidated.append({**item, "cmup": cmup})
        grand_total += remaining_amount

    return {
        "aggregate_data": aggregate_data,
        "consolidated_rows": consolidated,
        "grand_total": grand_total,
    }


@login_required
def consolidated_by_type_view(request: HttpRequest) -> HttpResponse:
    profile = _get_or_create_profile(request.user)
    if profile is None or profile.role != "admin":
        messages.error(request, "Acces reserve aux administrateurs.")
        return redirect("dashboard")

    all_types = StructureType.objects.all()
    all_fiscal_years = list(
        FiscalYear.objects.values_list("year", flat=True).distinct().order_by("-year")
    )

    type_param = request.GET.get("structure_type")
    fiscal_year_param = request.GET.get("fiscal_year")
    selected_type = None
    data = {"aggregate_data": [], "consolidated_rows": [], "grand_total": Decimal("0")}

    if type_param:
        try:
            selected_type = StructureType.objects.get(pk=int(type_param))
        except (StructureType.DoesNotExist, ValueError):
            selected_type = None

    selected_fiscal_year = int(fiscal_year_param) if fiscal_year_param else (all_fiscal_years[0] if all_fiscal_years else None)

    if selected_type and selected_fiscal_year:
        data = _get_consolidation_by_type_data(
            structure_type_id=selected_type.pk,
            fiscal_year=selected_fiscal_year,
        )

    return render(
        request,
        "inventory/consolidated_by_type.html",
        {
            "all_types": all_types,
            "all_fiscal_years": all_fiscal_years,
            "selected_type": selected_type,
            "selected_fiscal_year": selected_fiscal_year,
            "aggregate_data": data["aggregate_data"],
            "consolidated_rows": data["consolidated_rows"],
            "grand_total": data["grand_total"],
        },
    )


@login_required
def recap_pdf_view(request: HttpRequest) -> HttpResponse:
    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    headers = [
        "Compte",
        "Matiere",
        "Existant debut gestion",
        "Entree en cours gestion",
        "Qte sortie",
        "Total",
        "En attente affectation",
        "En service",
        "En sortie provisoire",
        "Total",
        "Montant stock",
        "CMUP",
    ]
    from decimal import Decimal
    ZERO = Decimal("0")
    rows = [
        [
            row.account_code,
            row.name,
            _fmt_pdf_int(row.qty_in_opening),
            _fmt_pdf_int(row.qty_in_period),
            _fmt_pdf_int(row.qty_out),
            _fmt_pdf_int(row.qty_in - row.qty_out),
            _fmt_pdf_int(row.pending_qty),
            _fmt_pdf_int(row.in_service_qty),
            _fmt_pdf_int(row.provisional_qty),
            _fmt_pdf_int(row.remaining_qty),
            _fmt_pdf_amount(row.remaining_amount),
            _fmt_pdf_amount(row.cmup),
        ]
        for row in summary
    ]
    rows.append([
        "",
        "TOTAL",
        _fmt_pdf_int(sum((r.qty_in_opening for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.qty_in_period for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.qty_out for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.qty_in - r.qty_out for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.pending_qty for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.in_service_qty for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.provisional_qty for r in summary), ZERO)),
        _fmt_pdf_int(sum((r.remaining_qty for r in summary), ZERO)),
        _fmt_pdf_amount(sum((r.remaining_amount for r in summary), ZERO)),
        "",
    ])
    return _build_pdf_response(
        title="Releve recapitulatif",
        headers=headers,
        rows=rows,
        filename="releve_recapitulatif.pdf",
        signatures=[
            "Ordonnateur des matieres",
            "Comptable des matieres",
        ],
        landscape_mode=True,
        context_lines=_pdf_context_lines(request),
        document_ref="Modele 19",
        subtitle="Etat de synthese des mouvements et du stock par compte matiere",
        logo_path=_pdf_logo_path(request),
    )


@login_required
def recap_excel_view(request: HttpRequest) -> HttpResponse:
    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    headers = [
        "Compte",
        "Matiere",
        "Existant debut gestion",
        "Entree en cours gestion",
        "Qte sortie",
        "Total",
        "En attente affectation",
        "En service",
        "En sortie provisoire",
        "Total",
        "Montant stock",
        "CMUP",
    ]
    from decimal import Decimal
    ZERO = Decimal("0")
    rows = [
        [
            row.account_code,
            row.name,
            row.qty_in_opening,
            row.qty_in_period,
            row.qty_out,
            row.qty_in - row.qty_out,
            row.pending_qty,
            row.in_service_qty,
            row.provisional_qty,
            row.remaining_qty,
            row.remaining_amount,
            row.cmup,
        ]
        for row in summary
    ]
    rows.append([
        "",
        "TOTAL",
        sum((r.qty_in_opening for r in summary), ZERO),
        sum((r.qty_in_period for r in summary), ZERO),
        sum((r.qty_out for r in summary), ZERO),
        sum((r.qty_in - r.qty_out for r in summary), ZERO),
        sum((r.pending_qty for r in summary), ZERO),
        sum((r.in_service_qty for r in summary), ZERO),
        sum((r.provisional_qty for r in summary), ZERO),
        sum((r.remaining_qty for r in summary), ZERO),
        sum((r.remaining_amount for r in summary), ZERO),
        "",
    ])
    return _build_excel_response(
        title="Releve recapitulatif",
        headers=headers,
        rows=rows,
        filename="releve_recapitulatif.xlsx",
        context_lines=_pdf_context_lines(request),
    )


@login_required
def pv_recensement_view(request: HttpRequest) -> HttpResponse:
    if not request.current_structure or not request.current_fiscal_year:
        messages.error(request, "Veuillez selectionner une structure et un exercice avant le PV de recensement.")
        return redirect("select_structure")

    if request.method == "POST" and _can_edit(request.user):
        if not _ensure_open_fiscal_year(request, redirect_to="pv_recensement"):
            return redirect("pv_recensement")
        def _apply_pv_display(row, inventory: PhysicalInventory) -> None:
            if not inventory.is_pv_filled:
                return
            row.pv_is_filled = True
            row.inventory_total_qty = (
                inventory.pending_qty
                + inventory.in_service_qty
                + inventory.provisional_qty
            )
            row.inventory_gap = row.inventory_total_qty - row.remaining_qty
            row.gap_plus = row.inventory_gap if row.inventory_gap > Decimal("0") else Decimal("0")
            row.gap_minus = -row.inventory_gap if row.inventory_gap < Decimal("0") else Decimal("0")
            row.inventory_gap_value = row.inventory_gap * row.cmup
        # Valider toutes les lignes avant sauvegarde pour eviter une perte apparente des saisies.
        materials_all = Material.objects.filter(
            structure=request.current_structure,
            fiscal_year=request.current_fiscal_year,
        ).order_by("account_code")
        forms_data: list[tuple[Material, PhysicalInventory, PhysicalInventoryForm]] = []
        has_errors = False
        for material in materials_all:
            inventory, _ = PhysicalInventory.objects.get_or_create(material=material)
            prefix = f"mat_{material.id}"
            form = PhysicalInventoryForm(request.POST, instance=inventory, prefix=prefix)
            forms_data.append((material, inventory, form))
            if not form.is_valid():
                has_errors = True

        if has_errors:
            messages.error(request, "Certaines lignes sont invalides. Corrigez puis enregistrez a nouveau.")
            summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
            summary_by_code = {row.account_code: row for row in summary}
            inventory_forms = []
            for material, _, form in forms_data:
                row = summary_by_code.get(material.account_code)
                if row is not None:
                    inventory = form.instance
                    if inventory is not None:
                        _apply_pv_display(row, inventory)
                    inventory_forms.append((row, form))
            return render(
                request,
                "inventory/pv_recensement.html",
                {
                    "inventory_forms": inventory_forms,
                    "can_edit": _can_edit(request.user),
                    "can_write": _can_edit(request.user) and not bool(request.current_fiscal_year and request.current_fiscal_year.is_closed),
                    "print_logo_url": request.current_structure.logo.url if request.current_structure and request.current_structure.logo else "",
                    **default_site_context(structure=request.current_structure, fiscal_year=request.current_fiscal_year),
                },
            )

        for _, inventory, form in forms_data:
            inventory.pending_qty = form.cleaned_data["pending_qty"]
            inventory.in_service_qty = form.cleaned_data["in_service_qty"]
            inventory.provisional_qty = form.cleaned_data["provisional_qty"]
            inventory.is_pv_filled = True
            inventory.save()
        _audit(
            request,
            action=AuditLog.ACTION_UPDATE,
            entity="PhysicalInventory",
            description="Enregistrement PV de recensement",
            metadata={"rows": len(forms_data)},
        )
        messages.success(request, "Proces verbal de recensement enregistre.")
        return redirect("pv_recensement")

    # GET : construire les formulaires avec les valeurs sauvegardées
    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    materials_by_code = {
        material.account_code: material
        for material in Material.objects.filter(
            structure=request.current_structure,
            fiscal_year=request.current_fiscal_year,
        )
    }
    inventory_forms = []
    for row in summary:
        material = materials_by_code.get(row.account_code)
        if material is None:
            continue
        inventory, _ = PhysicalInventory.objects.get_or_create(material=material)
        form_instance = inventory if inventory.is_pv_filled else None
        form = PhysicalInventoryForm(instance=form_instance, prefix=f"mat_{material.id}")
        if inventory.is_pv_filled:
            row.pv_is_filled = True
            row.inventory_total_qty = (
                inventory.pending_qty
                + inventory.in_service_qty
                + inventory.provisional_qty
            )
            row.inventory_gap = row.inventory_total_qty - row.remaining_qty
            row.gap_plus = row.inventory_gap if row.inventory_gap > Decimal("0") else Decimal("0")
            row.gap_minus = -row.inventory_gap if row.inventory_gap < Decimal("0") else Decimal("0")
            row.inventory_gap_value = row.inventory_gap * row.cmup
        inventory_forms.append((row, form))
    return render(
        request,
        "inventory/pv_recensement.html",
        {
            "inventory_forms": inventory_forms,
            "can_edit": _can_edit(request.user),
            "can_write": _can_edit(request.user) and not bool(request.current_fiscal_year and request.current_fiscal_year.is_closed),
            "print_logo_url": request.current_structure.logo.url if request.current_structure and request.current_structure.logo else "",
            **default_site_context(structure=request.current_structure, fiscal_year=request.current_fiscal_year),
        },
    )


@login_required
def pv_recensement_pdf_view(request: HttpRequest) -> HttpResponse:
    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    headers = [
        "Compte",
        "Matiere",
        "Qte restante",
        "Attente affectation",
        "En service",
        "Sortie provisoire",
        "Total PV",
        "Ecart +",
        "Ecart -",
        "CMUP",
        "Valeur difference",
    ]
    rows = [
        [
            row.account_code,
            row.name,
            _fmt_pdf_int(row.remaining_qty),
            _fmt_pdf_int(row.pending_qty if row.pv_is_filled else Decimal("0")),
            _fmt_pdf_int(row.in_service_qty if row.pv_is_filled else Decimal("0")),
            _fmt_pdf_int(row.provisional_qty if row.pv_is_filled else Decimal("0")),
            _fmt_pdf_int(row.inventory_total_qty if row.pv_is_filled else Decimal("0")),
            _fmt_pdf_int(row.gap_plus if row.pv_is_filled else Decimal("0")),
            _fmt_pdf_int(row.gap_minus if row.pv_is_filled else Decimal("0")),
            _fmt_pdf_amount(row.cmup),
            _fmt_pdf_amount(row.inventory_gap_value if row.pv_is_filled else Decimal("0")),
        ]
        for row in summary
    ]
    return _build_pdf_response(
        title="PV de recensement",
        headers=headers,
        rows=rows,
        filename="pv_recensement.pdf",
        signatures=["Nom, qualite et signature des membres de la commission"],
        landscape_mode=True,
        context_lines=_pdf_context_lines(request),
        document_ref="Modele 20",
        subtitle="Proces-verbal de recensement contradictoire des matieres",
        logo_path=_pdf_logo_path(request),
    )


@login_required
def pv_recensement_excel_view(request: HttpRequest) -> HttpResponse:
    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    headers = [
        "Compte",
        "Matiere",
        "Qte restante",
        "Attente affectation",
        "En service",
        "Sortie provisoire",
        "Total PV",
        "Ecart +",
        "Ecart -",
        "CMUP",
        "Valeur difference",
    ]
    rows = [
        [
            row.account_code,
            row.name,
            row.remaining_qty,
            row.pending_qty if row.pv_is_filled else Decimal("0"),
            row.in_service_qty if row.pv_is_filled else Decimal("0"),
            row.provisional_qty if row.pv_is_filled else Decimal("0"),
            row.inventory_total_qty if row.pv_is_filled else Decimal("0"),
            row.gap_plus if row.pv_is_filled else Decimal("0"),
            row.gap_minus if row.pv_is_filled else Decimal("0"),
            row.cmup,
            row.inventory_gap_value if row.pv_is_filled else Decimal("0"),
        ]
        for row in summary
    ]
    return _build_excel_response(
        title="PV de recensement",
        headers=headers,
        rows=rows,
        filename="pv_recensement.xlsx",
        context_lines=_pdf_context_lines(request),
    )


@login_required
def final_report_view(request: HttpRequest) -> HttpResponse:
    if not request.current_structure or not request.current_fiscal_year:
        messages.error(request, "Veuillez selectionner une structure et un exercice.")
        return redirect("select_structure")

    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)

    q = (request.GET.get("q") or "").strip().lower()
    pv_filter = (request.GET.get("pv") or "all").strip().lower()
    stock_filter = (request.GET.get("stock") or "all").strip().lower()
    pv_by_account: dict[str, tuple[Decimal, Decimal, Decimal]] = {}
    for material in Material.objects.filter(
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    ).select_related("physical_inventory"):
        inventory = getattr(material, "physical_inventory", None)
        if inventory and inventory.is_pv_filled:
            pv_by_account[material.account_code] = (
                inventory.pending_qty,
                inventory.in_service_qty,
                inventory.provisional_qty,
            )

    model_20_url = ""
    documents_path = Path(settings.MEDIA_ROOT) / "documents"
    if documents_path.exists():
        for file_path in documents_path.glob("*.docx"):
            name_lower = file_path.name.lower()
            if "modele 20" in name_lower:
                model_20_url = f"/media/documents/{file_path.name}"

    regulation_entries = []
    balance_rows = []
    plus_lines_count = 0
    minus_lines_count = 0
    plus_total_qty = Decimal("0")
    minus_total_qty = Decimal("0")
    plus_total_amount = Decimal("0")
    minus_total_amount = Decimal("0")
    total_qty_in = Decimal("0")
    total_qty_in_opening = Decimal("0")
    total_qty_in_period = Decimal("0")
    total_qty_out = Decimal("0")
    total_comptable_qty = Decimal("0")
    total_pending_qty = Decimal("0")
    total_in_service_qty = Decimal("0")
    total_provisional_qty = Decimal("0")
    total_pv_qty = Decimal("0")
    total_final_amount_display = Decimal("0")

    for row in summary:
        pv_pending, pv_in_service, pv_provisional = pv_by_account.get(
            row.account_code,
            (Decimal("0"), Decimal("0"), Decimal("0")),
        )
        pv_is_filled = row.account_code in pv_by_account
        row_total_pv_qty = pv_pending + pv_in_service + pv_provisional
        gap = row_total_pv_qty - row.remaining_qty if pv_is_filled else Decimal("0")
        regulation_qty = gap if pv_is_filled else Decimal("0")
        regulation_amount = regulation_qty * row.cmup if pv_is_filled else Decimal("0")
        final_qty = row.remaining_qty + regulation_qty
        final_amount = row.remaining_amount + regulation_amount

        if q and (q not in row.account_code.lower()) and (q not in row.name.lower()):
            continue
        if pv_filter == "filled" and not pv_is_filled:
            continue
        if pv_filter == "missing" and pv_is_filled:
            continue
        if stock_filter == "negative" and final_qty >= Decimal("0"):
            continue
        if stock_filter == "zero" and final_qty != Decimal("0"):
            continue
        if stock_filter == "positive" and final_qty <= Decimal("0"):
            continue

        total_qty_in += row.qty_in
        total_qty_in_opening += row.qty_in_opening
        total_qty_in_period += row.qty_in_period
        total_qty_out += row.qty_out
        total_comptable_qty += row.remaining_qty
        total_pending_qty += pv_pending if pv_is_filled else Decimal("0")
        total_in_service_qty += pv_in_service if pv_is_filled else Decimal("0")
        total_provisional_qty += pv_provisional if pv_is_filled else Decimal("0")
        total_pv_qty += row_total_pv_qty
        total_final_amount_display += (final_qty * row.cmup if row.cmup > Decimal("0") else Decimal("0"))

        if pv_is_filled and (regulation_qty != Decimal("0") or regulation_amount != Decimal("0")):
            gap_plus = regulation_qty if regulation_qty > Decimal("0") else Decimal("0")
            gap_minus = -regulation_qty if regulation_qty < Decimal("0") else Decimal("0")
            regulation_entries.append(
                {
                    "account_code": row.account_code,
                    "name": row.name,
                    "regulation_qty": regulation_qty,
                    "regulation_amount": regulation_amount,
                    "gap_plus": gap_plus,
                    "gap_minus": gap_minus,
                    "cmup": row.cmup,
                }
            )
            if gap_plus > Decimal("0"):
                plus_lines_count += 1
                plus_total_qty += gap_plus
                plus_total_amount += gap_plus * row.cmup
            if gap_minus > Decimal("0"):
                minus_lines_count += 1
                minus_total_qty += gap_minus
                minus_total_amount += gap_minus * row.cmup

        balance_rows.append(
            {
                "account_code": row.account_code,
                "name": row.name,
                "unit": row.unit,
                "qty_in": row.qty_in,
                "qty_in_opening": row.qty_in_opening,
                "qty_in_period": row.qty_in_period,
                "qty_out": row.qty_out,
                "comptable_qty": row.remaining_qty,
                "comptable_amount": row.remaining_amount,
                "pending_qty": pv_pending if pv_is_filled else Decimal("0"),
                "in_service_qty": pv_in_service if pv_is_filled else Decimal("0"),
                "provisional_qty": pv_provisional if pv_is_filled else Decimal("0"),
                "total_pv_qty": row_total_pv_qty,
                "gap": gap,
                "cmup": row.cmup,
                "physique_qty": row_total_pv_qty if pv_is_filled else Decimal("0"),
                "regulation_qty": regulation_qty,
                "regulation_amount": regulation_amount,
                "final_qty": final_qty,
                "final_amount": final_amount,
                "remaining_amount_final": final_qty * row.cmup if row.cmup > Decimal("0") else Decimal("0"),
                "pv_is_filled": pv_is_filled,
            }
        )

    context = {
        "regulation_entries": regulation_entries,
        "balance_rows": balance_rows,
        "model_20_url": model_20_url,
        "plus_lines_count": plus_lines_count,
        "minus_lines_count": minus_lines_count,
        "plus_total_qty": plus_total_qty,
        "minus_total_qty": minus_total_qty,
        "plus_total_amount": plus_total_amount,
        "minus_total_amount": minus_total_amount,
        "filters": {
            "q": q,
            "pv": pv_filter,
            "stock": stock_filter,
        },
        "filtered_rows_count": len(balance_rows),
        "totals": {
            "comptable_qty": total_comptable_qty,
            "final_amount_display": total_final_amount_display,
            "qty_in": total_qty_in,
            "qty_in_opening": total_qty_in_opening,
            "qty_in_period": total_qty_in_period,
            "qty_out": total_qty_out,
            "pending_qty": total_pending_qty,
            "in_service_qty": total_in_service_qty,
            "provisional_qty": total_provisional_qty,
            "total_pv_qty": total_pv_qty,
        },
        "print_logo_url": request.current_structure.logo.url if request.current_structure and request.current_structure.logo else "",
        "can_close_fiscal_year": _can_close_fiscal_year(request.user),
        **default_site_context(structure=request.current_structure, fiscal_year=request.current_fiscal_year),
    }
    return render(request, "inventory/final_report.html", context)


@login_required
def final_balance_pdf_view(request: HttpRequest) -> HttpResponse:
    if not request.current_structure or not request.current_fiscal_year:
        messages.error(request, "Veuillez selectionner une structure et un exercice.")
        return redirect("select_structure")

    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    pv_totals_by_account: dict[str, Decimal] = {}
    for material in Material.objects.filter(
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    ).select_related("physical_inventory"):
        inventory = getattr(material, "physical_inventory", None)
        if inventory and inventory.is_pv_filled:
            pv_totals_by_account[material.account_code] = (
                inventory.pending_qty + inventory.in_service_qty + inventory.provisional_qty
            )

    headers = [
        "Compte",
        "Matiere",
        "Existant debut gestion",
        "Entree en cours gestion",
        "Qte sortie",
        "Qte comptable",
        "Total PV",
        "CMUP",
        "Valeur stock",
    ]
    rows = []
    for row in summary:
        pv_total = pv_totals_by_account.get(row.account_code, Decimal("0"))
        pv_is_filled = row.account_code in pv_totals_by_account
        regulation_qty = (pv_total - row.remaining_qty) if pv_is_filled else Decimal("0")
        final_qty = row.remaining_qty + regulation_qty
        final_value = final_qty * row.cmup if row.cmup > Decimal("0") else Decimal("0")
        rows.append(
            [
                row.account_code,
                row.name,
                _fmt_pdf_int(row.qty_in_opening),
                _fmt_pdf_int(row.qty_in_period),
                _fmt_pdf_int(row.qty_out),
                _fmt_pdf_int(row.remaining_qty),
                _fmt_pdf_int(pv_total),
                _fmt_pdf_amount(row.cmup),
                _fmt_pdf_amount(final_value),
            ]
        )

    return _build_pdf_response(
        title="Balance generale des comptes",
        headers=headers,
        rows=rows,
        filename="balance_generale_comptes.pdf",
        signatures=["Ordonnateur des matieres", "Comptable des matieres"],
        landscape_mode=True,
        context_lines=_pdf_context_lines(request),
        document_ref="Modele 22",
        subtitle="Balance finale des comptes matieres par article",
        logo_path=_pdf_logo_path(request),
    )


@login_required
def final_balance_excel_view(request: HttpRequest) -> HttpResponse:
    if not request.current_structure or not request.current_fiscal_year:
        messages.error(request, "Veuillez selectionner une structure et un exercice.")
        return redirect("select_structure")

    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    pv_totals_by_account: dict[str, Decimal] = {}
    for material in Material.objects.filter(
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    ).select_related("physical_inventory"):
        inventory = getattr(material, "physical_inventory", None)
        if inventory and inventory.is_pv_filled:
            pv_totals_by_account[material.account_code] = (
                inventory.pending_qty + inventory.in_service_qty + inventory.provisional_qty
            )

    headers = [
        "Compte",
        "Matiere",
        "Existant debut gestion",
        "Entree en cours gestion",
        "Qte sortie",
        "Qte comptable",
        "Total PV",
        "CMUP",
        "Valeur stock",
    ]
    rows = []
    for row in summary:
        pv_total = pv_totals_by_account.get(row.account_code, Decimal("0"))
        pv_is_filled = row.account_code in pv_totals_by_account
        regulation_qty = (pv_total - row.remaining_qty) if pv_is_filled else Decimal("0")
        final_qty = row.remaining_qty + regulation_qty
        final_value = final_qty * row.cmup if row.cmup > Decimal("0") else Decimal("0")
        rows.append(
            [
                row.account_code,
                row.name,
                row.qty_in_opening,
                row.qty_in_period,
                row.qty_out,
                row.remaining_qty,
                pv_total,
                row.cmup,
                final_value,
            ]
        )

    if rows:
        rows.append(["" for _ in headers])
        rows.append([
            "Ordonnateur des matieres",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "Comptable des matieres",
        ])

    return _build_excel_response(
        title="Balance generale des comptes",
        headers=headers,
        rows=rows,
        filename="balance_generale_comptes.xlsx",
        context_lines=_pdf_context_lines(request),
    )


def _pdf_context_lines(request: HttpRequest) -> list[str]:
    site = default_site_context(structure=getattr(request, "current_structure", None), fiscal_year=getattr(request, "current_fiscal_year", None))
    structure = getattr(request, "current_structure", None)
    fiscal_year = getattr(request, "current_fiscal_year", None)
    structure_line = f"Service : {structure.code} - {structure.name}" if structure else "Service : -"
    year_label = f"Exercice : {fiscal_year.year}" if fiscal_year else "Exercice : -"
    generated_label = f"Date d'edition : {date.today().strftime('%d/%m/%Y')}"
    lines = [
        site.get("country", ""),
        "Un Peuple - Un But - Une Foi",
        site.get("ministry", ""),
        structure_line,
    ]
    if site.get("address", "").strip():
        lines.append(f"Adresse : {site['address']}")
    if site.get("phone", "").strip():
        lines.append(f"Telephone : {site['phone']}")
    if site.get("email", "").strip():
        lines.append(f"Email : {site['email']}")
    lines.extend([year_label, generated_label])
    return lines


def _pdf_logo_path(request: HttpRequest) -> str | None:
    structure = getattr(request, "current_structure", None)
    if not structure or not getattr(structure, "logo", None):
        return None
    try:
        return structure.logo.path
    except Exception:  # noqa: BLE001
        return None


def _build_pdf_response(
    *,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    filename: str,
    signatures: list[str] | None = None,
    landscape_mode: bool = False,
    context_lines: list[str] | None = None,
    document_ref: str = "",
    subtitle: str = "",
    logo_path: str | None = None,
    extra_paragraphs: list[str] | None = None,
) -> HttpResponse:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
        from reportlab.platypus import Image as RLImage
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except Exception as exc:  # noqa: BLE001
        return HttpResponse(
            f"Generation PDF indisponible: {exc}. Installez reportlab.",
            status=500,
            content_type="text/plain; charset=utf-8",
        )

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    page_size = landscape(A4) if landscape_mode else A4
    doc = SimpleDocTemplate(response, pagesize=page_size, leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=20)
    styles = getSampleStyleSheet()
    header_style = ParagraphStyle(
        "header_line",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=11,
        alignment=1,
    )
    title_style = ParagraphStyle(
        "report_title",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        alignment=1,
    )
    subtitle_style = ParagraphStyle(
        "report_subtitle",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9,
        leading=11,
        alignment=1,
    )
    right_meta_style = ParagraphStyle(
        "right_meta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        alignment=2,
    )
    signature_style = ParagraphStyle(
        "signature_style",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        alignment=1,
    )
    table_header_style = ParagraphStyle(
        "table_header",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=9,
        alignment=1,
    )

    def _wrap_header(label: str) -> str:
        if " " not in label:
            return label
        left, right = label.rsplit(" ", 1)
        return f"{left}<br/>{right}"

    header_cells = [Paragraph(_wrap_header(text), table_header_style) for text in headers]
    data = [header_cells] if headers else []
    if rows:
        data.extend(rows)
    else:
        data.append(["Aucune donnee"] + ([""] * (len(headers) - 1) if headers else []))

    available_width = doc.width
    col_count = len(headers) if headers else (len(data[0]) if data else 1)
    if col_count >= 2:
        weights = [0.8, 2.0] + [1.0] * (col_count - 2)
    else:
        weights = [1.0] * col_count
    total_weight = sum(weights) if weights else 1.0
    col_widths = [available_width * weight / total_weight for weight in weights]
    table = Table(data, repeatRows=1, colWidths=col_widths)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0ece4")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#b8aa95")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#faf8f3")]),
            ]
        )
    )

    story: list = []
    left_lines = [line for line in (context_lines or [])[:5] if line]
    right_lines = []
    if document_ref:
        right_lines.append(f"Reference : {document_ref}")
    if len(context_lines or []) > 5:
        right_lines.extend((context_lines or [])[5:])
    else:
        right_lines.append(f"Date d'edition : {date.today().strftime('%d/%m/%Y')}")

    header_left = "<br/>".join(left_lines) if left_lines else ""
    header_right = "<br/>".join(right_lines) if right_lines else ""
    left_cell = Paragraph(header_left, header_style)
    if logo_path:
        try:
            logo_img = RLImage(logo_path)
            logo_img._restrictSize(26 * mm, 26 * mm)
            left_cell = Table(
                [[logo_img, Paragraph(header_left, header_style)]],
                colWidths=[30 * mm, available_width * 0.62 - 30 * mm],
            )
            left_cell.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("ALIGN", (0, 0), (0, 0), "LEFT"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 0),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                    ]
                )
            )
        except Exception:  # noqa: BLE001
            left_cell = Paragraph(header_left, header_style)

    header_table = Table(
        [[left_cell, Paragraph(header_right, right_meta_style)]],
        colWidths=[available_width * 0.62, available_width * 0.38],
    )
    header_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (0, 0), (0, 0), "LEFT"),
                ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#9c8b73")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )

    story.extend([header_table, Spacer(1, 8), Paragraph(title, title_style)])
    if subtitle:
        story.extend([Spacer(1, 2), Paragraph(subtitle, subtitle_style)])
    story.extend([Spacer(1, 10), table])

    if extra_paragraphs:
        body_style = ParagraphStyle(
            "body_paragraph",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            alignment=0,
        )
        for paragraph in extra_paragraphs:
            if not (paragraph or "").strip():
                continue
            story.extend([Spacer(1, 8), Paragraph(paragraph.strip(), body_style)])

    if signatures:
        story.append(Spacer(1, 16))
        sig_table = Table([signatures], colWidths=[available_width / max(len(signatures), 1)] * max(len(signatures), 1))
        sig_table.setStyle(
            TableStyle(
                [
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 32),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.8, colors.black),
                ]
            )
        )
        story.append(sig_table)
        story.append(Spacer(1, 8))
        story.append(Paragraph("Date, cachet et signatures", signature_style))

    def _draw_page_footer(pdf_canvas: canvas.Canvas, _doc: SimpleDocTemplate) -> None:
        pdf_canvas.saveState()
        pdf_canvas.setFont("Helvetica", 8)
        footer_text = f"Page {_doc.page}"
        pdf_canvas.setFillColor(colors.HexColor("#555555"))
        pdf_canvas.drawRightString(page_size[0] - 20, 10 * mm, footer_text)
        pdf_canvas.restoreState()

    doc.build(story, onFirstPage=_draw_page_footer, onLaterPages=_draw_page_footer)
    return response


def _build_excel_response(
    *,
    title: str,
    headers: list[str],
    rows: list[list],
    filename: str,
    context_lines: list[str] | None = None,
) -> HttpResponse:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception as exc:  # noqa: BLE001
        return HttpResponse(
            f"Generation Excel indisponible: {exc}. Installez openpyxl.",
            status=500,
            content_type="text/plain; charset=utf-8",
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Export"

    current_row = 1
    context = [line for line in (context_lines or []) if line]
    if context:
        for line in context:
            sheet.cell(row=current_row, column=1, value=line)
            current_row += 1
        current_row += 1

    sheet.cell(row=current_row, column=1, value=title)
    sheet.cell(row=current_row, column=1).font = Font(bold=True, size=14)
    current_row += 2

    header_fill = PatternFill(start_color="F0ECE4", end_color="F0ECE4", fill_type="solid")
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=current_row, column=col_index, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    current_row += 1

    for row in rows:
        for col_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=current_row, column=col_index, value=value)
            if isinstance(value, Decimal):
                cell.number_format = "#,##0.00"
        current_row += 1

    for col_index in range(1, len(headers) + 1):
        max_length = 10
        for cell in sheet.iter_cols(min_col=col_index, max_col=col_index, min_row=1, max_row=current_row):
            for col_cell in cell:
                if col_cell.value is not None:
                    max_length = max(max_length, len(str(col_cell.value)))
        sheet.column_dimensions[sheet.cell(row=1, column=col_index).column_letter].width = min(max_length + 2, 45)

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


@login_required
def create_grouped_regulation_voucher_view(request: HttpRequest, direction: str) -> HttpResponse:
    if request.method != "POST":
        return redirect("final_report")

    if not _can_edit(request.user):
        messages.error(request, "Vous n'avez pas les droits pour creer des bons de regularisation.")
        return redirect("final_report")

    if not _ensure_open_fiscal_year(request, redirect_to="final_report"):
        return redirect("final_report")

    if not request.current_structure or not request.current_fiscal_year:
        messages.error(request, "Veuillez selectionner une structure et un exercice.")
        return redirect("select_structure")

    if direction not in {"plus", "minus"}:
        messages.error(request, "Type de regularisation invalide.")
        return redirect("final_report")

    summary = compute_stock_summary(structure=request.current_structure, fiscal_year=request.current_fiscal_year)
    materials_map = {
        m.account_code: m
        for m in Material.objects.filter(
            structure=request.current_structure,
            fiscal_year=request.current_fiscal_year,
        )
    }
    pv_totals_by_account: dict[str, Decimal] = {}
    for material in Material.objects.filter(
        structure=request.current_structure,
        fiscal_year=request.current_fiscal_year,
    ).select_related("physical_inventory"):
        inventory = getattr(material, "physical_inventory", None)
        if inventory and inventory.is_pv_filled:
            pv_totals_by_account[material.account_code] = (
                inventory.pending_qty + inventory.in_service_qty + inventory.provisional_qty
            )

    lines: list[VoucherLineInput] = []

    if direction == "plus":
        voucher_type = Voucher.TYPE_ENTRY
        reason = "Regularisation globale des ecarts en plus (suite certificat administratif modele 20)"
        for row in summary:
            pv_total = pv_totals_by_account.get(row.account_code)
            if pv_total is None:
                continue
            gap = pv_total - row.remaining_qty
            if gap <= Decimal("0"):
                continue
            material = materials_map.get(row.account_code)
            if material is None:
                continue
            lines.append(
                VoucherLineInput(
                    material_id=material.id,
                    specification="Regularisation ecart en plus",
                    quantity=gap,
                    unit_price=row.cmup,
                )
            )
    else:
        voucher_type = Voucher.TYPE_FINAL_EXIT
        reason = "Regularisation globale des ecarts en moins (suite certificat administratif modele 20)"
        for row in summary:
            pv_total = pv_totals_by_account.get(row.account_code)
            if pv_total is None:
                continue
            gap = pv_total - row.remaining_qty
            if gap >= Decimal("0"):
                continue
            material = materials_map.get(row.account_code)
            if material is None:
                continue
            lines.append(
                VoucherLineInput(
                    material_id=material.id,
                    specification="Regularisation ecart en moins",
                    quantity=-gap,
                    unit_price=row.cmup,
                )
            )

    if not lines:
        messages.error(request, "Aucune ligne a regulariser pour cette operation.")
        return redirect("final_report")

    created_vouchers: list[Voucher] = []
    try:
        vouchers = create_voucher(
            user=request.user,
            voucher_type=voucher_type,
            operation_date=date.today(),
            source_or_destination="Regularisation",
            comments=reason,
            lines=lines,
            structure=request.current_structure,
            fiscal_year=request.current_fiscal_year,
            apply_to_physical_inventory=False,
        )
        created_vouchers.extend(vouchers)
    except (ValueError, ValidationError) as error:
        messages.error(request, str(error))
        return redirect("final_report")

    total_lines = len(lines)
    nums = ", ".join(f"N° {v.number}" for v in created_vouchers)
    messages.success(
        request,
        f"{len(created_vouchers)} bon(s) de regularisation cree(s) : {nums} ({total_lines} ligne(s) au total).",
    )
    for v in created_vouchers:
        _audit(
            request,
            action=AuditLog.ACTION_CREATE,
            entity="Voucher",
            entity_id=str(v.id),
            description="Creation bon de regularisation global",
            metadata={"direction": direction, "lines": v.lines.count()},
        )
    _invalidate_central_recap_cache()
    if len(created_vouchers) == 1:
        return redirect("voucher_detail", pk=created_vouchers[0].pk)
    return redirect("final_report")


@login_required
def carry_forward_view(request: HttpRequest) -> HttpResponse:
    """Prévisualise et confirme le report de la balance vers l'exercice N+1."""
    if not _can_close_fiscal_year(request.user):
        messages.error(request, "Seuls les admins et comptables peuvent cloturer un exercice.")
        return redirect("final_report")

    if not request.current_structure or not request.current_fiscal_year:
        messages.error(request, "Veuillez selectionner une structure et un exercice.")
        return redirect("select_structure")

    structure = request.current_structure
    current_fy = request.current_fiscal_year

    if current_fy.is_closed:
        messages.error(request, f"L'exercice {current_fy.year} est deja cloture.")
        return redirect("final_report")

    summary = compute_stock_summary(structure=structure, fiscal_year=current_fy)

    preview_rows = []
    total_qty = Decimal("0")
    total_amount = Decimal("0")
    for row in summary:
        reg_qty = row.inventory_gap if row.pv_is_filled else Decimal("0")
        final_qty = row.remaining_qty + reg_qty
        valeur = final_qty * row.cmup if row.cmup > Decimal("0") else Decimal("0")
        preview_rows.append({
            "account_code": row.account_code,
            "name": row.name,
            "unit": row.unit,
            "final_qty": final_qty,
            "cmup": row.cmup,
            "valeur": valeur,
        })
        total_qty += final_qty
        total_amount += valeur

    checklist = {
        "has_materials": len(summary) > 0,
        "all_pv_filled": all(item.pv_is_filled for item in summary) if summary else False,
        "no_negative_stock": all(item.remaining_qty >= Decimal("0") for item in summary),
    }
    checklist_ok = all(checklist.values())

    if request.method == "POST":
        if not checklist_ok:
            messages.error(request, "Checklist de cloture non satisfaite. Corrigez les points obligatoires avant le report.")
            return redirect("carry_forward")
        try:
            new_fy, vouchers = carry_forward_to_new_year(
                structure=structure,
                current_fy=current_fy,
                user=request.user,
            )
            request.session["fiscal_year_id"] = str(new_fy.id)
            if vouchers:
                nums = ", ".join(str(v.number) for v in vouchers)
                _audit(
                    request,
                    action=AuditLog.ACTION_REPORT_YEAR,
                    entity="FiscalYear",
                    entity_id=str(new_fy.id),
                    description=f"Report exercice {current_fy.year} vers {new_fy.year}",
                    metadata={
                        "opening_voucher_numbers": nums,
                        "opening_voucher_count": len(vouchers),
                    },
                )
                messages.success(
                    request,
                    f"Exercice {current_fy.year} cloture. Exercice {new_fy.year} cree avec les Bons N°{nums} (balance d'entree, {len(preview_rows)} lignes).",
                )
                _invalidate_central_recap_cache()
                if len(vouchers) == 1:
                    return redirect("voucher_detail", pk=vouchers[0].pk)
                return redirect("voucher_list")
            else:
                _audit(
                    request,
                    action=AuditLog.ACTION_REPORT_YEAR,
                    entity="FiscalYear",
                    entity_id=str(new_fy.id),
                    description=f"Report exercice {current_fy.year} vers {new_fy.year} sans lignes d'entree",
                )
                messages.success(
                    request,
                    f"Exercice {current_fy.year} cloture. Exercice {new_fy.year} cree (aucune ligne a reporter).",
                )
                _invalidate_central_recap_cache()
                return redirect("screen_operations")
        except ValueError as exc:
            messages.error(request, str(exc))
            return redirect("carry_forward")

    context = {
        "structure": structure,
        "current_fy": current_fy,
        "new_year": current_fy.year + 1,
        "preview_rows": preview_rows,
        "total_qty": total_qty,
        "total_amount": total_amount,
        "checklist": checklist,
        "checklist_ok": checklist_ok,
    }
    return render(request, "inventory/carry_forward.html", context)


_ALLOWED_LOGO_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"}
_ALLOWED_LOGO_CONTENT_TYPES = {
    "image/png",
    "image/jpeg",
    "image/gif",
    "image/webp",
    "image/svg+xml",
}
_MAX_LOGO_UPLOAD_BYTES = 2 * 1024 * 1024  # 2 MB


@login_required
def settings_view(request: HttpRequest) -> HttpResponse:
    if not _is_admin_user(request.user):
        messages.error(request, "Seuls les admins peuvent modifier les parametres.")
        return redirect("screen_admin")

    editable_keys = [
        "country",
        "ministry",
        "structure_label",
        "address",
        "phone",
        "email",
        "fiscal_year",
    ]
    if request.method == "POST":
        if request.current_structure is None:
            messages.error(request, "Selectionnez d'abord un service avant d'enregistrer les parametres.")
            return redirect("select_structure")

        for key in editable_keys:
            set_setting(key, request.POST.get(key, ""), structure=request.current_structure)

        uploaded_logo = request.FILES.get("structure_logo")
        remove_logo = request.POST.get("remove_logo") == "1"
        if remove_logo:
            if request.current_structure.logo:
                request.current_structure.logo.delete(save=False)
            request.current_structure.logo = ""
            request.current_structure.save(update_fields=["logo"])
        elif uploaded_logo:
            logo_ext = Path(uploaded_logo.name).suffix.lower()
            logo_content_type = (uploaded_logo.content_type or "").lower()
            if logo_ext not in _ALLOWED_LOGO_EXTENSIONS:
                messages.error(request, f"Extension logo non autorisee : {logo_ext}. Extensions acceptees : {', '.join(sorted(_ALLOWED_LOGO_EXTENSIONS))}")
                return redirect("settings")
            if logo_content_type and logo_content_type not in _ALLOWED_LOGO_CONTENT_TYPES:
                messages.error(request, f"Type MIME logo non autorise : {logo_content_type}.")
                return redirect("settings")
            if uploaded_logo.size > _MAX_LOGO_UPLOAD_BYTES:
                messages.error(request, f"Logo trop volumineux. Taille max : {_MAX_LOGO_UPLOAD_BYTES // (1024 * 1024)} MB.")
                return redirect("settings")
            safe_logo_name = _safe_filename(uploaded_logo.name)
            if not safe_logo_name:
                messages.error(request, "Nom de fichier logo invalide.")
                return redirect("settings")
            uploaded_logo.name = safe_logo_name
            request.current_structure.logo = uploaded_logo
            request.current_structure.save(update_fields=["logo"])

        messages.success(request, "Parametres enregistres.")
        return redirect("settings")
    values = {
        key: get_setting(
            key,
            str(request.current_fiscal_year.year) if key == "fiscal_year" and request.current_fiscal_year else "",
            structure=request.current_structure,
        )
        for key in editable_keys
    }
    return render(
        request,
        "inventory/settings.html",
        {
            "values": values,
            "current_structure": request.current_structure,
        },
    )


@login_required
def select_structure_view(request: HttpRequest) -> HttpResponse:
    """Allow users to select their current structure and fiscal year."""
    profile = _get_or_create_profile(request.user)
    if profile is None:
        messages.error(request, "Profil utilisateur introuvable.")
        return redirect("menu")
    is_superuser = request.user.is_superuser
    
    # Get available structures according to profile scope
    available_structures = profile.accessible_structures_qs().order_by("code")

    selected_structure = request.current_structure
    if selected_structure and not available_structures.filter(id=selected_structure.id).exists():
        selected_structure = None
    if selected_structure is None:
        selected_structure = available_structures.first()
    
    if request.method == "POST":
        action = (request.POST.get("action") or "").strip()
        structure_id = request.POST.get("structure_id")
        fiscal_year_id = request.POST.get("fiscal_year_id")

        if action == "reopen_fiscal_year":
            if not is_superuser:
                messages.error(request, "Seul un superuser peut reouvrir un exercice cloture.")
                return redirect("select_structure")

            reopen_fiscal_year_id = request.POST.get("reopen_fiscal_year_id")
            try:
                selected_structure = available_structures.get(id=structure_id)
                fiscal_year = FiscalYear.objects.get(id=reopen_fiscal_year_id, structure=selected_structure)
            except (Structure.DoesNotExist, FiscalYear.DoesNotExist, ValueError, TypeError):
                messages.error(request, "Selection invalide pour la reouverture de l'exercice.")
                return redirect("select_structure")

            if not fiscal_year.is_closed:
                messages.info(request, f"L'exercice {fiscal_year.year} est deja ouvert.")
                return redirect("select_structure")

            fiscal_year.is_closed = False
            fiscal_year.is_active = True
            fiscal_year.save(update_fields=["is_closed", "is_active", "updated_at"])
            request.session["structure_id"] = str(selected_structure.id)
            request.session["fiscal_year_id"] = str(fiscal_year.id)
            _audit(
                request,
                action=AuditLog.ACTION_UPDATE,
                entity="FiscalYear",
                entity_id=str(fiscal_year.id),
                description=f"Reouverture exercice {fiscal_year.year}",
                metadata={"structure_id": selected_structure.id, "year": fiscal_year.year},
            )
            messages.success(request, f"Exercice {fiscal_year.year} reouvert avec succes.")
            return redirect("select_structure")
        
        try:
            selected_structure = available_structures.get(id=structure_id)

            # Check permissions (already filtered by available_structures, but explicit for safety)
            if not is_superuser and not available_structures.filter(id=structure_id).exists():
                messages.error(request, "Vous n'avez pas acces a ce service.")
                return redirect("select_structure")
            
            request.session["structure_id"] = structure_id
            
            if fiscal_year_id:
                fiscal_year = FiscalYear.objects.get(id=fiscal_year_id, structure=selected_structure)
                request.session["fiscal_year_id"] = fiscal_year_id
            
            messages.success(request, f"Service {selected_structure.code} selectionne.")
            return redirect("screen_admin")
        except (Structure.DoesNotExist, FiscalYear.DoesNotExist):
            messages.error(request, "Selection invalide.")
    
    fiscal_years_by_structure: dict[str, list[dict[str, str | int | bool]]] = {}
    for structure in available_structures:
        # Tous les utilisateurs voient tous les exercices (clos ou non) de leur propre structure
        fiscal_years_qs = FiscalYear.objects.filter(structure=structure).order_by("-year")
        fiscal_years_by_structure[str(structure.id)] = [
            {
                "id": fy.id,
                "year": fy.year,
                "start_date": fy.start_date.strftime("%d/%m/%Y"),
                "end_date": fy.end_date.strftime("%d/%m/%Y"),
                "is_active": fy.is_active,
                "is_closed": fy.is_closed,
            }
            for fy in fiscal_years_qs
        ]

    fiscal_years = fiscal_years_by_structure.get(str(selected_structure.id), []) if selected_structure else []
    closed_fiscal_years = (
        FiscalYear.objects.filter(structure=selected_structure, is_closed=True).order_by("-year")
        if selected_structure
        else FiscalYear.objects.none()
    )
    
    context = {
        "available_structures": available_structures,
        "fiscal_years": fiscal_years,
        "fiscal_years_by_structure": fiscal_years_by_structure,
        "closed_fiscal_years": closed_fiscal_years,
        "current_structure": selected_structure,
        "current_fiscal_year": request.current_fiscal_year,
        "is_superuser": is_superuser,
    }
    return render(request, "inventory/select_structure.html", context)


@login_required
def central_recap_view(request: HttpRequest) -> HttpResponse:
    """Balance generale consolidee (services + ministere) pour les admins."""
    profile = _get_or_create_profile(request.user)
    if profile is None:
        messages.error(request, "Profil utilisateur introuvable.")
        return redirect("dashboard")
    if profile.role != "admin":
        messages.error(request, "Seuls les admins peuvent voir le rapport central.")
        return redirect("dashboard")

    fiscal_year_param = request.GET.get("fiscal_year")
    structure_type_param = (request.GET.get("structure_type") or "").strip()
    structure_type_id: int | None = None
    if structure_type_param:
        try:
            structure_type_id = int(structure_type_param)
        except ValueError:
            structure_type_id = None

    central_data = _get_cached_central_recap_data(
        fiscal_year=fiscal_year_param,
        structure_type_id=structure_type_id,
    )

    context = {
        "aggregate_data": central_data["aggregate_data"],
        "aggregate_totals": central_data["aggregate_totals"],
        "ministry_balance_rows": central_data["ministry_balance_rows"],
        "ministry_totals": central_data["ministry_totals"],
        "all_fiscal_years": central_data["all_fiscal_years"],
        "all_structure_types": central_data["all_structure_types"],
        "selected_fiscal_year": str(central_data["selected_fiscal_year"]) if central_data["selected_fiscal_year"] is not None else "",
        "selected_structure_type_id": central_data["selected_structure_type_id"],
        "selected_structure_type": central_data["selected_structure_type"],
        "total_materials": central_data["total_materials"],
        "total_vouchers": central_data["total_vouchers"],
        "total_comptable_value": central_data["total_comptable_value"],
        "total_regulation_value": central_data["total_regulation_value"],
        "total_final_value": central_data["total_final_value"],
    }
    return render(request, "inventory/central_recap.html", context)


@login_required
def central_recap_excel_view(request: HttpRequest) -> HttpResponse:
    profile = _get_or_create_profile(request.user)
    if profile is None:
        messages.error(request, "Profil utilisateur introuvable.")
        return redirect("dashboard")
    if profile.role != "admin":
        messages.error(request, "Seuls les admins peuvent exporter le rapport central.")
        return redirect("dashboard")

    fiscal_year_param = request.GET.get("fiscal_year")
    structure_type_param = (request.GET.get("structure_type") or "").strip()
    structure_type_id: int | None = None
    if structure_type_param:
        try:
            structure_type_id = int(structure_type_param)
        except ValueError:
            structure_type_id = None

    central_data = _get_cached_central_recap_data(
        fiscal_year=fiscal_year_param,
        structure_type_id=structure_type_id,
    )

    return _build_central_excel_response(
        central_data=central_data,
        filename="balance_generale_consolidee_ministere.xlsx",
        fiscal_year_label=str(central_data["selected_fiscal_year"]) if central_data["selected_fiscal_year"] is not None else "N/A",
        structure_type_label=(
            f"{central_data['selected_structure_type'].code} - {central_data['selected_structure_type'].name}"
            if central_data["selected_structure_type"]
            else "Tous"
        ),
    )


def _get_cached_central_recap_data(*, fiscal_year: str | None, structure_type_id: int | None) -> dict:
    timeout = int(getattr(settings, "CENTRAL_RECAP_CACHE_TIMEOUT_SECONDS", 60))
    if timeout <= 0:
        return _build_central_recap_data(fiscal_year=fiscal_year, structure_type_id=structure_type_id)

    cache_version = _get_central_recap_cache_version()
    cache_key = f"central_recap:v{cache_version}:{fiscal_year or 'auto'}:type:{structure_type_id or 'all'}"
    cached = cache.get(cache_key)
    if cached is not None:
        return cached

    central_data = _build_central_recap_data(fiscal_year=fiscal_year, structure_type_id=structure_type_id)
    cache.set(cache_key, central_data, timeout=timeout)
    return central_data


def _build_central_recap_data(*, fiscal_year: str | None, structure_type_id: int | None) -> dict:
    structures_qs = Structure.objects.filter(is_active=True)
    all_structure_types = list(StructureType.objects.order_by("code"))
    selected_structure_type = None
    if structure_type_id is not None:
        structures_qs = structures_qs.filter(structure_type_id=structure_type_id)
        selected_structure_type = next((t for t in all_structure_types if t.id == structure_type_id), None)
        if selected_structure_type is None:
            structure_type_id = None
            structures_qs = Structure.objects.filter(is_active=True)

    structures = list(structures_qs.order_by("code"))
    aggregate_data = []
    ministry_balance_by_account: dict[str, dict] = {}

    years_by_structure: dict[int, set[int]] = defaultdict(set)
    fiscal_year_pairs = FiscalYear.objects.filter(structure__in=structures).values_list("structure_id", "year")
    for structure_id, year in fiscal_year_pairs:
        years_by_structure[structure_id].add(year)

    common_years: list[int] = []
    if years_by_structure:
        year_sets = list(years_by_structure.values())
        common = set.intersection(*year_sets) if all(year_sets) else set()
        common_years = sorted(common, reverse=True)

    # Fallback: si l'intersection est vide (services incomplets), proposer les annees existantes.
    if not common_years:
        all_years = set(FiscalYear.objects.values_list("year", flat=True))
        common_years = sorted(all_years, reverse=True)

    selected_year: int | None = None
    if fiscal_year:
        try:
            requested_year = int(fiscal_year)
        except (TypeError, ValueError):
            requested_year = None
        if requested_year not in common_years:
            fy_by_id = FiscalYear.objects.filter(id=fiscal_year).first()
            if fy_by_id:
                requested_year = fy_by_id.year
        if requested_year in common_years:
            selected_year = requested_year

    if selected_year is None and common_years:
        selected_year = common_years[0]

    fiscal_year_by_structure: dict[int, FiscalYear] = {}
    if selected_year is not None:
        for fy in FiscalYear.objects.filter(structure__in=structures, year=selected_year).order_by("structure_id", "-id"):
            if fy.structure_id not in fiscal_year_by_structure:
                fiscal_year_by_structure[fy.structure_id] = fy

    voucher_counts: dict[tuple[int, int], int] = {}
    if fiscal_year_by_structure:
        voucher_rows = (
            Voucher.objects.filter(fiscal_year__in=fiscal_year_by_structure.values())
            .values("structure_id", "fiscal_year_id")
            .annotate(total=Count("id"))
        )
        for row in voucher_rows:
            voucher_counts[(row["structure_id"], row["fiscal_year_id"])] = row["total"]

    for structure in structures:
        if selected_year is None:
            continue

        fy = fiscal_year_by_structure.get(structure.id)
        if not fy:
            continue

        summary = compute_stock_summary(structure=structure, fiscal_year=fy)
        comptable_qty = Decimal("0")
        comptable_value = Decimal("0")
        pending_qty = Decimal("0")
        in_service_qty = Decimal("0")
        provisional_qty = Decimal("0")
        regulation_qty = Decimal("0")
        regulation_value = Decimal("0")
        final_qty = Decimal("0")
        final_value = Decimal("0")

        for item in summary:
            item_reg_qty = item.inventory_gap if item.pv_is_filled else Decimal("0")
            item_reg_value = item.inventory_gap_value if item.pv_is_filled else Decimal("0")
            item_final_qty = item.remaining_qty + item_reg_qty
            item_final_value = item.remaining_amount + item_reg_value

            comptable_qty += item.remaining_qty
            comptable_value += item.remaining_amount
            pending_qty += item.pending_qty if item.pv_is_filled else Decimal("0")
            in_service_qty += item.in_service_qty if item.pv_is_filled else Decimal("0")
            provisional_qty += item.provisional_qty if item.pv_is_filled else Decimal("0")
            regulation_qty += item_reg_qty
            regulation_value += item_reg_value
            final_qty += item_final_qty
            final_value += item_final_value

            if item.account_code not in ministry_balance_by_account:
                ministry_balance_by_account[item.account_code] = {
                    "account_code": item.account_code,
                    "name": item.name,
                    "unit": item.unit,
                    "qty_in": Decimal("0"),
                    "qty_in_opening": Decimal("0"),
                    "qty_in_period": Decimal("0"),
                    "qty_out": Decimal("0"),
                    "remaining_qty": Decimal("0"),
                    "pending_qty": Decimal("0"),
                    "in_service_qty": Decimal("0"),
                    "provisional_qty": Decimal("0"),
                    "total_pv_qty": Decimal("0"),
                    "final_amount": Decimal("0"),
                }

            row = ministry_balance_by_account[item.account_code]
            row["qty_in"] += item.qty_in
            row["qty_in_opening"] += item.qty_in_opening
            row["qty_in_period"] += item.qty_in_period
            row["qty_out"] += item.qty_out
            row["remaining_qty"] += item.remaining_qty
            row["pending_qty"] += item.pending_qty if item.pv_is_filled else Decimal("0")
            row["in_service_qty"] += item.in_service_qty if item.pv_is_filled else Decimal("0")
            row["provisional_qty"] += item.provisional_qty if item.pv_is_filled else Decimal("0")
            row["final_amount"] += item_final_value

        aggregate_data.append(
            {
                "structure": structure,
                "fiscal_year": fy,
                "materials_count": len(summary),
                "comptable_qty": comptable_qty,
                "comptable_value": comptable_value,
                "pending_qty": pending_qty,
                "in_service_qty": in_service_qty,
                "provisional_qty": provisional_qty,
                "regulation_qty": regulation_qty,
                "regulation_value": regulation_value,
                "final_qty": final_qty,
                "final_value": final_value,
                "vouchers_count": voucher_counts.get((structure.id, fy.id), 0),
                "detail_rows": [
                    {
                        "account_code": item.account_code,
                        "name": item.name,
                        "unit": item.unit,
                        "qty_in": item.qty_in,
                        "qty_in_opening": item.qty_in_opening,
                        "qty_in_period": item.qty_in_period,
                        "qty_out": item.qty_out,
                        "remaining_qty": item.remaining_qty,
                        "total_pv_qty": item.inventory_total_qty if item.pv_is_filled else Decimal("0"),
                        "gap": item.inventory_gap if item.pv_is_filled else Decimal("0"),
                        "final_qty": item.remaining_qty + (item.inventory_gap if item.pv_is_filled else Decimal("0")),
                        "cmup": item.cmup,
                        "remaining_amount_final": (item.remaining_qty + (item.inventory_gap if item.pv_is_filled else Decimal("0"))) * item.cmup if item.cmup > Decimal("0") else Decimal("0"),
                        "pv_is_filled": item.pv_is_filled,
                    }
                    for item in summary
                ],
            }
        )

    ministry_balance_rows = []
    for row in ministry_balance_by_account.values():
        row["total_pv_qty"] = row["pending_qty"] + row["in_service_qty"] + row["provisional_qty"]
        ministry_balance_rows.append(row)

    ministry_balance_rows = sorted(ministry_balance_rows, key=lambda item: item["account_code"])

    aggregate_totals = {
        "materials_count": sum((item["materials_count"] for item in aggregate_data), 0),
        "vouchers_count": sum((item["vouchers_count"] for item in aggregate_data), 0),
        "comptable_qty": sum((item["comptable_qty"] for item in aggregate_data), Decimal("0")),
        "comptable_value": sum((item["comptable_value"] for item in aggregate_data), Decimal("0")),
        "pending_qty": sum((item["pending_qty"] for item in aggregate_data), Decimal("0")),
        "in_service_qty": sum((item["in_service_qty"] for item in aggregate_data), Decimal("0")),
        "provisional_qty": sum((item["provisional_qty"] for item in aggregate_data), Decimal("0")),
        "regulation_value": sum((item["regulation_value"] for item in aggregate_data), Decimal("0")),
        "final_value": sum((item["final_value"] for item in aggregate_data), Decimal("0")),
    }

    ministry_totals = {
        "qty_in": sum((item["qty_in"] for item in ministry_balance_rows), Decimal("0")),
        "qty_in_opening": sum((item["qty_in_opening"] for item in ministry_balance_rows), Decimal("0")),
        "qty_in_period": sum((item["qty_in_period"] for item in ministry_balance_rows), Decimal("0")),
        "qty_out": sum((item["qty_out"] for item in ministry_balance_rows), Decimal("0")),
        "remaining_qty": sum((item["remaining_qty"] for item in ministry_balance_rows), Decimal("0")),
        "pending_qty": sum((item["pending_qty"] for item in ministry_balance_rows), Decimal("0")),
        "in_service_qty": sum((item["in_service_qty"] for item in ministry_balance_rows), Decimal("0")),
        "provisional_qty": sum((item["provisional_qty"] for item in ministry_balance_rows), Decimal("0")),
        "total_pv_qty": sum((item["total_pv_qty"] for item in ministry_balance_rows), Decimal("0")),
        "final_amount": sum((item["final_amount"] for item in ministry_balance_rows), Decimal("0")),
    }

    return {
        "aggregate_data": aggregate_data,
        "aggregate_totals": aggregate_totals,
        "ministry_balance_rows": ministry_balance_rows,
        "ministry_totals": ministry_totals,
        "all_fiscal_years": common_years,
        "all_structure_types": all_structure_types,
        "selected_fiscal_year": selected_year,
        "selected_structure_type_id": structure_type_id,
        "selected_structure_type": selected_structure_type,
        "total_materials": sum((item["materials_count"] for item in aggregate_data), 0),
        "total_vouchers": sum((item["vouchers_count"] for item in aggregate_data), 0),
        "total_comptable_value": sum((item["comptable_value"] for item in aggregate_data), Decimal("0")),
        "total_regulation_value": sum((item["regulation_value"] for item in aggregate_data), Decimal("0")),
        "total_final_value": sum((item["final_value"] for item in aggregate_data), Decimal("0")),
    }


def _build_central_excel_response(
    *,
    central_data: dict,
    filename: str,
    fiscal_year_label: str,
    structure_type_label: str,
) -> HttpResponse:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except Exception as exc:  # noqa: BLE001
        return HttpResponse(
            f"Generation Excel indisponible: {exc}. Installez openpyxl.",
            status=500,
            content_type="text/plain; charset=utf-8",
        )

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Synthese ministere"

    title_fill = PatternFill(start_color="D7B98C", end_color="D7B98C", fill_type="solid")
    header_fill = PatternFill(start_color="F0ECE4", end_color="F0ECE4", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="B8AA95"),
        right=Side(style="thin", color="B8AA95"),
        top=Side(style="thin", color="B8AA95"),
        bottom=Side(style="thin", color="B8AA95"),
    )

    context_lines = [
        "Rapport central ministeriel",
        f"Exercice filtre : {fiscal_year_label}",
        f"Type de service : {structure_type_label}",
        f"Services consolides : {len(central_data['aggregate_data'])}",
        f"Date d'edition : {date.today().strftime('%d/%m/%Y')}",
    ]
    row_cursor = 1
    for line in context_lines:
        summary_sheet.cell(row=row_cursor, column=1, value=line)
        row_cursor += 1
    row_cursor += 1

    summary_sheet.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=8)
    title_cell = summary_sheet.cell(row=row_cursor, column=1, value="Synthese par service")
    title_cell.font = Font(bold=True, size=12)
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="left")
    row_cursor += 1

    aggregate_headers = [
        "Service",
        "Exercice",
        "Bons",
        "Qte comptable",
        "En attente d'affectation",
        "En service",
        "En sortie provisoire",
        "Valeur du stock",
    ]
    for col_idx, header in enumerate(aggregate_headers, start=1):
        cell = summary_sheet.cell(row=row_cursor, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
    row_cursor += 1

    for item in central_data["aggregate_data"]:
        row_values = [
            f"{item['structure'].code} - {item['structure'].name}",
            item["fiscal_year"].year,
            item["vouchers_count"],
            item["comptable_qty"],
            item["pending_qty"],
            item["in_service_qty"],
            item["provisional_qty"],
            item["final_value"],
        ]
        for col_idx, value in enumerate(row_values, start=1):
            cell = summary_sheet.cell(row=row_cursor, column=col_idx, value=value)
            cell.border = thin_border
            if isinstance(value, Decimal):
                cell.number_format = "#,##0.00"
        row_cursor += 1

    aggregate_totals = central_data["aggregate_totals"]
    aggregate_total_row = [
        "TOTAL",
        fiscal_year_label,
        aggregate_totals["vouchers_count"],
        aggregate_totals["comptable_qty"],
        aggregate_totals["pending_qty"],
        aggregate_totals["in_service_qty"],
        aggregate_totals["provisional_qty"],
        aggregate_totals["final_value"],
    ]
    for col_idx, value in enumerate(aggregate_total_row, start=1):
        cell = summary_sheet.cell(row=row_cursor, column=col_idx, value=value)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = thin_border
        if isinstance(value, Decimal):
            cell.number_format = "#,##0.00"
    row_cursor += 1

    row_cursor += 2
    summary_sheet.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=11)
    ministry_title = summary_sheet.cell(row=row_cursor, column=1, value="Balance generale consolidee ministere")
    ministry_title.font = Font(bold=True, size=12)
    ministry_title.fill = title_fill
    ministry_title.alignment = Alignment(horizontal="left")
    row_cursor += 1

    balance_headers = [
        "Compte",
        "Matiere",
        "Unite",
        "Existant debut gestion",
        "Entree en cours gestion",
        "Qte sortie",
        "Qte comptable",
        "Total PV",
        "Valeur stock",
    ]
    for col_idx, header in enumerate(balance_headers, start=1):
        cell = summary_sheet.cell(row=row_cursor, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
    row_cursor += 1

    for row in central_data["ministry_balance_rows"]:
        row_values = [
            row["account_code"],
            row["name"],
            row["unit"],
            row["qty_in_opening"],
            row["qty_in_period"],
            row["qty_out"],
            row["remaining_qty"],
            row["total_pv_qty"],
            row["final_amount"],
        ]
        for col_idx, value in enumerate(row_values, start=1):
            cell = summary_sheet.cell(row=row_cursor, column=col_idx, value=value)
            cell.border = thin_border
            if isinstance(value, Decimal):
                cell.number_format = "#,##0.00"
        row_cursor += 1

    ministry_totals = central_data["ministry_totals"]
    ministry_total_row = [
        "",
        "TOTAL",
        "",
        ministry_totals["qty_in_opening"],
        ministry_totals["qty_in_period"],
        ministry_totals["qty_out"],
        ministry_totals["remaining_qty"],
        ministry_totals["total_pv_qty"],
        ministry_totals["final_amount"],
    ]
    for col_idx, value in enumerate(ministry_total_row, start=1):
        cell = summary_sheet.cell(row=row_cursor, column=col_idx, value=value)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = thin_border
        if isinstance(value, Decimal):
            cell.number_format = "#,##0.00"
    row_cursor += 1

    summary_sheet.freeze_panes = "A9"
    _autosize_sheet_columns(summary_sheet, max_col=8)

    used_sheet_names = {summary_sheet.title}
    for item in central_data["aggregate_data"]:
        base_name = f"{item['structure'].code}_{item['fiscal_year'].year}"
        sheet_name = _excel_safe_sheet_name(base_name, used_sheet_names)
        used_sheet_names.add(sheet_name)
        detail_sheet = workbook.create_sheet(title=sheet_name)

        detail_sheet.cell(row=1, column=1, value="Service").font = Font(bold=True)
        detail_sheet.cell(row=1, column=2, value=f"{item['structure'].code} - {item['structure'].name}")
        detail_sheet.cell(row=2, column=1, value="Exercice").font = Font(bold=True)
        detail_sheet.cell(row=2, column=2, value=item["fiscal_year"].year)
        detail_sheet.cell(row=3, column=1, value="Date d'edition").font = Font(bold=True)
        detail_sheet.cell(row=3, column=2, value=date.today().strftime("%d/%m/%Y"))

        start_row = 5
        for col_idx, header in enumerate(balance_headers, start=1):
            cell = detail_sheet.cell(row=start_row, column=col_idx, value=header)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        row_cursor = start_row + 1
        for row in item["detail_rows"]:
            row_values = [
                row["account_code"],
                row["name"],
                row["unit"],
                row["qty_in"],
                row["qty_out"],
                row["remaining_qty"],
                row["total_pv_qty"],
                row["remaining_amount_final"],
            ]
            for col_idx, value in enumerate(row_values, start=1):
                cell = detail_sheet.cell(row=row_cursor, column=col_idx, value=value)
                cell.border = thin_border
                if isinstance(value, Decimal):
                    cell.number_format = "#,##0.00"
            row_cursor += 1

        detail_sheet.freeze_panes = "A6"
        _autosize_sheet_columns(detail_sheet, max_col=8)

    kpi_sheet = workbook.create_sheet(title="KPI qualite")
    kpi_sheet.cell(row=1, column=1, value="Indicateurs qualite des donnees").font = Font(bold=True, size=13)
    kpi_sheet.cell(row=2, column=1, value=f"Exercice filtre : {fiscal_year_label}")
    kpi_sheet.cell(row=3, column=1, value=f"Date d'edition : {date.today().strftime('%d/%m/%Y')}")

    kpi_headers = [
        "Service",
        "Exercice",
        "Matieres",
        "PV renseignes",
        "PV manquants",
        "Ecarts non nuls",
        "Stocks negatifs",
        "Taux PV",
    ]
    kpi_good_fill = PatternFill(start_color="DFF0D8", end_color="DFF0D8", fill_type="solid")
    kpi_warn_fill = PatternFill(start_color="FFF4CC", end_color="FFF4CC", fill_type="solid")
    kpi_bad_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")

    def _kpi_fill_for_rate(rate: Decimal) -> PatternFill:
        if rate >= Decimal("0.95"):
            return kpi_good_fill
        if rate >= Decimal("0.80"):
            return kpi_warn_fill
        return kpi_bad_fill
    start_row = 5
    for col_idx, header in enumerate(kpi_headers, start=1):
        cell = kpi_sheet.cell(row=start_row, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    row_cursor = start_row + 1
    recommendation_rows: list[dict[str, str]] = []
    totals = {
        "materials": 0,
        "pv_filled": 0,
        "pv_missing": 0,
        "nonzero_gap": 0,
        "negative_stock": 0,
    }
    for item in central_data["aggregate_data"]:
        detail_rows = item.get("detail_rows", [])
        materials_count = len(detail_rows)
        pv_filled = sum(1 for row in detail_rows if row.get("pv_is_filled"))
        pv_missing = materials_count - pv_filled
        nonzero_gap = sum(1 for row in detail_rows if row.get("gap") != Decimal("0"))
        negative_stock = sum(1 for row in detail_rows if row.get("final_qty", Decimal("0")) < Decimal("0"))
        pv_rate = (Decimal(pv_filled) / Decimal(materials_count)) if materials_count else Decimal("0")

        values = [
            item["structure"].code,
            item["fiscal_year"].year,
            materials_count,
            pv_filled,
            pv_missing,
            nonzero_gap,
            negative_stock,
            pv_rate,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = kpi_sheet.cell(row=row_cursor, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx == 8:
                cell.number_format = "0.00%"

        # Couleurs d'aide a la decision sur les indicateurs de qualite.
        kpi_sheet.cell(row=row_cursor, column=5).fill = kpi_bad_fill if pv_missing > 0 else kpi_good_fill
        kpi_sheet.cell(row=row_cursor, column=6).fill = kpi_warn_fill if nonzero_gap > 0 else kpi_good_fill
        kpi_sheet.cell(row=row_cursor, column=7).fill = kpi_bad_fill if negative_stock > 0 else kpi_good_fill
        kpi_sheet.cell(row=row_cursor, column=8).fill = _kpi_fill_for_rate(pv_rate)

        if pv_missing > 0:
            recommendation_rows.append(
                {
                    "service": item["structure"].code,
                    "exercice": str(item["fiscal_year"].year),
                    "priorite": "Haute",
                    "constat": f"{pv_missing} matiere(s) sans PV renseigne",
                    "action": "Completer les PV manquants avant cloture.",
                    "delai": "Sous 7 jours",
                }
            )
        if negative_stock > 0:
            recommendation_rows.append(
                {
                    "service": item["structure"].code,
                    "exercice": str(item["fiscal_year"].year),
                    "priorite": "Haute",
                    "constat": f"{negative_stock} stock(s) negatif(s)",
                    "action": "Verifier les sorties, corriger les ecritures et regularisations.",
                    "delai": "Sous 7 jours",
                }
            )
        if nonzero_gap > 0:
            recommendation_rows.append(
                {
                    "service": item["structure"].code,
                    "exercice": str(item["fiscal_year"].year),
                    "priorite": "Moyenne",
                    "constat": f"{nonzero_gap} ecart(s) de recensement non nul(s)",
                    "action": "Documenter et traiter les ecarts avant validation finale.",
                    "delai": "Sous 30 jours",
                }
            )
        if pv_rate < Decimal("0.95") and pv_missing == 0:
            recommendation_rows.append(
                {
                    "service": item["structure"].code,
                    "exercice": str(item["fiscal_year"].year),
                    "priorite": "Moyenne",
                    "constat": f"Taux PV a {pv_rate * Decimal('100'):.2f}%",
                    "action": "Planifier un controle qualite mensuel des fiches de stock.",
                    "delai": "Sous 30 jours",
                }
            )
        row_cursor += 1

        totals["materials"] += materials_count
        totals["pv_filled"] += pv_filled
        totals["pv_missing"] += pv_missing
        totals["nonzero_gap"] += nonzero_gap
        totals["negative_stock"] += negative_stock

    global_rate = (Decimal(totals["pv_filled"]) / Decimal(totals["materials"])) if totals["materials"] else Decimal("0")
    total_values = [
        "TOTAL",
        "-",
        totals["materials"],
        totals["pv_filled"],
        totals["pv_missing"],
        totals["nonzero_gap"],
        totals["negative_stock"],
        global_rate,
    ]
    for col_idx, value in enumerate(total_values, start=1):
        cell = kpi_sheet.cell(row=row_cursor, column=col_idx, value=value)
        cell.font = Font(bold=True)
        cell.fill = title_fill
        cell.border = thin_border
        if col_idx == 8:
            cell.number_format = "0.00%"

    kpi_sheet.cell(row=row_cursor, column=8).fill = _kpi_fill_for_rate(global_rate)

    legend_title_row = row_cursor + 2
    kpi_sheet.merge_cells(start_row=legend_title_row, start_column=1, end_row=legend_title_row, end_column=4)
    legend_title_cell = kpi_sheet.cell(row=legend_title_row, column=1, value="Legende des seuils KPI")
    legend_title_cell.font = Font(bold=True)
    legend_title_cell.fill = title_fill
    legend_title_cell.alignment = Alignment(horizontal="left")
    legend_title_cell.border = thin_border

    legend_rows = [
        ("Vert", "Conforme", "Taux PV >= 95% et aucun incident critique", kpi_good_fill),
        ("Orange", "Vigilance", "Taux PV entre 80% et 95% ou ecarts detectes", kpi_warn_fill),
        ("Rouge", "A corriger", "Taux PV < 80% ou PV manquant/stock negatif", kpi_bad_fill),
    ]
    legend_start_row = legend_title_row + 1
    for index, legend in enumerate(legend_rows):
        color_label, status_label, criteria_label, color_fill = legend
        row_index = legend_start_row + index
        color_cell = kpi_sheet.cell(row=row_index, column=1, value=color_label)
        color_cell.fill = color_fill
        color_cell.font = Font(bold=True)
        color_cell.border = thin_border
        status_cell = kpi_sheet.cell(row=row_index, column=2, value=status_label)
        status_cell.border = thin_border
        criteria_cell = kpi_sheet.cell(row=row_index, column=3, value=criteria_label)
        criteria_cell.border = thin_border
        kpi_sheet.merge_cells(start_row=row_index, start_column=3, end_row=row_index, end_column=8)

    kpi_sheet.freeze_panes = "A6"
    _autosize_sheet_columns(kpi_sheet, max_col=8)

    reco_sheet = workbook.create_sheet(title="Recommandations")
    reco_sheet.cell(row=1, column=1, value="Plan d'actions recommande").font = Font(bold=True, size=13)
    reco_sheet.cell(row=2, column=1, value=f"Exercice filtre : {fiscal_year_label}")
    reco_sheet.cell(row=3, column=1, value=f"Date d'edition : {date.today().strftime('%d/%m/%Y')}")

    reco_headers = ["Service", "Exercice", "Priorite", "Constat", "Action recommandee", "Delai suggere"]
    reco_header_row = 5
    for col_idx, header in enumerate(reco_headers, start=1):
        cell = reco_sheet.cell(row=reco_header_row, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    if not recommendation_rows:
        recommendation_rows.append(
            {
                "service": "Global",
                "exercice": fiscal_year_label,
                "priorite": "Information",
                "constat": "Aucune anomalie critique detectee",
                "action": "Maintenir le controle periodique des donnees.",
                "delai": "Suivi trimestriel",
            }
        )

    priority_order = {"Haute": 0, "Moyenne": 1, "Information": 2}
    recommendation_rows.sort(
        key=lambda row: (priority_order.get(row.get("priorite", "Information"), 9), row.get("service", ""), row.get("exercice", ""))
    )

    reco_row_cursor = reco_header_row + 1
    for row in recommendation_rows:
        values = [row["service"], row["exercice"], row["priorite"], row["constat"], row["action"], row["delai"]]
        for col_idx, value in enumerate(values, start=1):
            cell = reco_sheet.cell(row=reco_row_cursor, column=col_idx, value=value)
            cell.border = thin_border
            if col_idx == 3:
                if value == "Haute":
                    cell.fill = kpi_bad_fill
                elif value == "Moyenne":
                    cell.fill = kpi_warn_fill
                elif value == "Information":
                    cell.fill = kpi_good_fill
        reco_row_cursor += 1

    reco_sheet.freeze_panes = "A6"
    _autosize_sheet_columns(reco_sheet, max_col=6)

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


def _autosize_sheet_columns(sheet, *, max_col: int) -> None:
    from openpyxl.utils import get_column_letter

    for col_idx in range(1, max_col + 1):
        max_length = 10
        for cells in sheet.iter_cols(min_col=col_idx, max_col=col_idx, min_row=1, max_row=sheet.max_row):
            for cell in cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 45)


def _excel_safe_sheet_name(base_name: str, used_names: set[str]) -> str:
    sanitized = "".join(ch for ch in base_name if ch not in "[]:*?/\\") or "Sheet"
    sanitized = sanitized[:31]
    candidate = sanitized
    suffix = 1
    while candidate in used_names:
        suffix_label = f"_{suffix}"
        candidate = f"{sanitized[:31 - len(suffix_label)]}{suffix_label}"
        suffix += 1
    return candidate
