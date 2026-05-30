from __future__ import annotations

from io import BytesIO
from io import StringIO
from pathlib import Path
import shutil
import tempfile
from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from openpyxl import load_workbook

from .forms import InternalMovementLineInput, VoucherLineInput
from .models import AuditLog, FiscalYear, Material, NomenclatureItem, PhysicalInventory, Structure, Voucher, VoucherLine
from .models import InternalMovement, Location
from .services import build_location_inventory_report, create_internal_movement, create_voucher


User = get_user_model()


class InventoryViewsTests(TestCase):
    def setUp(self) -> None:
        self.user = User.objects.create_user(username="tester", password="pass12345")
        self.user.profile.role = "admin"
        self.user.profile.save()

        self.structure = Structure.objects.create(name="Service Test", code="ST01", is_active=True)
        self.fiscal_year = FiscalYear.objects.create(
            structure=self.structure,
            year=2026,
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
            is_active=True,
        )

        self.material_1 = Material.objects.create(
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            account_code="6010",
            name="Papier",
            unit="Rame",
            group_code="60",
        )
        self.material_2 = Material.objects.create(
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            account_code="6020",
            name="Encre",
            unit="Boite",
            group_code="60",
        )

        voucher_1 = Voucher.objects.create(
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            number=1,
            voucher_type=Voucher.TYPE_ENTRY,
            operation_date=date(2026, 1, 5),
            source_or_destination="Fournisseur A",
            comments="Initial",
            created_by=self.user,
        )
        VoucherLine.objects.create(
            voucher=voucher_1,
            material=self.material_1,
            material_name=self.material_1.name,
            specification="Lot 1",
            quantity=Decimal("10"),
            unit=self.material_1.unit,
            unit_price=Decimal("500"),
            amount=Decimal("5000"),
        )

        voucher_2 = Voucher.objects.create(
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            number=2,
            voucher_type=Voucher.TYPE_ENTRY,
            operation_date=date(2026, 1, 10),
            source_or_destination="Fournisseur B",
            comments="Second",
            created_by=self.user,
        )
        VoucherLine.objects.create(
            voucher=voucher_2,
            material=self.material_2,
            material_name=self.material_2.name,
            specification="Lot 2",
            quantity=Decimal("7"),
            unit=self.material_2.unit,
            unit_price=Decimal("300"),
            amount=Decimal("2100"),
        )

        self.client.login(username="tester", password="pass12345")
        session = self.client.session
        session["structure_id"] = self.structure.id
        session["fiscal_year_id"] = self.fiscal_year.id
        session.save()

    def test_excel_exports_are_available(self) -> None:
        for url_name in ["recap_excel", "pv_recensement_excel", "final_balance_excel", "central_recap_excel"]:
            response = self.client.get(reverse(url_name))
            self.assertEqual(response.status_code, 200)
            self.assertIn(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                response["Content-Type"],
            )

    def test_entry_voucher_updates_pending_inventory_without_overwriting_origin(self) -> None:
        material = Material.objects.create(
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            account_code="1001",
            name="Ordinateur portable",
            unit="Unite",
            group_code="1",
        )

        voucher = create_voucher(
            user=self.user,
            voucher_type=Voucher.TYPE_ENTRY,
            operation_date=date(2026, 2, 1),
            source_or_destination="Fournisseur A",
            comments="Reception lot 3",
            lines=[
                VoucherLineInput(
                    material_id=material.id,
                    specification="Serie A",
                    quantity=Decimal("3"),
                    unit_price=Decimal("250000"),
                )
            ],
            structure=self.structure,
            fiscal_year=self.fiscal_year,
        )[0]

        inventory = material.physical_inventory
        self.assertEqual(voucher.source_or_destination, "Fournisseur A")
        self.assertEqual(inventory.pending_qty, Decimal("3"))
        self.assertEqual(inventory.in_service_qty, Decimal("0"))

    def test_assignment_internal_movement_updates_physical_inventory(self) -> None:
        store = Location.objects.create(
            structure=self.structure,
            name="Magasin de materiel",
            location_type=Location.TYPE_STORE_EQUIPMENT,
            responsible_name="Magasinier",
        )
        office = Location.objects.create(
            structure=self.structure,
            name="Bureau du directeur",
            location_type=Location.TYPE_OFFICE,
            responsible_name="Directeur",
        )
        PhysicalInventory.objects.create(
            material=self.material_1,
            pending_qty=Decimal("5"),
            in_service_qty=Decimal("0"),
            provisional_qty=Decimal("0"),
            unit_price=Decimal("500"),
            is_pv_filled=True,
        )

        movement = create_internal_movement(
            user=self.user,
            movement_type=InternalMovement.TYPE_ASSIGNMENT,
            operation_date=date(2026, 2, 5),
            from_location=store,
            to_location=office,
            reason="Affectation initiale",
            lines=[
                InternalMovementLineInput(
                    material_id=self.material_1.id,
                    inventory_code="1001-2026-01",
                    quantity=Decimal("2"),
                    unit_price=Decimal("500"),
                    observations="RAS",
                )
            ],
            structure=self.structure,
            fiscal_year=self.fiscal_year,
        )

        inventory = PhysicalInventory.objects.get(material=self.material_1)
        self.assertEqual(movement.lines.count(), 1)
        self.assertEqual(inventory.pending_qty, Decimal("3"))
        self.assertEqual(inventory.in_service_qty, Decimal("2"))

    def test_store_inventory_amount_sums_line_amounts(self) -> None:
        store = Location.objects.create(
            structure=self.structure,
            name="Magasin principal",
            location_type=Location.TYPE_STORE_EQUIPMENT,
            responsible_name="Magasinier",
        )

        create_voucher(
            user=self.user,
            voucher_type=Voucher.TYPE_ENTRY,
            operation_date=date(2026, 2, 1),
            source_or_destination="Fournisseur A",
            storage_location=store,
            comments="Lot 1",
            lines=[
                VoucherLineInput(
                    material_id=self.material_1.id,
                    specification="Lot 1",
                    quantity=Decimal("5"),
                    unit_price=Decimal("1000"),
                )
            ],
            structure=self.structure,
            fiscal_year=self.fiscal_year,
        )

        create_voucher(
            user=self.user,
            voucher_type=Voucher.TYPE_ENTRY,
            operation_date=date(2026, 2, 10),
            source_or_destination="Fournisseur B",
            storage_location=store,
            comments="Lot 2",
            lines=[
                VoucherLineInput(
                    material_id=self.material_1.id,
                    specification="Lot 2",
                    quantity=Decimal("5"),
                    unit_price=Decimal("1500"),
                )
            ],
            structure=self.structure,
            fiscal_year=self.fiscal_year,
        )

        report = build_location_inventory_report(store, self.fiscal_year)
        row = next(item for item in report if item.material_id == self.material_1.id)
        self.assertEqual(row.quantity, Decimal("10"))
        self.assertEqual(row.amount, Decimal("12500"))

    def test_contradictory_inventory_view_lists_location_holdings(self) -> None:
        first_group_material = Material.objects.create(
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            account_code="1002",
            name="Armoire metallique",
            unit="Unite",
            group_code="1",
        )
        store = Location.objects.create(
            structure=self.structure,
            name="Magasin de materiel",
            location_type=Location.TYPE_STORE_EQUIPMENT,
            responsible_name="Magasinier",
        )
        office = Location.objects.create(
            structure=self.structure,
            name="Bureau informatique",
            location_type=Location.TYPE_OFFICE,
            responsible_name="Chef service informatique",
        )
        PhysicalInventory.objects.create(
            material=first_group_material,
            pending_qty=Decimal("4"),
            in_service_qty=Decimal("0"),
            provisional_qty=Decimal("0"),
            unit_price=Decimal("500"),
            is_pv_filled=True,
        )
        create_internal_movement(
            user=self.user,
            movement_type=InternalMovement.TYPE_ASSIGNMENT,
            operation_date=date(2026, 3, 1),
            from_location=store,
            to_location=office,
            reason="Dotation bureau",
            lines=[
                InternalMovementLineInput(
                    material_id=first_group_material.id,
                    inventory_code="INV-001",
                    quantity=Decimal("1"),
                    unit_price=Decimal("500"),
                    observations="PC portable",
                )
            ],
            structure=self.structure,
            fiscal_year=self.fiscal_year,
        )

        response = self.client.get(reverse("contradictory_inventory"), {"location_id": office.id})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Bureau informatique")
        self.assertContains(response, "INV-001")
        self.assertContains(response, first_group_material.account_code)

    def test_central_excel_export_contains_summary_and_detail_sheets(self) -> None:
        response = self.client.get(reverse("central_recap_excel"))
        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(filename=BytesIO(response.content))
        self.assertIn("Synthese ministere", workbook.sheetnames)
        self.assertIn("KPI qualite", workbook.sheetnames)
        self.assertIn("Recommandations", workbook.sheetnames)
        self.assertTrue(any(name.startswith("ST01_2026") for name in workbook.sheetnames))

        summary_sheet = workbook["Synthese ministere"]
        self.assertEqual(summary_sheet[1][0].value, "Rapport central ministeriel")

        kpi_sheet = workbook["KPI qualite"]
        self.assertEqual(kpi_sheet[1][0].value, "Indicateurs qualite des donnees")
        self.assertEqual(kpi_sheet[5][0].value, "Service")
        self.assertEqual(kpi_sheet[5][7].value, "Taux PV")
        self.assertEqual(kpi_sheet[6][7].fill.fill_type, "solid")
        first_col_values = [
            kpi_sheet.cell(row=row_idx, column=1).value
            for row_idx in range(1, kpi_sheet.max_row + 1)
        ]
        self.assertIn("Legende des seuils KPI", first_col_values)
        self.assertIn("Vert", first_col_values)

        reco_sheet = workbook["Recommandations"]
        self.assertEqual(reco_sheet[1][0].value, "Plan d'actions recommande")
        self.assertEqual(reco_sheet[5][0].value, "Service")
        self.assertEqual(reco_sheet[5][5].value, "Delai suggere")
        self.assertEqual(reco_sheet[6][2].value, "Haute")

    def test_central_excel_export_works_without_common_year_across_all_services(self) -> None:
        other_structure = Structure.objects.create(name="Service Sans 2026", code="ST99", is_active=True)
        FiscalYear.objects.create(
            structure=other_structure,
            year=2025,
            start_date=date(2025, 1, 1),
            end_date=date(2025, 12, 31),
            is_active=True,
        )

        response = self.client.get(reverse("central_recap_excel"))
        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(filename=BytesIO(response.content))
        self.assertIn("Synthese ministere", workbook.sheetnames)
        self.assertTrue(any(name.startswith("ST01_2026") for name in workbook.sheetnames))

    def test_final_report_advanced_filters(self) -> None:
        response = self.client.get(
            reverse("final_report"),
            {
                "q": "papier",
                "pv": "all",
                "stock": "positive",
            },
        )
        self.assertEqual(response.status_code, 200)
        rows = response.context["balance_rows"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["account_code"], "6010")

    def test_nomenclature_add_propagates_to_all_open_services(self) -> None:
        second_structure = Structure.objects.create(name="Service B", code="ST02", is_active=True)
        second_fiscal_year = FiscalYear.objects.create(
            structure=second_structure,
            year=2026,
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
            is_active=True,
            is_closed=False,
        )

        response = self.client.post(
            reverse("nomenclature_admin"),
            {
                "action": "add",
                "account_code": "6030",
                "name": "Classeur",
                "unit": "Unite",
                "group_code": "60",
            },
            follow=False,
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(NomenclatureItem.objects.filter(account_code="6030").exists())
        self.assertTrue(
            Material.objects.filter(
                structure=self.structure,
                fiscal_year=self.fiscal_year,
                account_code="6030",
            ).exists()
        )
        self.assertTrue(
            Material.objects.filter(
                structure=second_structure,
                fiscal_year=second_fiscal_year,
                account_code="6030",
            ).exists()
        )

    def test_nomenclature_duplicate_account_code_shows_validation_error(self) -> None:
        NomenclatureItem.objects.create(
            account_code="6040",
            name="Agrafes",
            unit="Boite",
            group_code="60",
        )

        response = self.client.post(
            reverse("nomenclature_admin"),
            {
                "action": "add",
                "account_code": "6040",
                "name": "Agrafes bis",
                "unit": "Boite",
                "group_code": "60",
            },
            follow=False,
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ce code de compte existe deja dans la nomenclature.")
        self.assertEqual(NomenclatureItem.objects.filter(account_code="6040").count(), 1)

    def test_exit_voucher_uses_cmup_as_unit_price(self) -> None:
        response = self.client.post(
            reverse("voucher_create"),
            {
                "voucher_type": Voucher.TYPE_FINAL_EXIT,
                "operation_date": "2026-02-01",
                "source_or_destination": "Consommation service",
                "comments": "Test CMUP",
                "line_material_0": str(self.material_1.id),
                "line_specification_0": "Sortie test",
                "line_quantity_0": "2",
                "line_unit_price_0": "9999",
            },
            follow=False,
        )
        self.assertEqual(response.status_code, 302)

        created_voucher = Voucher.objects.get(number=3, structure=self.structure, fiscal_year=self.fiscal_year)
        self.assertEqual(created_voucher.voucher_type, Voucher.TYPE_FINAL_EXIT)
        line = created_voucher.lines.get(material=self.material_1)
        self.assertEqual(line.unit_price, Decimal("500"))
        self.assertEqual(line.amount, Decimal("1000"))

    def test_temporary_exit_does_not_decrease_existing_stock(self) -> None:
        response = self.client.post(
            reverse("voucher_create"),
            {
                "voucher_type": Voucher.TYPE_TEMP_EXIT,
                "operation_date": "2026-02-02",
                "source_or_destination": "Pret atelier A",
                "comments": "Sortie provisoire",
                "line_material_0": str(self.material_1.id),
                "line_specification_0": "Pret",
                "line_quantity_0": "3",
                "line_unit_price_0": "1",
            },
            follow=False,
        )
        self.assertEqual(response.status_code, 302)

        recap_response = self.client.get(reverse("recap"))
        self.assertEqual(recap_response.status_code, 200)
        summary = recap_response.context["summary"]
        papier_row = next(row for row in summary if row.account_code == "6010")
        self.assertEqual(papier_row.remaining_qty, Decimal("10"))
        self.assertEqual(papier_row.qty_out, Decimal("0"))
        self.assertEqual(papier_row.qty_out_provisional, Decimal("3"))

    def test_grand_ledger_shows_origin_destination_for_temporary_exit(self) -> None:
        create_response = self.client.post(
            reverse("voucher_create"),
            {
                "voucher_type": Voucher.TYPE_TEMP_EXIT,
                "operation_date": "2026-02-03",
                "source_or_destination": "Mise a disposition bureau B",
                "comments": "Sortie provisoire",
                "line_material_0": str(self.material_1.id),
                "line_specification_0": "Pret",
                "line_quantity_0": "2",
                "line_unit_price_0": "1",
            },
            follow=False,
        )
        self.assertEqual(create_response.status_code, 302)

        ledger_response = self.client.get(reverse("ledger"), {"material": "6010"})
        self.assertEqual(ledger_response.status_code, 200)
        self.assertContains(ledger_response, "Mise a disposition bureau B")

    def test_temporary_exit_status_defaults_to_en_cours(self) -> None:
        create_response = self.client.post(
            reverse("voucher_create"),
            {
                "voucher_type": Voucher.TYPE_TEMP_EXIT,
                "operation_date": "2026-02-04",
                "source_or_destination": "Atelier reparation",
                "comments": "Sortie provisoire",
                "line_material_0": str(self.material_1.id),
                "line_specification_0": "Pret reparation",
                "line_quantity_0": "1",
                "line_unit_price_0": "1",
            },
            follow=False,
        )
        self.assertEqual(create_response.status_code, 302)

        voucher = Voucher.objects.get(number=3, structure=self.structure, fiscal_year=self.fiscal_year)
        response = self.client.get(reverse("voucher_detail", kwargs={"pk": voucher.pk}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Statut retour")
        self.assertContains(response, "En cours")

    def test_can_mark_temporary_exit_as_retourne_with_date_and_observation(self) -> None:
        create_response = self.client.post(
            reverse("voucher_create"),
            {
                "voucher_type": Voucher.TYPE_TEMP_EXIT,
                "operation_date": "2026-02-05",
                "source_or_destination": "Atelier reparation",
                "comments": "Sortie provisoire",
                "line_material_0": str(self.material_1.id),
                "line_specification_0": "Pret reparation",
                "line_quantity_0": "1",
                "line_unit_price_0": "1",
            },
            follow=False,
        )
        self.assertEqual(create_response.status_code, 302)

        voucher = Voucher.objects.get(number=3, structure=self.structure, fiscal_year=self.fiscal_year)
        response = self.client.post(
            reverse("voucher_detail", kwargs={"pk": voucher.pk}),
            {
                "action": "update_temp_return_status",
                "temp_return_status": "retourne",
                "temp_return_date": "2026-02-12",
                "temp_return_note": "Retourne apres reparation",
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Statut de retour mis a jour")
        voucher.refresh_from_db()
        self.assertEqual(voucher.temp_return_status, "retourne")
        self.assertEqual(voucher.temp_return_note, "Retourne apres reparation")
        self.assertEqual(voucher.temp_return_date, date(2026, 2, 12))

    def test_carry_forward_is_blocked_without_complete_checklist(self) -> None:
        response = self.client.post(reverse("carry_forward"), follow=True)
        self.assertEqual(response.status_code, 200)
        self.fiscal_year.refresh_from_db()
        self.assertFalse(self.fiscal_year.is_closed)
        self.assertContains(response, "Checklist de cloture non satisfaite", status_code=200)

    def test_carry_forward_is_accessible_for_comptable(self) -> None:
        comptable = User.objects.create_user(username="comptable", password="pass12345")
        comptable.profile.role = "comptable"
        comptable.profile.default_structure = self.structure
        comptable.profile.save()
        comptable.profile.assigned_structures.add(self.structure)

        self.client.login(username="comptable", password="pass12345")
        session = self.client.session
        session["structure_id"] = self.structure.id
        session["fiscal_year_id"] = self.fiscal_year.id
        session.save()

        response = self.client.get(reverse("carry_forward"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Clôture de l'exercice")

    def test_carry_forward_creates_new_year_and_opening_voucher(self) -> None:
        PhysicalInventory.objects.create(
            material=self.material_1,
            pending_qty=Decimal("10"),
            in_service_qty=Decimal("0"),
            provisional_qty=Decimal("0"),
            unit_price=Decimal("500"),
            is_pv_filled=True,
        )
        PhysicalInventory.objects.create(
            material=self.material_2,
            pending_qty=Decimal("7"),
            in_service_qty=Decimal("0"),
            provisional_qty=Decimal("0"),
            unit_price=Decimal("300"),
            is_pv_filled=True,
        )

        response = self.client.post(reverse("carry_forward"), follow=False)
        self.assertEqual(response.status_code, 302)

        self.fiscal_year.refresh_from_db()
        self.assertTrue(self.fiscal_year.is_closed)

        new_fy = FiscalYear.objects.get(structure=self.structure, year=2027)
        opening_voucher = Voucher.objects.get(structure=self.structure, fiscal_year=new_fy, number=1)
        self.assertEqual(opening_voucher.voucher_type, Voucher.TYPE_ENTRY)
        self.assertEqual(opening_voucher.lines.count(), 2)

    def test_voucher_list_filters_by_number(self) -> None:
        response = self.client.get(reverse("voucher_list"), {"q": "2"})
        self.assertEqual(response.status_code, 200)
        vouchers = list(response.context["vouchers"].object_list)
        self.assertEqual(len(vouchers), 1)
        self.assertEqual(vouchers[0].number, 2)

    def test_voucher_list_filters_by_temporary_return_status(self) -> None:
        provisional_en_cours = Voucher.objects.create(
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            number=3,
            voucher_type=Voucher.TYPE_TEMP_EXIT,
            operation_date=date(2026, 2, 20),
            source_or_destination="Pret A",
            comments="Sortie provisoire en cours",
            created_by=self.user,
        )
        VoucherLine.objects.create(
            voucher=provisional_en_cours,
            material=self.material_1,
            material_name=self.material_1.name,
            specification="Pret",
            quantity=Decimal("1"),
            unit=self.material_1.unit,
            unit_price=Decimal("500"),
            amount=Decimal("500"),
        )

        provisional_retourne = Voucher.objects.create(
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            number=4,
            voucher_type=Voucher.TYPE_TEMP_EXIT,
            operation_date=date(2026, 2, 21),
            source_or_destination="Pret B",
            comments="Sortie provisoire retournee",
            temp_return_date=date(2026, 2, 28),
            temp_return_note="Retourne apres reparation",
            created_by=self.user,
        )
        VoucherLine.objects.create(
            voucher=provisional_retourne,
            material=self.material_1,
            material_name=self.material_1.name,
            specification="Pret",
            quantity=Decimal("1"),
            unit=self.material_1.unit,
            unit_price=Decimal("500"),
            amount=Decimal("500"),
        )

        response_en_cours = self.client.get(reverse("voucher_list"), {"temp_return_status": "en_cours"})
        self.assertEqual(response_en_cours.status_code, 200)
        vouchers_en_cours = list(response_en_cours.context["vouchers"].object_list)
        self.assertEqual(len(vouchers_en_cours), 1)
        self.assertEqual(vouchers_en_cours[0].id, provisional_en_cours.id)

        response_retourne = self.client.get(reverse("voucher_list"), {"temp_return_status": "retourne"})
        self.assertEqual(response_retourne.status_code, 200)
        vouchers_retourne = list(response_retourne.context["vouchers"].object_list)
        self.assertEqual(len(vouchers_retourne), 1)
        self.assertEqual(vouchers_retourne[0].id, provisional_retourne.id)

    def test_materials_pagination_and_filter(self) -> None:
        for index in range(3, 36):
            Material.objects.create(
                structure=self.structure,
                fiscal_year=self.fiscal_year,
                account_code=f"7{index:03d}",
                name=f"Article {index}",
                unit="U",
                group_code="70",
            )

        response_page_2 = self.client.get(reverse("materials"), {"page": 2})
        self.assertEqual(response_page_2.status_code, 200)
        self.assertEqual(response_page_2.context["materials"].number, 2)

        response_filter = self.client.get(reverse("materials"), {"q": "Encre"})
        self.assertEqual(response_filter.status_code, 200)
        results = list(response_filter.context["materials"].object_list)
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0].account_code, "6020")

    def test_admin_can_view_document_security_logs_with_filters(self) -> None:
        AuditLog.objects.create(
            user=self.user,
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            action=AuditLog.ACTION_UPDATE,
            entity="DocumentSecurity",
            entity_id="file1.pdf",
            description="Rejet upload document: mime_not_allowed",
            metadata={
                "reason": "mime_not_allowed",
                "filename": "file1.pdf",
                "content_type": "application/x-msdownload",
                "size": 123,
                "ip": "127.0.0.1",
            },
        )
        AuditLog.objects.create(
            user=self.user,
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            action=AuditLog.ACTION_UPDATE,
            entity="DocumentSecurity",
            entity_id="file2.pdf",
            description="Rejet upload document: size_exceeded",
            metadata={
                "reason": "size_exceeded",
                "filename": "file2.pdf",
                "content_type": "application/pdf",
                "size": 99999999,
                "ip": "10.0.0.8",
            },
        )

        response = self.client.get(reverse("document_security_logs"), {"reason": "mime_not_allowed"})
        self.assertEqual(response.status_code, 200)
        logs = list(response.context["logs"].object_list)
        self.assertEqual(len(logs), 1)
        self.assertEqual(logs[0].metadata.get("reason"), "mime_not_allowed")

    def test_non_admin_cannot_view_document_security_logs(self) -> None:
        comptable = User.objects.create_user(username="comptable_logs", password="pass12345")
        comptable.profile.role = "comptable"
        comptable.profile.default_structure = self.structure
        comptable.profile.save()
        comptable.profile.assigned_structures.add(self.structure)

        self.client.login(username="comptable_logs", password="pass12345")
        session = self.client.session
        session["structure_id"] = self.structure.id
        session["fiscal_year_id"] = self.fiscal_year.id
        session.save()

        response = self.client.get(reverse("document_security_logs"), follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Seuls les admins peuvent consulter ce journal.")

    def test_admin_can_export_document_security_logs_csv(self) -> None:
        AuditLog.objects.create(
            user=self.user,
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            action=AuditLog.ACTION_UPDATE,
            entity="DocumentSecurity",
            entity_id="file3.pdf",
            description="Rejet upload document: extension_not_allowed",
            metadata={
                "reason": "extension_not_allowed",
                "filename": "file3.pdf",
                "content_type": "application/octet-stream",
                "size": 42,
                "ip": "127.0.0.1",
            },
        )

        response = self.client.get(reverse("document_security_logs_csv"), {"reason": "extension_not_allowed"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        self.assertIn("attachment; filename=\"journal_securite_documents.csv\"", response["Content-Disposition"])

        content = response.content.decode("utf-8")
        self.assertIn("Date;Utilisateur;Raison;Fichier;MIME;Taille;IP", content)
        self.assertIn("extension_not_allowed", content)
        self.assertIn("file3.pdf", content)

    def test_document_security_logs_csv_escapes_formula_like_values(self) -> None:
        AuditLog.objects.create(
            user=self.user,
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            action=AuditLog.ACTION_UPDATE,
            entity="DocumentSecurity",
            entity_id="file_formula.csv",
            description="Rejet upload document: mime_not_allowed",
            metadata={
                "reason": "mime_not_allowed",
                "filename": "=2+2",
                "content_type": "text/csv",
                "size": 1,
                "ip": "127.0.0.1",
            },
        )

        response = self.client.get(reverse("document_security_logs_csv"))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        self.assertIn("'=2+2", content)

    def test_non_admin_cannot_export_document_security_logs_csv(self) -> None:
        comptable = User.objects.create_user(username="comptable_logs_csv", password="pass12345")
        comptable.profile.role = "comptable"
        comptable.profile.default_structure = self.structure
        comptable.profile.save()
        comptable.profile.assigned_structures.add(self.structure)

        self.client.login(username="comptable_logs_csv", password="pass12345")
        session = self.client.session
        session["structure_id"] = self.structure.id
        session["fiscal_year_id"] = self.fiscal_year.id
        session.save()

        response = self.client.get(reverse("document_security_logs_csv"), follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Seuls les admins peuvent exporter ce journal.")

    def test_voucher_detail_restricted_to_current_structure(self) -> None:
        other_structure = Structure.objects.create(name="Service Autre", code="ST02", is_active=True)
        other_fy = FiscalYear.objects.create(
            structure=other_structure,
            year=2026,
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
            is_active=True,
        )
        other_material = Material.objects.create(
            structure=other_structure,
            fiscal_year=other_fy,
            account_code="6090",
            name="Autre",
            unit="U",
            group_code="60",
        )
        other_voucher = Voucher.objects.create(
            structure=other_structure,
            fiscal_year=other_fy,
            number=1,
            voucher_type=Voucher.TYPE_ENTRY,
            operation_date=date(2026, 4, 1),
            source_or_destination="Fournisseur X",
            comments="Autre",
            created_by=self.user,
        )
        VoucherLine.objects.create(
            voucher=other_voucher,
            material=other_material,
            material_name=other_material.name,
            specification="Lot",
            quantity=Decimal("1"),
            unit=other_material.unit,
            unit_price=Decimal("100"),
            amount=Decimal("100"),
        )

        response = self.client.get(reverse("voucher_detail", args=[other_voucher.pk]))
        self.assertEqual(response.status_code, 404)

    def test_document_download_requires_existing_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            media_root = Path(tmp_dir)
            documents_dir = media_root / "documents"
            documents_dir.mkdir(parents=True, exist_ok=True)
            file_path = documents_dir / "guide.pdf"
            file_path.write_bytes(b"test document")

            with override_settings(MEDIA_ROOT=media_root):
                response = self.client.get(reverse("document_download", args=["guide.pdf"]))
                self.assertEqual(response.status_code, 200)
                self.assertIn("attachment;", response["Content-Disposition"])
                self.assertIn("guide.pdf", response["Content-Disposition"])
                response.close()

    def test_screen_admin_shows_recent_security_rejections_badge(self) -> None:
        AuditLog.objects.create(
            user=self.user,
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            action=AuditLog.ACTION_UPDATE,
            entity="DocumentSecurity",
            entity_id="file4.pdf",
            description="Rejet upload document: mime_not_allowed",
            metadata={
                "reason": "mime_not_allowed",
                "filename": "file4.pdf",
                "content_type": "application/x-msdownload",
                "size": 10,
                "ip": "127.0.0.1",
            },
        )

        response = self.client.get(reverse("screen_admin"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Journal de sécurité")
        self.assertContains(response, "1 / 24h")

    def test_admin_can_purge_old_document_security_logs(self) -> None:
        old_log = AuditLog.objects.create(
            user=self.user,
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            action=AuditLog.ACTION_UPDATE,
            entity="DocumentSecurity",
            entity_id="old.pdf",
            description="Rejet ancien",
            metadata={"reason": "size_exceeded", "ip": "127.0.0.1"},
        )
        recent_log = AuditLog.objects.create(
            user=self.user,
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            action=AuditLog.ACTION_UPDATE,
            entity="DocumentSecurity",
            entity_id="recent.pdf",
            description="Rejet recent",
            metadata={"reason": "mime_not_allowed", "ip": "127.0.0.1"},
        )
        AuditLog.objects.filter(pk=old_log.pk).update(created_at=timezone.now() - timezone.timedelta(days=400))
        AuditLog.objects.filter(pk=recent_log.pk).update(created_at=timezone.now() - timezone.timedelta(days=10))

        response = self.client.post(reverse("document_security_logs_purge"), {"retention_days": "180"}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Purge terminee")
        self.assertFalse(AuditLog.objects.filter(pk=old_log.pk).exists())
        self.assertTrue(AuditLog.objects.filter(pk=recent_log.pk).exists())

    def test_non_admin_cannot_purge_document_security_logs(self) -> None:
        comptable = User.objects.create_user(username="comptable_logs_purge", password="pass12345")
        comptable.profile.role = "comptable"
        comptable.profile.default_structure = self.structure
        comptable.profile.save()
        comptable.profile.assigned_structures.add(self.structure)

        self.client.login(username="comptable_logs_purge", password="pass12345")
        session = self.client.session
        session["structure_id"] = self.structure.id
        session["fiscal_year_id"] = self.fiscal_year.id
        session.save()

        response = self.client.post(reverse("document_security_logs_purge"), {"retention_days": "180"}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Seuls les admins peuvent purger ce journal.")

    def test_purge_rejects_invalid_retention_days(self) -> None:
        response = self.client.post(reverse("document_security_logs_purge"), {"retention_days": "0"}, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "La retention doit etre &gt;= 1 jour.")

    @override_settings(DOCUMENT_SECURITY_LOG_RETENTION_DAYS=45)
    def test_document_security_logs_uses_configured_default_retention_days(self) -> None:
        response = self.client.get(reverse("document_security_logs"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="retention_days"')
        self.assertContains(response, 'value="45"')

    @override_settings(DOCUMENT_SECURITY_LOG_RETENTION_DAYS=180)
    def test_management_command_purge_document_security_logs(self) -> None:
        old_log = AuditLog.objects.create(
            user=self.user,
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            action=AuditLog.ACTION_UPDATE,
            entity="DocumentSecurity",
            entity_id="old-cmd.pdf",
            description="Old log",
            metadata={"reason": "size_exceeded"},
        )
        recent_log = AuditLog.objects.create(
            user=self.user,
            structure=self.structure,
            fiscal_year=self.fiscal_year,
            action=AuditLog.ACTION_UPDATE,
            entity="DocumentSecurity",
            entity_id="recent-cmd.pdf",
            description="Recent log",
            metadata={"reason": "mime_not_allowed"},
        )
        AuditLog.objects.filter(pk=old_log.pk).update(created_at=timezone.now() - timezone.timedelta(days=250))
        AuditLog.objects.filter(pk=recent_log.pk).update(created_at=timezone.now() - timezone.timedelta(days=5))

        out = StringIO()
        call_command("purge_document_security_logs", stdout=out)
        output = out.getvalue()

        self.assertIn("Purged", output)
        self.assertFalse(AuditLog.objects.filter(pk=old_log.pk).exists())
        self.assertTrue(AuditLog.objects.filter(pk=recent_log.pk).exists())
        self.assertTrue(
            AuditLog.objects.filter(
                entity="DocumentSecurity",
                entity_id="bulk_purge_command",
                action=AuditLog.ACTION_DELETE,
            ).exists()
        )

    def test_central_cache_version_is_invalidated_after_voucher_create(self) -> None:
        cache.set("central_recap:version", 1, timeout=None)

        response = self.client.post(
            reverse("voucher_create"),
            {
                "voucher_type": Voucher.TYPE_ENTRY,
                "operation_date": "2026-02-10",
                "source_or_destination": "Fournisseur C",
                "comments": "Cache invalidation",
                "line_material_0": str(self.material_1.id),
                "line_specification_0": "Lot cache",
                "line_quantity_0": "1",
                "line_unit_price_0": "100",
            },
            follow=False,
        )
        self.assertEqual(response.status_code, 302)
        self.assertGreater(int(cache.get("central_recap:version") or 0), 1)

    def test_voucher_detail_entry_shows_quantity_total_and_words(self) -> None:
        voucher = Voucher.objects.get(number=1, structure=self.structure, fiscal_year=self.fiscal_year)

        response = self.client.get(reverse("voucher_detail", kwargs={"pk": voucher.pk}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "<td colspan=\"3\">Total</td>", html=False)
        self.assertContains(response, "10")
        self.assertContains(response, "(dix)")
        self.assertContains(response, "(cinq mille)")

    def test_voucher_detail_final_exit_shows_words_in_declaration(self) -> None:
        create_response = self.client.post(
            reverse("voucher_create"),
            {
                "voucher_type": Voucher.TYPE_FINAL_EXIT,
                "operation_date": "2026-02-15",
                "source_or_destination": "Consommation test",
                "comments": "Sortie pour declaration",
                "line_material_0": str(self.material_1.id),
                "line_specification_0": "Sortie test",
                "line_quantity_0": "2",
                "line_unit_price_0": "1",
            },
            follow=False,
        )
        self.assertEqual(create_response.status_code, 302)

        voucher = Voucher.objects.get(number=3, structure=self.structure, fiscal_year=self.fiscal_year)
        response = self.client.get(reverse("voucher_detail", kwargs={"pk": voucher.pk}))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "diminuer ses prises en charge")
        self.assertContains(response, "(deux)")
        self.assertContains(response, "(mille)")

    def test_superuser_can_reopen_closed_fiscal_year_from_select_structure(self) -> None:
        closed_fy = FiscalYear.objects.create(
            structure=self.structure,
            year=2025,
            start_date=date(2025, 1, 1),
            end_date=date(2025, 12, 31),
            is_active=True,
            is_closed=True,
        )

        superuser = User.objects.create_superuser(username="root", email="root@example.com", password="pass12345")
        self.client.login(username="root", password="pass12345")

        session = self.client.session
        session["structure_id"] = self.structure.id
        session["fiscal_year_id"] = self.fiscal_year.id
        session.save()

        response = self.client.post(
            reverse("select_structure"),
            {
                "action": "reopen_fiscal_year",
                "structure_id": str(self.structure.id),
                "reopen_fiscal_year_id": str(closed_fy.id),
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Exercice 2025 reouvert avec succes.")
        closed_fy.refresh_from_db()
        self.assertFalse(closed_fy.is_closed)

    def test_non_superuser_cannot_reopen_closed_fiscal_year(self) -> None:
        closed_fy = FiscalYear.objects.create(
            structure=self.structure,
            year=2025,
            start_date=date(2025, 1, 1),
            end_date=date(2025, 12, 31),
            is_active=True,
            is_closed=True,
        )

        response = self.client.post(
            reverse("select_structure"),
            {
                "action": "reopen_fiscal_year",
                "structure_id": str(self.structure.id),
                "reopen_fiscal_year_id": str(closed_fy.id),
            },
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Seul un superuser peut reouvrir un exercice cloture.")
        closed_fy.refresh_from_db()
        self.assertTrue(closed_fy.is_closed)


class DocumentManagementTests(TestCase):
    def setUp(self) -> None:
        self.media_root = tempfile.mkdtemp(prefix="compta_media_tests_")
        self.media_override = override_settings(MEDIA_ROOT=self.media_root)
        self.media_override.enable()

        self.admin_user = User.objects.create_user(username="admin_doc", password="pass12345")
        self.admin_user.profile.role = "admin"
        self.admin_user.profile.save()

        self.structure = Structure.objects.create(name="Service Doc", code="SD01", is_active=True)
        self.fiscal_year = FiscalYear.objects.create(
            structure=self.structure,
            year=2026,
            start_date=date(2026, 1, 1),
            end_date=date(2026, 12, 31),
            is_active=True,
        )

    def tearDown(self) -> None:
        self.media_override.disable()
        shutil.rmtree(self.media_root, ignore_errors=True)

    def _set_active_context(self) -> None:
        session = self.client.session
        session["structure_id"] = self.structure.id
        session["fiscal_year_id"] = self.fiscal_year.id
        session.save()

    def test_admin_can_upload_and_delete_document(self) -> None:
        self.client.login(username="admin_doc", password="pass12345")
        self._set_active_context()

        upload = SimpleUploadedFile("modele.pdf", b"sample content", content_type="application/pdf")
        upload_response = self.client.post(reverse("screen_documents"), {"document": upload}, follow=True)
        self.assertEqual(upload_response.status_code, 200)

        uploaded_path = Path(self.media_root) / "documents" / "modele.pdf"
        self.assertTrue(uploaded_path.exists())

        delete_response = self.client.post(reverse("document_delete"), {"filename": "modele.pdf"}, follow=True)
        self.assertEqual(delete_response.status_code, 200)
        self.assertFalse(uploaded_path.exists())

    def test_non_admin_cannot_upload_or_delete_document(self) -> None:
        comptable = User.objects.create_user(username="compte_doc", password="pass12345")
        comptable.profile.role = "comptable"
        comptable.profile.default_structure = self.structure
        comptable.profile.save()
        comptable.profile.assigned_structures.add(self.structure)

        self.client.login(username="compte_doc", password="pass12345")
        self._set_active_context()

        upload = SimpleUploadedFile("interdit.pdf", b"sample", content_type="application/pdf")
        upload_response = self.client.post(reverse("screen_documents"), {"document": upload}, follow=True)
        self.assertEqual(upload_response.status_code, 200)
        self.assertContains(upload_response, "Seuls les admins peuvent ajouter des documents.")
        self.assertFalse((Path(self.media_root) / "documents" / "interdit.pdf").exists())

        docs_path = Path(self.media_root) / "documents"
        docs_path.mkdir(parents=True, exist_ok=True)
        existing_path = docs_path / "manuel.pdf"
        existing_path.write_bytes(b"existing")

        delete_response = self.client.post(reverse("document_delete"), {"filename": "manuel.pdf"}, follow=True)
        self.assertEqual(delete_response.status_code, 200)
        self.assertContains(delete_response, "Seuls les admins peuvent supprimer des documents.")
        self.assertTrue(existing_path.exists())

    def test_document_upload_rejects_file_larger_than_limit(self) -> None:
        self.client.login(username="admin_doc", password="pass12345")
        self._set_active_context()

        oversized = SimpleUploadedFile(
            "gros.pdf",
            b"a" * (10 * 1024 * 1024 + 1),
            content_type="application/pdf",
        )
        response = self.client.post(reverse("screen_documents"), {"document": oversized}, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Fichier trop volumineux")
        self.assertFalse((Path(self.media_root) / "documents" / "gros.pdf").exists())

    def test_document_upload_rejects_disallowed_mime_type(self) -> None:
        self.client.login(username="admin_doc", password="pass12345")
        self._set_active_context()

        bad_mime = SimpleUploadedFile(
            "doc.pdf",
            b"fake",
            content_type="application/x-msdownload",
        )
        response = self.client.post(reverse("screen_documents"), {"document": bad_mime}, follow=True)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Type MIME non autorisé")
        self.assertFalse((Path(self.media_root) / "documents" / "doc.pdf").exists())

    def test_document_rejection_is_audited(self) -> None:
        self.client.login(username="admin_doc", password="pass12345")
        self._set_active_context()

        bad_mime = SimpleUploadedFile(
            "intrus.pdf",
            b"fake",
            content_type="application/x-msdownload",
        )
        response = self.client.post(reverse("screen_documents"), {"document": bad_mime}, follow=True)

        self.assertEqual(response.status_code, 200)
        rejection_log = AuditLog.objects.filter(entity="DocumentSecurity").order_by("-created_at").first()
        self.assertIsNotNone(rejection_log)
        self.assertEqual(rejection_log.action, AuditLog.ACTION_UPDATE)
        self.assertEqual(rejection_log.metadata.get("reason"), "mime_not_allowed")
